"""Variant numbers, the scoring rule (same as the web page) and the Excel export."""
import io, math

from .courses import keys


def variant(vset, st):
    """Variant of a student: set*100 + 3*(team-1) + member for members 1-3; set*100 + 18 + team for member 4."""
    if not vset or not st.get('team') or not st.get('member'):
        return None
    t, m = int(st['team']), int(st['member'])
    return int(vset) * 100 + (3 * (t - 1) + m if m <= 3 else 18 + t)


def rule_of(K, q):
    """Tolerance of one question: its own {full, half} (5th item of the question) or the course rule."""
    return q[4] if len(q) > 4 and isinstance(q[4], dict) else K['rule']


def qvalue(q, values, i):
    """The number a question is marked on. Usually the submitted value; a question with {'calc': 'rc_fc', 'from': j} is
    marked on fc = 1 / (2 pi R C) in kHz, from R (ohm, question j) and C (nF, this question): Lab 2 Task 1 (design)."""
    x = float(values[i])
    r = q[4] if len(q) > 4 and isinstance(q[4], dict) else {}
    if r.get('calc') == 'rc_fc':
        R = float(values[r['from']])
        return 1 / (2 * math.pi * R * x * 1e-9) / 1e3
    return x


def score(C, lab, st, K=None):
    """Marks of one student for one lab. C = one class as returned by db.load(). K = keys pack (with overrides)."""
    K = K or keys(C['course']); lab = str(lab); L = K['labs'][lab]; v = variant(C['vset'], st)
    rec = C['labs'].get(lab, {}).get(st['sid'], {})
    if rec.get('override') is not None:          # lecturer's manual score wins
        return float(rec['override'])
    total = 0.0
    if v and str(v) in L['variants'] and rec.get('tick'):
        ans = L['variants'][str(v)]['keys']
        for i, q in enumerate(L['questions']):
            if not q[3]:                         # input only (no marks), e.g. R of Lab 2 Task 1
                continue
            try:
                x = qvalue(q, rec.get('values') or [None] * len(ans), i)
            except (TypeError, ValueError, IndexError, ZeroDivisionError):
                continue
            err = abs(x / ans[i] - 1)
            rule = rule_of(K, q)
            total += q[3] if err <= rule['full'] else (q[3] / 2 if err <= rule['half'] else 0)
    for m in L['manual']:
        if m[3] == 'student':
            total += float((rec.get('manual') or {}).get(m[0], 0) or 0)
        elif st.get('team') and rec.get('tick'):     # a team mark counts only for a member the lecturer ticked (was there)
            total += float(C.get('team', {}).get(lab, {}).get(str(st['team']), {}).get(m[0], 0) or 0)
    return total


def export_xlsx(classes, keypacks=None):
    """One sheet per class. keypacks = {course: keys_live(...)} when overrides exist."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    for C in classes:
        K = (keypacks or {}).get(C['course']) or keys(C['course']); labs = list(K['labs']); full = sum(K['labs'][l]['total'] for l in labs)
        ws = wb.create_sheet(f"{C['name']} {C['term']}"[:30].replace('/', '-'))
        ws.append(['MSSV', 'Họ tên', 'Nhóm', 'TV', 'Mã đề'] + [f'Lab {l} (/{K["labs"][l]["total"]})' for l in labs] + [f'Tổng /{full:g}', 'Quy về /10'])
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='549E39'); c.alignment = Alignment(horizontal='center', wrap_text=True)
        for st in sorted(C['students'], key=lambda s: (int(s.get('team') or 99), int(s.get('member') or 9), s['name'])):
            sc = [score(C, l, st, K) for l in labs]
            ws.append([st['sid'], st['name'], st.get('team'), st.get('member'), variant(C['vset'], st)] + sc + [sum(sc), round(sum(sc) / full * 10, 2)])
        for col, w in zip('ABCDEFGHIJK', [14, 28, 7, 5, 8, 11, 11, 11, 11, 10, 12]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'C2'
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
