"""HTTP handler: static files + JSON API (/api/...). Runs on the laptop (server.py) and on Vercel (api/index.py)."""
import base64, http.server, io, json, os, re, socketserver, urllib.parse

from .config import ASSETS, CLASSES, KEYS, LOCK, PORT, STATIC
from .db import (allow, add_hit, as_int, assign_random, connect, count_hits, get_pin, h, h_student, load, new_session,
                 new_student_session, now, reset_student_pin, set_config, set_gradebook, set_sync, set_team_mark, upsert_student)
from .network import lan_ips, qr_svg
from .scoring import export_xlsx, score, variant

PUBLIC_URL = os.environ.get('PUBLIC_URL', '').rstrip('/')
ON_VERCEL = bool(os.environ.get('VERCEL'))
TOO_MANY = {'error': 'Too many tries. Wait a few minutes.', 'error_vi': 'Thử quá nhiều lần. Đợi vài phút.'}


class H(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *a, **k):
        super().__init__(*a, directory=STATIC, **k)

    def log_message(self, *a):
        pass

    # ---------------- helpers ----------------
    def ip(self):
        fwd = self.headers.get('X-Forwarded-For') if ON_VERCEL else None   # only trust the header behind Vercel's proxy
        return (fwd.split(',')[0].strip() if fwd else self.client_address[0]) or '?'

    def https(self):
        return ON_VERCEL or self.headers.get('X-Forwarded-Proto') == 'https'

    def cookie(self, name):
        c = self.headers.get('Cookie', '')
        return dict(p.strip().split('=', 1) for p in c.split(';') if '=' in p).get(name)

    def authed(self, con):
        tok = self.cookie('tok')
        return bool(tok) and con.execute('SELECT 1 FROM sessions WHERE token=?', (tok,)).fetchone() is not None

    def student(self, con):
        """The signed-in student (row with cls, sid, name, team, member) or None."""
        tok = self.cookie('stu')
        if not tok:
            return None
        r = con.execute('SELECT s.cls, s.sid, s.name, s.team, s.member FROM stu_sessions x JOIN students s ON s.cls=x.cls AND s.sid=x.sid '
                        'WHERE x.token=?', (tok,)).fetchone()
        return dict(r) if r else None

    def set_cookie(self, name, value, age, same='Strict'):
        flags = f'Path=/; Max-Age={age}; SameSite={same}; HttpOnly' + ('; Secure' if self.https() else '')
        return f'{name}={value}; {flags}'

    def send_json(self, obj, code=200, cookie=None):
        b = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        if cookie:
            self.send_header('Set-Cookie', cookie)
        self.send_header('Content-Length', str(len(b)))
        self.end_headers(); self.wfile.write(b)

    def send_bytes(self, b, ctype, cache='private, max-age=3600'):
        self.send_response(200); self.send_header('Content-Type', ctype); self.send_header('Cache-Control', cache)
        self.send_header('Content-Length', str(len(b))); self.end_headers(); self.wfile.write(b)

    def base_url(self):
        if PUBLIC_URL:
            return PUBLIC_URL
        if ON_VERCEL:
            return 'https://' + (self.headers.get('X-Forwarded-Host') or self.headers.get('Host') or '')
        ips = lan_ips()
        return f'http://{ips[0]}:{PORT}' if ips else ''

    # ---------------- GET ----------------
    def do_GET(self):
        path = urllib.parse.urlparse(self.path).path
        if path.startswith('/api/public/'):     # student page – the student's own information only
            return self.public_get(path)
        if path == '/api/qr.svg':               # QR code of the page address (no secret inside)
            target = '/student/' if 'student' in urllib.parse.urlparse(self.path).query else ''
            base = self.base_url()
            if not base:
                return self.send_json({'error': 'No Wi-Fi address found on this computer', 'error_vi': 'Không tìm thấy địa chỉ wifi của máy'}, 404)
            try:
                b = qr_svg(base + target)
            except ImportError:
                return self.send_json({'error': 'The segno library is missing. Start the app with start.command to install it.', 'error_vi': 'Thiếu thư viện segno. Chạy bằng start.command để tự cài.'}, 500)
            self.send_response(200)
            self.send_header('Content-Type', 'image/svg+xml'); self.send_header('Cache-Control', 'no-store')
            self.send_header('X-Url', base + target); self.send_header('Content-Length', str(len(b)))
            self.end_headers(); self.wfile.write(b); return
        if path.startswith('/api/'):
            with connect() as con:
                if path == '/api/status':
                    return self.send_json({'hasPin': bool(get_pin(con)), 'authed': self.authed(con), 'pinFromServer': bool(os.environ.get('LECTURER_PIN'))})
                if not self.authed(con):
                    return self.send_json({'error': 'auth'}, 401)
                if path == '/api/keys':
                    return self.send_json(KEYS)
                m = re.fullmatch(r'/api/photo/(\d+)\.jpg', path)
                if m:
                    r = con.execute('SELECT data FROM photos WHERE id=?', (int(m.group(1)),)).fetchone()
                    return self.send_bytes(bytes(r['data']), 'image/jpeg') if r and r['data'] else self.send_json({'error': 'not found'}, 404)
                m = re.fullmatch(r'/api/lab3/(\d+)\.png', path)
                if m:                           # scope screen of a Lab 3 variant (lecturer: always visible)
                    f = os.path.join(ASSETS, 'lab3', f'v{m.group(1)}.png')
                    return self.send_bytes(open(f, 'rb').read(), 'image/png') if os.path.exists(f) else self.send_json({'error': 'not found'}, 404)
                if path == '/api/data':
                    return self.send_json(load(con))
                if path == '/api/export.xlsx':
                    try:
                        b = export_xlsx(load(con))
                    except ImportError:
                        return self.send_json({'error': 'The openpyxl library is missing. Start the app with start.command to install it.', 'error_vi': 'Thiếu thư viện openpyxl. Chạy bằng start.command để tự cài.'}, 500)
                    self.send_response(200)
                    self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                    self.send_header('Content-Disposition', 'attachment; filename="Diem_Electric_Circuit_Lab_HCM.xlsx"')
                    self.send_header('Content-Length', str(len(b)))
                    self.end_headers(); self.wfile.write(b); return
                return self.send_json({'error': 'not found'}, 404)
        if path == '/':
            self.path = '/index.html'
        return super().do_GET()

    # ---------------- POST ----------------
    def do_POST(self):
        path = urllib.parse.urlparse(self.path).path
        n = int(self.headers.get('Content-Length', '0') or 0)
        if n > 6_000_000:
            return self.send_json({'error': 'too large'}, 413)
        try:
            body = json.loads(self.rfile.read(n) or b'{}')
        except json.JSONDecodeError:
            return self.send_json({'error': 'bad json'}, 400)
        with LOCK, connect() as con:
            if path.startswith('/api/public/'):
                return self.public_post(con, path, body)
            if path == '/api/setpin':
                if os.environ.get('LECTURER_PIN'):
                    return self.send_json({'error': 'The PIN is set on the server.', 'error_vi': 'PIN được đặt trên server.'}, 403)
                if get_pin(con) and not self.authed(con):
                    return self.send_json({'error': 'The PIN is already set', 'error_vi': 'PIN đã được đặt'}, 403)
                pin = str(body.get('pin', ''))
                if len(pin) < 8:
                    return self.send_json({'error': 'The PIN needs at least 8 characters', 'error_vi': 'PIN cần ít nhất 8 ký tự'}, 400)
                set_config(con, 'pin', h(pin))
                return self.send_json({'ok': True}, cookie=self.set_cookie('tok', new_session(con), 15552000))
            if path == '/api/login':
                key = 'login-fail:' + self.ip()
                if count_hits(con, key, 900) >= 8:          # 8 wrong PINs -> wait 15 minutes
                    return self.send_json(TOO_MANY, 429)
                if get_pin(con) and h(str(body.get('pin', ''))) == get_pin(con):
                    return self.send_json({'ok': True}, cookie=self.set_cookie('tok', new_session(con), 15552000))
                add_hit(con, key)
                return self.send_json({'error': 'Wrong PIN', 'error_vi': 'Sai PIN'}, 403)
            if not self.authed(con):
                return self.send_json({'error': 'auth'}, 401)
            if path == '/api/logout':
                con.execute('DELETE FROM sessions WHERE token=?', (self.cookie('tok'),))
                return self.send_json({'ok': True}, cookie=self.set_cookie('tok', '', 0))
            if path == '/api/readsheet':        # .xlsx -> rows (first sheet), for the roster upload
                try:
                    from openpyxl import load_workbook
                except ImportError:
                    return self.send_json({'error': 'openpyxl is missing. Start with start.command, or save the file as CSV.', 'error_vi': 'Thiếu openpyxl. Chạy bằng start.command, hoặc lưu file dạng CSV.'}, 500)
                try:
                    wb = load_workbook(io.BytesIO(base64.b64decode(body['b64'])), read_only=True, data_only=True)
                    rows = [['' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)) for v in r]
                            for r in wb.worksheets[0].iter_rows(values_only=True)]
                except Exception as e:
                    return self.send_json({'error': f'Could not read the file: {e}', 'error_vi': f'Không đọc được file: {e}'}, 400)
                return self.send_json({'rows': [r for r in rows if any(x.strip() for x in r)][:2000]})
            cls = str(body.get('cls', ''))
            if cls not in CLASSES:
                return self.send_json({'error': 'class'}, 400)
            if path == '/api/class':            # name + full roster (replaces the list)
                if 'name' in body:
                    con.execute('UPDATE classes SET name=? WHERE cls=?', (str(body['name'])[:40], cls))
                if 'students' in body:
                    keep = []
                    for s in body['students']:
                        sid = str(s.get('sid', '')).strip()
                        if not sid: continue
                        keep.append(sid)
                        upsert_student(con, cls, sid, str(s.get('name', '')).strip(), as_int(s.get('team')), as_int(s.get('member')))
                    q = ','.join('?' * len(keep))
                    con.execute(f'DELETE FROM students WHERE cls=? AND sid NOT IN ({q})' if keep else 'DELETE FROM students WHERE cls=?', (cls, *keep))
            elif path == '/api/roster_merge':   # add / update students, keep existing team & member
                if body.get('name'):
                    con.execute('UPDATE classes SET name=? WHERE cls=?', (str(body['name'])[:40], cls))
                for s in body.get('students', []):
                    sid = str(s.get('sid', '')).strip()
                    if not sid: continue
                    old = con.execute('SELECT team, member FROM students WHERE cls=? AND sid=?', (cls, sid)).fetchone()
                    team = as_int(s.get('team')) or (old['team'] if old else None)
                    member = as_int(s.get('member')) or (old['member'] if old else None)
                    upsert_student(con, cls, sid, str(s.get('name', '')).strip(), team, member)
                assign_random(con, cls)
            elif path == '/api/shuffle':        # re-draw all teams of a class
                assign_random(con, cls, reshuffle=True)
            elif path == '/api/student_pin_reset':   # student forgot the PIN: they choose a new one at the next sign-in
                reset_student_pin(con, cls, str(body.get('sid', '')))
            elif path == '/api/record':         # one student's record for one lab
                lab = str(body['lab']); sid = str(body['sid'])
                con.execute('INSERT INTO records (cls, lab, sid, tick, manual) VALUES (?,?,?,0,?) ON CONFLICT DO NOTHING', (cls, lab, sid, '{}'))
                if 'tick' in body: con.execute('UPDATE records SET tick=? WHERE cls=? AND lab=? AND sid=?', (1 if body['tick'] else 0, cls, lab, sid))
                if 'values' in body: con.execute('UPDATE records SET vals=? WHERE cls=? AND lab=? AND sid=?', (json.dumps(body['values']), cls, lab, sid))
                if 'manual' in body: con.execute('UPDATE records SET manual=? WHERE cls=? AND lab=? AND sid=?', (json.dumps(body['manual']), cls, lab, sid))
                if 'src' in body: con.execute('UPDATE records SET src=? WHERE cls=? AND lab=? AND sid=?', (body['src'], cls, lab, sid))
                if 'override' in body:          # manual score (None = back to the automatic score)
                    ov = body['override']
                    if ov is not None:
                        ov = float(ov)
                        if not 0 <= ov <= KEYS['labs'][lab]['total']:
                            return self.send_json({'error': f"Score must be 0–{KEYS['labs'][lab]['total']}", 'error_vi': f"Điểm phải từ 0 đến {KEYS['labs'][lab]['total']}"}, 400)
                    con.execute('UPDATE records SET override=?, reason=? WHERE cls=? AND lab=? AND sid=?', (ov, str(body.get('reason', ''))[:200], cls, lab, sid))
                con.execute('UPDATE records SET updated=? WHERE cls=? AND lab=? AND sid=?', (now(), cls, lab, sid))
            elif path == '/api/sync':           # remember what was sent to Canvas (+ the gradebook file)
                t = now()
                for lab, sc in body.get('scores', {}).items():
                    set_sync(con, cls, lab, t, sc)
                if body.get('gb'):
                    set_gradebook(con, cls, t, str(body.get('gbname', ''))[:80], body['gb'][:3000])
            elif path == '/api/publish':        # what students may see for one class + lab
                lab = str(body['lab'])
                con.execute('INSERT INTO publish (cls, lab) VALUES (?, ?) ON CONFLICT DO NOTHING', (cls, lab))
                for k in ('card', 'marks', 'submit'):
                    if k in body:
                        con.execute(f'UPDATE publish SET {k}=? WHERE cls=? AND lab=?', (1 if body[k] else 0, cls, lab))
            elif path == '/api/team':           # team-level manual marks (Lab 4 demo)
                if body.get('value') is None:   # tapping the selected mark again clears it
                    con.execute('DELETE FROM team_marks WHERE cls=? AND lab=? AND team=? AND field=?', (cls, str(body['lab']), str(body['team']), body['field']))
                else:
                    set_team_mark(con, cls, body['lab'], body['team'], body['field'], body['value'])
            else:
                return self.send_json({'error': 'not found'}, 404)
            con.commit()
            return self.send_json(load(con))

    # ---------------- student page ----------------
    def public_get(self, path):
        with connect() as con:
            st = self.student(con)
            if not allow(con, 'get:' + self.ip(), 240):
                return self.send_json(TOO_MANY, 429)
            if path == '/api/public/me':
                if not st:
                    return self.send_json({'error': 'signin'}, 401)
                return self.send_json(student_view(st['cls'], st, load(con)))
            m = re.fullmatch(r'/api/public/photo/(\d+)\.jpg', path)
            if m:                               # only the student's own photos
                r = con.execute('SELECT sid, cls, data FROM photos WHERE id=?', (int(m.group(1)),)).fetchone()
                if not st or not r or r['data'] is None or (r['cls'], r['sid']) != (st['cls'], st['sid']):
                    return self.send_json({'error': 'not found'}, 404)
                return self.send_bytes(bytes(r['data']), 'image/jpeg', 'no-store')
            m = re.fullmatch(r'/api/public/lab3/(\d+)\.png', path)
            if m:                               # only the student's own scope screen, when the Lab 3 card is open
                v = m.group(1)
                pub = con.execute("SELECT card FROM publish WHERE cls=? AND lab='3'", (st['cls'],)).fetchone() if st else None
                f = os.path.join(ASSETS, 'lab3', f'v{v}.png')
                if not st or str(variant(st['cls'], st)) != v or not pub or not pub['card'] or not os.path.exists(f):
                    return self.send_json({'error': 'not available'}, 404)
                return self.send_bytes(open(f, 'rb').read(), 'image/png')
        return self.send_json({'error': 'not found'}, 404)

    def public_post(self, con, path, body):
        if path == '/api/public/login':
            return self.student_login(con, body)
        if path == '/api/public/logout':
            con.execute('DELETE FROM stu_sessions WHERE token=?', (self.cookie('stu'),))
            return self.send_json({'ok': True}, cookie=self.set_cookie('stu', '', 0, 'Lax'))
        st = self.student(con)
        if not st:
            return self.send_json({'error': 'Please sign in again.', 'error_vi': 'Đăng nhập lại.'}, 401)
        if not allow(con, 'post:' + st['sid'], 40):
            return self.send_json(TOO_MANY, 429)
        if path == '/api/public/submit':
            return self.public_submit(con, st, body)
        if path in ('/api/public/photo', '/api/public/photo_delete'):
            return self.public_photo(con, st, body, delete=path.endswith('delete'))
        return self.send_json({'error': 'not found'}, 404)

    def student_login(self, con, body):
        """Student ID + 4-digit PIN. First time: the student chooses the PIN. Forgot it: the lecturer resets it."""
        if not allow(con, 'stu-login:' + self.ip(), 30, 600):
            return self.send_json(TOO_MANY, 429)
        sid = str(body.get('sid', '')).strip().lower()
        row = con.execute('SELECT cls, sid, name, team, member, pin FROM students WHERE lower(sid)=?', (sid,)).fetchone()
        if not row:
            return self.send_json({'error': 'Student ID not found. Check it (e.g. AUS15128).', 'error_vi': 'Không tìm thấy MSSV. Kiểm tra lại (vd AUS15128).'}, 404)
        fail_key = f"stu-fail:{row['cls']}:{row['sid']}"
        if count_hits(con, fail_key, 900) >= 5:
            return self.send_json({'error': 'Too many wrong PINs. Wait 15 minutes or ask the lecturer to reset your PIN.',
                                   'error_vi': 'Sai PIN quá nhiều lần. Đợi 15 phút hoặc nhờ giảng viên đặt lại PIN.'}, 429)
        pin = str(body.get('pin', '')).strip()
        if not row['pin']:                      # first sign-in: choose a PIN
            if not pin:
                return self.send_json({'need': 'new', 'name': row['name']})
            if not re.fullmatch(r'\d{4}', pin):
                return self.send_json({'error': 'The PIN must be 4 digits.', 'error_vi': 'PIN phải gồm 4 chữ số.'}, 400)
            con.execute('UPDATE students SET pin=? WHERE cls=? AND sid=?', (h_student(row['cls'], row['sid'], pin), row['cls'], row['sid']))
        elif not pin:
            return self.send_json({'need': 'pin'})
        elif h_student(row['cls'], row['sid'], pin) != row['pin']:
            add_hit(con, fail_key)
            return self.send_json({'error': 'Wrong PIN.', 'error_vi': 'Sai PIN.'}, 403)
        tok = new_student_session(con, row['cls'], row['sid'])
        return self.send_json({'ok': True}, cookie=self.set_cookie('stu', tok, 15552000, 'Lax'))

    def _open_lab(self, con, st, lab):
        """Checks before a student changes one lab. Returns True or sends the error."""
        if lab not in KEYS['labs']:
            self.send_json({'error': 'bad lab'}, 400); return False
        pub = con.execute('SELECT submit FROM publish WHERE cls=? AND lab=?', (st['cls'], lab)).fetchone()
        if not pub or not pub['submit']:
            self.send_json({'error': 'Submissions for this lab are closed.', 'error_vi': 'Lab này đang không nhận bài.'}, 403); return False
        rec = con.execute('SELECT tick FROM records WHERE cls=? AND lab=? AND sid=?', (st['cls'], lab, st['sid'])).fetchone()
        if rec and rec['tick']:
            self.send_json({'error': 'Already checked by the lecturer – locked.', 'error_vi': 'Giảng viên đã xác nhận – đã khóa.'}, 409); return False
        return True

    def public_submit(self, con, st, body):
        lab = str(body.get('lab', ''))
        if not self._open_lab(con, st, lab):
            return
        cls, sid = st['cls'], st['sid']
        n = len(KEYS['labs'][lab]['questions']); vals = []
        for x in (body.get('values') or [])[:n]:
            try:
                vals.append(None if x in (None, '') else float(str(x).replace(',', '.')))
            except ValueError:
                return self.send_json({'error': 'Numbers only, please.', 'error_vi': 'Chỉ nhập số.'}, 400)
        vals += [None] * (n - len(vals))
        if all(v is None for v in vals):
            return self.send_json({'error': 'Type at least one value.', 'error_vi': 'Nhập ít nhất một số.'}, 400)
        t = now()
        con.execute('INSERT INTO records (cls, lab, sid, tick, manual) VALUES (?,?,?,0,?) ON CONFLICT DO NOTHING', (cls, lab, sid, '{}'))
        con.execute("UPDATE records SET vals=?, src='web', updated=? WHERE cls=? AND lab=? AND sid=?", (json.dumps(vals), t, cls, lab, sid))
        con.execute('INSERT INTO submissions (cls, lab, sid, vals, time, ip, agent) VALUES (?,?,?,?,?,?,?)',
                    (cls, lab, sid, json.dumps(vals), t, self.ip(), (self.headers.get('User-Agent') or '')[:200]))
        return self.send_json({'ok': True, 'time': t})

    def public_photo(self, con, st, body, delete=False):
        lab = str(body.get('lab', ''))
        if not self._open_lab(con, st, lab):
            return
        cls, sid = st['cls'], st['sid']
        if delete:
            con.execute('DELETE FROM photos WHERE id=? AND cls=? AND lab=? AND sid=?', (as_int(body.get('id')) or 0, cls, lab, sid))
            return self.send_json({'ok': True})
        n = con.execute('SELECT COUNT(*) FROM photos WHERE cls=? AND lab=? AND sid=?', (cls, lab, sid)).fetchone()[0]
        if n >= 3:
            return self.send_json({'error': 'Maximum 3 photos. Delete one first.', 'error_vi': 'Tối đa 3 ảnh. Xóa bớt một ảnh trước.'}, 400)
        try:
            raw = base64.b64decode(str(body.get('data', '')).split(',')[-1], validate=True)
        except Exception:
            raw = b''
        if not raw.startswith(b'\xff\xd8') or len(raw) > 3_000_000:
            return self.send_json({'error': 'Photo not accepted (JPEG, max 3 MB).', 'error_vi': 'Ảnh không hợp lệ (JPEG, tối đa 3 MB).'}, 400)
        t = now()
        pid = con.execute("INSERT INTO photos (cls, lab, sid, file, time, ip, data) VALUES (?,?,?,'',?,?,?) RETURNING id",
                          (cls, lab, sid, t, self.ip(), raw)).fetchone()['id']
        return self.send_json({'ok': True, 'id': pid, 'time': t})


def student_view(cls, st, data):
    """Only this student's own information. Never the answer keys."""
    C = data['classes'][cls]
    st = next(s for s in C['students'] if s['sid'] == st['sid'])
    v = variant(cls, st)
    mates = [s['name'] for s in C['students'] if st.get('team') and s.get('team') == st['team'] and s['sid'] != st['sid']]
    labs = []
    for lab in ('1', '2', '3', '4'):
        pub = C.get('publish', {}).get(lab, {})
        L = KEYS['labs'][lab]
        item = {'lab': lab, 'title': L['title'], 'total': L['total'], 'card': None, 'figs': [], 'image': None, 'marks': None, 'submit': None}
        if pub.get('card') and v and str(v) in L['variants']:
            item['card'] = L['variants'][str(v)].get('card')
            item['figs'] = L['variants'][str(v)].get('figs', [])
            if lab == '3':
                item['image'] = f'/api/public/lab3/{v}.png'
        if pub.get('marks'):
            item['marks'] = score(cls, lab, st, data)
        if pub.get('submit') and v:
            rec = C['labs'].get(lab, {}).get(st['sid'], {})
            item['submit'] = {'questions': [[q[1], q[2]] for q in L['questions']],   # label + unit only
                              'values': rec.get('values') if rec.get('src') == 'web' else None,
                              'time': rec.get('updated') if rec.get('src') == 'web' else None,
                              'locked': bool(rec.get('tick')),
                              'photos': [{'id': p['id'], 'url': f"/api/public/photo/{p['id']}.jpg", 'time': p['time']}
                                         for p in C.get('photos', {}).get(lab, {}).get(st['sid'], [])]}
        labs.append(item)
    return {'name': st['name'], 'sid': st['sid'], 'className': C['name'], 'team': st.get('team'), 'member': st.get('member'),
            'variant': v, 'teammates': mates, 'labs': labs}


class Server(socketserver.ThreadingMixIn, http.server.HTTPServer):
    daemon_threads = True
    allow_reuse_address = True
