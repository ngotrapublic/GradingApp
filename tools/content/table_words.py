"""The HCM lab has plain tables and chairs, no lab benches: say "table" instead of "bench" in the slide text and the
speaker notes of the 4 decks (bench -> table, Bench -> Table, benches -> tables). Run first, so the scripts that find
their slides by title ("Find your table", "Before you leave – make your table look like this") find the old ones,
then run again at the end. Safe to run many times."""
import re
from pptx import Presentation

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
WORD = re.compile(r'(?<![\w-])([Bb])ench(es)?\b(?!_)')


def fix(t):
    return WORD.sub(lambda m: ('T' if m.group(1) == 'B' else 't') + 'able' + ('s' if m.group(2) else ''), t)


def frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:                                       # group
            yield from frames(sh.shapes)
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, 'has_table', False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame


def run():
    for n in (1, 2, 3, 4):
        path = f'{HCM}/Lab{n}/HCM Electronic Circuit Lab {n}.pptx'
        p = Presentation(path); k = 0
        for s in p.slides:
            tfs = list(frames(s.shapes)) + ([s.notes_slide.notes_text_frame] if s.has_notes_slide else [])
            for tf in tfs:
                for para in tf.paragraphs:
                    for r in para.runs:
                        new = fix(r.text)
                        if new != r.text:
                            r.text = new; k += 1
        p.save(path); print(f'Lab {n}: {k} runs changed')


if __name__ == '__main__':
    run()
