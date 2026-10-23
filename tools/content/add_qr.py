"""Put the student-page QR code on the slides where students open the student page (run on the final decks)."""
import segno
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

URL = 'https://grading-app-tau.vercel.app/student/'
HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
QR = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/qr_student.png'
segno.make(URL, error='m').save(QR, scale=20, border=2)


def find(slide, start):
    return next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith(start))


def add_qr(slide, left, top, size):
    for sh in list(slide.shapes):                       # running twice replaces the old code
        if sh.name in ('Student QR', 'Student QR caption'):
            sh._element.getparent().remove(sh._element)
    pic = slide.shapes.add_picture(QR, Inches(left), Inches(top), Inches(size), Inches(size)); pic.name = 'Student QR'
    tb = slide.shapes.add_textbox(Inches(left + size / 2 - 1.4), Inches(top + size + 0.02), Inches(2.8), Inches(0.5)); tb.name = 'Student QR caption'
    tf = tb.text_frame; tf.word_wrap = True
    for i, (txt, bold, sz) in enumerate([('Student page', True, 14), (URL.replace('https://', '').rstrip('/'), False, 10)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = 'Arial'; r.font.color.rgb = RGBColor(0x3F, 0x7A, 0x2A) if bold else RGBColor(0x55, 0x55, 0x55)


def shrink_tables(slide, width):
    for sh in slide.shapes:
        if sh.has_table:
            k = Inches(width) / sh.width
            for c in sh.table.columns:
                c.width = int(c.width * k)
            sh.width = Inches(width)


# Lab 1: sign-in slide + what to submit
p = Presentation(f'{HCM}/Lab1/HCM Electronic Circuit Lab 1.pptx')
s = p.slides[5]; find(s, 'Your team stays the same').width = Inches(9.6); add_qr(s, 10.75, 4.62, 1.85)
s = p.slides[85]; find(s, 'Before you leave').width = Inches(9.0); add_qr(s, 10.95, 5.2, 1.4)
p.save(f'{HCM}/Lab1/HCM Electronic Circuit Lab 1.pptx')
# Lab 2, 3: what to submit (free space on the right)
for lab, idx in ((2, 44), (3, 36)):
    p = Presentation(f'{HCM}/Lab{lab}/HCM Electronic Circuit Lab {lab}.pptx')
    shrink_tables(p.slides[idx], 9.9); add_qr(p.slides[idx], 11.0, 1.35, 1.95)
    p.save(f'{HCM}/Lab{lab}/HCM Electronic Circuit Lab {lab}.pptx')
# Lab 4: what is marked (narrow the tables first)
p = Presentation(f'{HCM}/Lab4/HCM Electronic Circuit Lab 4.pptx')
s = p.slides[84]; shrink_tables(s, 9.7); add_qr(s, 10.95, 1.35, 1.95)
p.save(f'{HCM}/Lab4/HCM Electronic Circuit Lab 4.pptx')
print('ok')
