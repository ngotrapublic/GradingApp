"""Lab 2 (lecturer, 29/10/2026): the oscilloscope and function generator lesson moves from Lab 1 to Lab 2. Lab 1 is not
changed (two classes already had it; the lecturer just skips that part). The Lab 1 slides from "Overview of AC lab
instruments" to the last "Function Generator to Oscilloscope" slide are copied into Lab 2, after "Devices and components
for this session" and before the RC part. Their exercises (Lab 1 numbers 8, 9, 10, 10) become Exercise A, B, C, D, so the
Lab 2 exercises keep their numbers 1-7. The notes of a copied slide are the text of the slide (the Lab 1 slides have none).
The copies carry the slide name below (cSld name): running again replaces them."""
import copy, io, re
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
SRC = f'{HCM}/Lab1/HCM Electronic Circuit Lab 1.pptx'
DECK = f'{HCM}/Lab2/HCM Electronic Circuit Lab 2.pptx'
NAME = 'From Lab 1: AC instruments'
FIRST, LAST = 'Overview of AC lab instruments', 'Step 5. Rotate the Multipurpose knob'
AFTER = 'Devices and components for this session'
LETTERS = iter('ABCD')
OVERVIEW_LINE = ('First: the oscilloscope and the function generator (Exercises A–D). Before you measure: check 1X / 10X '
                 'on the oscilloscope, and set the function generator to High-Z.')
OVERVIEW_SAY = ('First: the oscilloscope and the function generator, Exercises A to D. Before you measure: check 1X or 10X '
                'on the oscilloscope, and set the function generator to High-Z.')
R_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def text(s):
    return '\n'.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)


def copy_slide(dst, src_slide, layout):
    s = dst.slides.add_slide(layout)
    for el in list(s.shapes._spTree)[2:]:                 # the layout placeholders of the new slide
        s.shapes._spTree.remove(el)
    rid = {}
    for r_id, rel in src_slide.part.rels.items():
        if rel.reltype == RT.IMAGE:
            _, rid[r_id] = s.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        elif rel.is_external:
            rid[r_id] = s.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
    sld = copy.deepcopy(src_slide._element)
    for el in sld.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith(R_EMBED) and v in rid:
                el.set(k, rid[v])
    tree = s.shapes._spTree                                # keep the same spTree object (slide.shapes is cached)
    for el in list(tree):
        tree.remove(el)
    for el in sld.cSld.spTree:
        tree.append(el)
    bg = sld.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if bg is not None:
        s._element.cSld.insert(0, bg)
    for tag in ('transition', 'timing'):
        el = sld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}' + tag)
        if el is not None:
            s._element.append(el)
    s._element.cSld.set('name', NAME)
    return s


def renumber(s):
    """Exercise 8, 9, 10, 10 of Lab 1 -> Exercise A, B, C, D (the number can be split over runs)."""
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            m = re.search(r'Exercise (\d+)', para.text)
            if not m:
                continue
            new = para.text.replace(m.group(0), f'Exercise {next(LETTERS)}', 1)
            runs = para.runs
            runs[0].text = new
            for r in runs[1:]:
                r._r.getparent().remove(r._r)


def run():
    src, p = Presentation(SRC), Presentation(DECK)
    lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if sl._element.cSld.get('name') == NAME:
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    slides = list(src.slides)
    a = next(i for i, s in enumerate(slides) if FIRST in text(s))
    b = next(i for i, s in enumerate(slides) if LAST in text(s))
    layout = next(l for l in p.slide_layouts if l.name == 'Blank')
    new = []
    for s0 in slides[a:b + 1]:
        s = copy_slide(p, s0, layout)
        renumber(s)
        lines = [ln.strip() for sh in s.shapes if sh.has_text_frame for ln in sh.text_frame.text.split('\n') if ln.strip()]
        s.notes_slide.notes_text_frame.text = '\n'.join(lines)
        new.append(s)
    at = next(i for i, s in enumerate(p.slides) if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == AFTER
                                                         for sh in s.shapes)) + 1
    sids = [lst[list(p.slides).index(x)] for x in new]
    for sid in sids: lst.remove(sid)
    for k, sid in enumerate(sids): lst.insert(at + k, sid)
    # the overview: the reminder line now says that the lab starts with the instruments
    ov = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == 'Lab 2 overview' for sh in s.shapes))
    box = next(sh for sh in ov.shapes if sh.has_text_frame and sh.text_frame.text.startswith(('Reminder:', 'First:')))
    runs = box.text_frame.paragraphs[0].runs
    runs[0].text = OVERVIEW_LINE
    for r in runs[1:]: r._r.getparent().remove(r._r)
    nt = ov.notes_slide.notes_text_frame
    nt.text = '\n'.join([ln for ln in nt.text.split('\n') if not ln.startswith(('Reminder:', 'First:'))] + [OVERVIEW_SAY])
    p.save(DECK)
    print(f'Lab 2: {len(new)} slides from Lab 1 ({a + 1}-{b + 1}) at {at + 1}-{at + len(new)} | {len(p.slides)} slides')


if __name__ == '__main__':
    run()
