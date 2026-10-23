import runpy, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
g=runpy.run_path('bom2.py'); need,spare,reuse=g['need'],g['spare'],g['reuse']
S=runpy.run_path('shop.py')['S']; U='https://www.thegioiic.com/'
key=lambda k:{'R':(f'Điện trở {k[1]/1000:g}kΩ' if k[1]>=1000 else f'Điện trở {k[1]:g}Ω')}.get(k[0])
M={('Cf','2.2nF'):'Tụ film (polyester) 2.2nF',('Cf','10nF'):'Tụ film (polyester) 10nF',('Cf','22nF'):'Tụ film (polyester) 22nF',('Cf','33nF'):'Tụ film (polyester) 33nF',('Cf','47nF'):'Tụ film (polyester) 47nF',('Cf','100nF'):'Tụ film (polyester) 100nF',('Cf','220nF'):'Tụ film (polyester) 220nF',('Cc','100nF'):'Tụ gốm 100nF',('Ce','1µF 50V'):'Tụ hóa 1µF 50V',('Ce','10µF 25V'):'Tụ hóa 10µF 25V',('Ce','470µF 25V'):'Tụ hóa 470µF 25V',('MIC','Electret'):'Micro Electret 6x5mm chân dài',('LED','LED 5mm'):'LED 5mm đỏ'}
S[('Cc','1nF')]=dict(id=992,name='Tụ Gốm Tròn 1nF 50V',moq=20,mult=10,inv=2250,tiers=[(1,144)],url=U+'tu-gom-tron-1nf-50v')
S[('L','1mH 0410-102')]=dict(id=4226,name='Cuộn Cảm Xuyên Lỗ 0410-102 1000uH (1mH) 85mA 1/2W, ±10%',moq=10,mult=10,inv=1151,tiers=[(1,456),(100,433)],url=U+'cuon-cam-xuyen-lo-0410-102-1000uh-1mh-85ma-1-2w')
S[('IC','LM358L (DIP-8)')]=dict(id=36272,name='LM358L IC OPAMP 2 Circuit 8-DIP (UTC)',moq=1,mult=1,inv=482,tiers=[(1,2850)],url=U+'lm358l-ic-opamp-general-purpose-amplifier-2-circuit-8-dip')
def shop(k):
    if k in S: return S[k]
    return S[key(k) if k[0]=='R' else M[k]]
disp=lambda k: key(k) if k[0]=='R' else {'Cf':'Tụ film ','Cc':'Tụ gốm ','Ce':'Tụ hóa ','L':'Cuộn cảm ','IC':'IC ','MIC':'Micro ','LED':''}[k[0]]+k[1]
order={'R':0,'Cc':1,'Cf':2,'Ce':3,'L':4,'IC':5,'MIC':6,'LED':7}
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); thin=Side(style='thin',color='BBBBBB'); B=Border(top=thin,bottom=thin,left=thin,right=thin)
YEL=PatternFill('solid',fgColor='FFF2CC'); GRN=PatternFill('solid',fgColor='E2EFDA'); RED=PatternFill('solid',fgColor='F8CBAD')
wb=Workbook(); ws=wb.active; ws.title='Đơn hàng thegioiic'
ws.append(['Danh sách mua – Electric Circuit Lab HCM – 3 lớp × 18 SV (6 cụm × 3) – giá thegioiic.com tra ngày 20/10/2026'])
ws['A1'].font=Font(bold=True,size=14)
ws.append(['Quy tắc: chỉ thu hồi linh kiện nhìn là phân biệt được ngay (nền xanh: LM358, micro, tụ 470µF lon to, LED) – phát đầu giờ, thu lại cuối giờ, dùng chung 3 lớp. Còn lại là tiêu hao (nền trắng), tính riêng cho từng lớp. Số đặt đã làm tròn theo Min/Bội số của shop. Giá chưa gồm VAT/ship.'])
ws.append([])
hdr=['STT','Linh kiện','Loại','N1','N2','N3','N4','Thi','Tổng cần','% dư','Muốn mua','Min','Bội số','SỐ LƯỢNG ĐẶT','Đơn giá (đ)','Thành tiền (đ)','Tồn kho shop','Mã TGIC','Tên trên thegioiic','Link','Ghi chú']
ws.append(hdr)
for c in ws[4]: c.font=H; c.fill=HF; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
notes={('R',3900):'Shop không có 3.6k chân cắm → cặp 3k6/4k7 đổi thành 3k9/4k7',('MIC','Electret'):'Chân ngắn, cần hàn thêm chân (như ảnh thầy/cô đã làm). Shop còn ít',('L','1mH 0410-102'):'Ghép với 22nF // 2.2nF → f0 ≈ 32.3 kHz'}
r0=5; tot=0
keys=sorted(need,key=lambda k:(order[k[0]],k[1] if k[0]=='R' else str(k[1])))
for i,k in enumerate(keys):
    c=need[k]; t=sum(c.values()); want=math.ceil(t*(1+spare[k[0]])); s=shop(k)
    buy=max(s['moq'],math.ceil(want/s['mult'])*s['mult']); price=[p for q,p in s['tiers'] if buy>=q][-1]; tot+=buy*price
    rr=r0+i; ru=k[0] in reuse or k==('Ce','470µF 25V')
    ws.append([i+1,disp(k),'Thu hồi' if ru else 'Tiêu hao',c['N1'],c['N2'],c['N3'],c['N4'],c['Thi'],t,f'{int(spare[k[0]]*100)}%',want,s['moq'],s['mult'],buy,price,f'=N{rr}*O{rr}',s['inv'],s['id'],s['name'],s['url'],notes.get(k,'')])
    ws.cell(rr,20).hyperlink=s['url']; ws.cell(rr,20).font=Font(color='0563C1',underline='single')
    ws.cell(rr,14).font=Font(bold=True); ws.cell(rr,14).fill=YEL
    if ru:
        for cc in range(1,4): ws.cell(rr,cc).fill=GRN
    if k in notes: ws.cell(rr,21).fill=RED
last=r0+len(keys)-1
ws.append(['','TỔNG TIỀN']+['']*13+[f'=SUM(P{r0}:P{last})']); t=last+1
ws.cell(t,2).font=Font(bold=True); ws.cell(t,16).font=Font(bold=True,size=12); ws.cell(t,16).fill=YEL
for row in ws.iter_rows(min_row=4,max_row=t):
    for c in row: c.border=B
for rr in range(r0,t+1):
    for cc in (15,16,17): ws.cell(rr,cc).number_format='#,##0'
for col,w in zip('ABCDEFGHIJKLMNOPQRSTU',[5,22,9,5,5,5,5,5,8,6,8,6,7,11,10,13,9,8,44,40,55]): ws.column_dimensions[col].width=w
ws.freeze_panes='C5'
o=t+2
for a in ['Không cần mua: loa 8Ω, breadboard, nguồn ± (có trên Global PB-503A); điện trở 1MΩ (không cần cầu chia tạo ±6V).',
          'Ngày 3: mỗi SV 1 LM358 (thu hồi, dùng lại ngày 4 và cho lớp sau). Muốn tiết kiệm thêm: 2 SV dùng chung 1 IC (mỗi bạn 1 op-amp).']:
    ws.cell(o,2,a); o+=1
w2=wb.create_sheet('Cách tính')
lines=[('Ngày 1 (mỗi SV)','6 điện trở theo mã đề + 1×2.2k + 1×4.7k (cầu chia áp/dòng). LED: thu hồi'),
('Ngày 2 (mỗi SV)','RC khởi động: 1R (1k/2.2k/4.7k) + 1C (10/22/47nF). RLC: 1mH (tiêu hao – nhìn giống điện trở) + 22nF // 2.2nF + R ghép từ 2×47Ω, 2×100Ω, 1×220Ω'),
('Ngày 3 (mỗi SV)','1 LM358 (thu hồi) + cặp R theo mã đề (220/1.2k, 1.2k/47k, 3.9k/4.7k) + 2×100nF gốm. Mỗi cụm 2×470µF (thu hồi)'),
('Ngày 4 (mỗi cụm)','Mạch mic→loa bản v2: 5×10k, 2×2.2k, 1×4.7k, 1×47k, 1nF, 4×100nF gốm (tiêu hao); 1µF, 10µF (tiêu hao); LM358, electret, 4×470µF (thu hồi)'),
('Thi (mỗi đội, 18 đội)','Điện trở và tụ film (kẹo xanh nhìn giống nhau): tiêu hao'),
('% dư','Điện trở 15% (shop bán bội 50 nên tự dư), tụ 20%, tụ hóa 25%, LM358 50%, micro 30%, cuộn cảm/LED 10%')]
w2.append(['Phần','Linh kiện dùng'])
for c in w2[1]: c.font=H; c.fill=HF
for l in lines: w2.append(list(l))
w2.column_dimensions['A'].width=20; w2.column_dimensions['B'].width=130
wb.save('/Users/mac/Tra/SE9_ElecCircuitLab/HCM/docs/HCM_Linh_kien_can_mua.xlsx')
print('TOTAL',tot)
