"""Lab 2 (lecturer, 29/10/2026: students are slower than planned): Exercises 6-7 (band-pass table, bandwidth) become a
lecturer demo. Task 2 is the same measurement with the team's C and R, so the students do it only once, in Task 2.
Renames the headings and the references ("as in Exercise 6", "the method of Exercise 7") on the slides and in the notes.
Run after instrument_slides.py and lab2_task2.py, before subscripts.py. Safe to run again."""
from pptx import Presentation

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/HCM Electronic Circuit Lab 2.pptx'
REPL = [('Exercise 6: Measure a band-pass RLC filter and draw its graph (15 minutes)',
         'Demo by the lecturer: a band-pass RLC filter'),
        ('Demo by the lecturer: a band-pass RLC filter. Watch, and write the values in your notebook.',
         'Demo by the lecturer: a band-pass RLC filter'),
        ('Exercise 7 (15 minutes). Keep your circuit of Exercise 6.',
         'Demo by the lecturer, with the same circuit. You will do this yourself in Task 2.'),
        ('Exercise 7 – measure the bandwidth', 'Demo – measure the bandwidth'),
        ('Exercise 7 – answer', 'Demo – the result'),
        ('Is your BW close to the theory', 'Is the BW close to the theory'),
        ('the same circuit as in Exercise 6', 'the same circuit as in the demo'),
        ('the method of Exercise 7', 'the method of the demo'),
        ('Exercise 7', 'the demo')]                           # anything left (notes)


FIRST_LINE = 'Demo by the lecturer, with the same circuit. You will do this yourself in Task 2.'


def fix(text, slide=False):
    for a, b in (REPL[:-1] if slide else REPL):
        text = text.replace(a, b)
    return text


def run():
    p = Presentation(DECK); n = 0
    for s in p.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                if para.text.startswith(('Exercise 7 (15 minutes).', 'the demo (15 minutes).')):   # bold run + plain run
                    runs = para.runs; runs[0].text = FIRST_LINE
                    for r in runs[1:]: r._r.getparent().remove(r._r)
                    n += 1; continue
                for r in para.runs:
                    if fix(r.text, slide=True) != r.text:
                        r.text = fix(r.text, slide=True); n += 1
                left = [a for a, _ in REPL[:-1] if a in para.text]
                assert not left, (left, para.text)           # an old phrase split over runs
        if s.has_notes_slide:
            tf = s.notes_slide.notes_text_frame
            if fix(tf.text) != tf.text:
                tf.text = fix(tf.text); n += 1
    p.save(DECK); print(f'Lab 2: {n} texts changed (Exercises 6-7 -> lecturer demo)')


if __name__ == '__main__':
    run()
