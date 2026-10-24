"""Lab 3: the photo of "Connect the rest of the circuit" was the Hanoi build with its own ±6 V supply (lecturer,
24/10/2026: not right for HCM, where the ±6 V comes from the PB-503A). It is replaced by the wiring diagram
fig/p3_wiring.png (fig_practice.py), fitted into the same box. Safe to run again."""
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab3/HCM Electronic Circuit Lab 3.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/p3_wiring.png'
NAME = 'Lab3 wiring'


def run():
    p = Presentation(DECK)
    s = next(s for s in p.slides if any(sh.has_text_frame and 'Connect the rest of the circuit' in sh.text_frame.text for sh in s.shapes))
    pics = [sh for sh in s.shapes if sh.shape_type == 13 and (sh.name == NAME or sh.width > Emu(914400 * 4))]
    old = max(pics, key=lambda sh: sh.width * sh.height)
    x, y, w, h = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    iw, ih = Image.open(FIG).size; k = min(w / iw, h / ih); pw, ph = int(iw * k), int(ih * k)
    pic = s.shapes.add_picture(FIG, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph); pic.name = NAME
    p.save(DECK); print('Lab 3: wiring diagram on slide', list(p.slides).index(s) + 1)


if __name__ == '__main__':
    run()
