import json, io, contextlib, runpy, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
with contextlib.redirect_stdout(io.StringIO()):
    runpy.run_path('bom.py'); g=runpy.run_path('design.py')
S=runpy.run_path('shop.py')['S']
rows=json.load(open('bom.json'))
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); thin=Side(style='thin',color='BBBBBB'); B=Border(top=thin,bottom=thin,left=thin,right=thin)
YEL=PatternFill('solid',fgColor='FFF2CC'); RED=PatternFill('solid',fgColor='F8CBAD')
wb=Workbook(); ws=wb.active; ws.title='Đơn hàng thegioiic'
ws.append(['Danh sách mua – Electric Circuit Lab HCM (6 cụm × 3 SV) – giá tra trên thegioiic.com ngày 20/10/2026'])
ws['A1'].font=Font(bold=True,size=14)
ws.append(['"Cần" = đúng theo bài, mỗi buổi dùng linh kiện mới. "Muốn mua" = Cần + % dư. "Số lượng đặt" đã làm tròn theo số lượng tối thiểu (Min) và Bội số của shop. Đơn giá theo bậc số lượng của shop. Giá chưa gồm VAT/phí ship.'])
ws.append([])
hdr=['STT','Linh kiện','N1','N2','N3','N4','Thi','Tổng cần','% dư','Muốn mua','Min','Bội số','SỐ LƯỢNG ĐẶT','Đơn giá (đ)','Thành tiền (đ)','Tồn kho shop','Mã TGIC','Tên trên thegioiic','Link','Ghi chú']
ws.append(hdr)
for c in ws[4]: c.font=H; c.fill=HF; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
notes={'Điện trở 3.9kΩ':'Shop KHÔNG có 3.6k chân cắm → đổi cặp 3k6/4k7 thành 3k9/4k7 (độ lợi −1.21)',
'Cuộn cảm 10mH':'Cuộn lọc UU10.5 4 chân (2 cuộn 10mH) – chỉ dùng 1 cuộn (2 chân cùng phía). Shop không có cuộn 10mH dạng 2 chân',
'Micro Electret 6x5mm chân dài':'Chân ngắn → có thể cần hàn thêm dây/chân để cắm breadboard. Shop chỉ còn 14 cái',
'Tụ gốm NP0/C0G 1nF':'Dùng tụ film kẹo xanh (ổn định hơn gốm thường)',
'Tụ film (polyester) 100nF':'Tồn kho thấp (345) – nên đặt sớm'}
r0=5
for i,r in enumerate(rows):
    name,n1,n2,n3,n4,thi,tot,sp,want=r
    s=S[name]; buy=max(s['moq'],math.ceil(want/s['mult'])*s['mult'])
    price=[p for q,p in s['tiers'] if buy>=q][-1]
    rr=r0+i
    disp=name.replace('Micro Electret 6x5mm chân dài','Micro electret').replace('Tụ gốm NP0/C0G 1nF','Tụ 1nF').replace('Tụ gốm NP0/C0G 270pF','Tụ gốm 270pF')
    ws.append([i+1,disp,n1,n2,n3,n4,thi,tot,sp,want,s['moq'],s['mult'],buy,price,f'=M{rr}*N{rr}',s['inv'],s['id'],s['name'],s['url'],notes.get(name,'')])
    ws.cell(rr,19).hyperlink=s['url']; ws.cell(rr,19).font=Font(color='0563C1',underline='single')
    ws.cell(rr,13).font=Font(bold=True); ws.cell(rr,13).fill=YEL
    if notes.get(name): ws.cell(rr,20).fill=RED
last=r0+len(rows)-1
ws.append(['','TỔNG TIỀN LINH KIỆN']+['']*12+[f'=SUM(O{r0}:O{last})'])
t=last+1
ws.cell(t,2).font=Font(bold=True); ws.cell(t,15).font=Font(bold=True,size=12); ws.cell(t,15).fill=YEL
for row in ws.iter_rows(min_row=4,max_row=t):
    for c in row: c.border=B
for rr in range(r0,t+1):
    for cc in (14,15,16): ws.cell(rr,cc).number_format='#,##0'
for col,w in zip('ABCDEFGHIJKLMNOPQRST',[5,26,5,5,5,5,5,8,6,8,6,7,11,10,13,9,8,44,40,60]): ws.column_dimensions[col].width=w
ws.freeze_panes='C5'
o=t+2
ws.cell(o,2,'Không cần mua').font=Font(bold=True); o+=1
for a in ['Loa 8Ω, nguồn ±, breadboard: đã có trên máy Global PB-503 (6 máy)','Điện trở 1MΩ: bỏ vì PB-503 có sẵn nguồn ± (không cần cầu chia tạo ±6V)']:
    ws.cell(o,2,a); o+=1

w2=wb.create_sheet('Cách tính theo ngày')
lines=[('Ngày 1 (mỗi SV)','6 điện trở theo mã đề (3 điện trở "lạ", mỗi con ghép từ 2 điện trở chuẩn) + 1×2.2k + 1×4.7k cho cầu chia áp/dòng + 1 LED'),
('Ngày 2 (mỗi SV)','RC khởi động: 1R (1k…6.8k) + 1C (10/22/47nF). RLC: 1 cuộn 10mH + 2.2nF // 270pF (≈2.47nF → f0 ≈ 32 kHz) + R theo Q (Q=1: 1k+1k, 1.5: 1k+330, 2: 1k, 3: 680, 4: 470, 5: 390) + 47Ω, 100Ω để bù'),
('Ngày 3 (mỗi SV)','1 LM358 + cặp R theo mã đề (220/1.2k, 1.2k/47k, 3.9k/4.7k – mỗi cặp 6 SV) + 2×100nF gốm decoupling; mỗi cụm 2×470µF trên đường nguồn ± của PB-503'),
('Ngày 4 (mỗi cụm)','Mạch mic→loa bản v2 của HN: 1 LM358, 1 electret, 5×10k, 2×2.2k, 1×4.7k, 1×47k (bản v1), 1µF, 10µF, 1nF, 4×100nF gốm, 4×470µF (2 nguồn + 2 thử loa), 2×100nF film (thử loa)'),
('Thi (mỗi đội)','Khay chọn tự do cho Section A (440Ω, 5k, 6.9k) và Section B (RC 100 Hz / 5 kHz): R 100Ω…15k, C 10n…1µF'),
('% dư','Điện trở, LED, LM358, micro: 50%. Tụ: 30%. Cuộn cảm: 20%. Sau đó làm tròn theo Min/Bội số của shop')]
w2.append(['Phần','Linh kiện dùng'])
for c in w2[1]: c.font=H; c.fill=HF
for l in lines: w2.append(list(l))
w2.column_dimensions['A'].width=18; w2.column_dimensions['B'].width=130
for r in w2.iter_rows(min_row=2):
    for c in r: c.alignment=Alignment(wrap_text=True,vertical='top')

w3=wb.create_sheet('Mã đề Ngày 1 (nháp)')
w3.append(['Mạch: (RA // RB) nối tiếp RC, nguồn 5V. X = nút giữa khối song song và RC.'])
w3.append(['Mã đề','RA (Ω)','Ghép RA','RB (Ω)','Ghép RB','RC (Ω)','Ghép RC','Rtđ (Ω)','I tổng (mA)','I qua RA (mA)','VX (V)'])
for c in w3[2]: c.font=H; c.fill=HF
f=lambda r: f'{r/1000:g}k' if r>=1000 else f'{r:g}'
for i,d in enumerate(g['D1'],1):
    row=[f'N1-{i:02d}']
    for k in 'ABC':
        v,tt,a,b=d[k]; row+=[round(v),f'{f(a)} {"nối tiếp" if tt=="+" else "//"} {f(b)}']
    row+=[round(d['Req']),round(d['I']*1e3,3),round(d['IA']*1e3,3),round(d['VX'],3)]
    w3.append(row)
for col,w in zip('ABCDEFGHIJK',[9,9,16,9,16,9,16,10,11,13,9]): w3.column_dimensions[col].width=w
wb.save('/Users/mac/Tra/SE9_ElecCircuitLab/HCM/docs/HCM_Linh_kien_can_mua.xlsx')
# print summary
tot=0
for i,r in enumerate(rows):
    s=S[r[0]]; buy=max(s['moq'],math.ceil(r[8]/s['mult'])*s['mult']); p=[p for q,p in s['tiers'] if buy>=q][-1]; tot+=buy*p
    print(r[0],buy,p,buy*p)
print('TOTAL',tot)
