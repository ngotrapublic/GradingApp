"""Lab 1, title slide: a notice that students do not unpack yet (they sit anywhere during the introduction and move
to their bench after "Find your bench"). Phones stay out: they sign in on the student page at the team slide.
Running again replaces the notice."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab1/HCM Electronic Circuit Lab 1.pptx'
NAME = 'Welcome Notice'
DARK, ORANGE, FILL = RGBColor(0x22, 0x22, 0x22), RGBColor(0xC2, 0x41, 0x0C), RGBColor(0xFF, 0xF7, 0xF2)
LINES = ['**Please do not unpack yet.** Keep laptops, notebooks and bags off the bench.',
         'Keep your **phone** – you need it in a few minutes.']
SAY = ('Please do not unpack yet. Keep your laptops, notebooks and bags off the bench. You will move to your bench '
       'after the introduction. Keep your phone: you need it to sign in on the student page.')


def run():
    p = Presentation(DECK); s = p.slides[0]
    for sh in list(s.shapes):
        if sh.name == NAME:
            sh._element.getparent().remove(sh._element)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.4), Inches(5.3), Inches(9.2), Inches(1.0))
    box.name = NAME; box.adjustments[0] = 0.15; box.shadow.inherit = False
    box.fill.solid(); box.fill.fore_color.rgb = FILL; box.line.color.rgb = ORANGE; box.line.width = Pt(2)
    style = box._element.find(qn('p:style'))
    if style is not None:
        box._element.remove(style)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.25)
    for i, text in enumerate(LINES):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.alignment = PP_ALIGN.LEFT
        for k, part in enumerate(text.split('**')):
            if part:
                r = para.add_run(); r.text = part; r.font.size = Pt(20); r.font.name = 'Arial'
                r.font.bold = k % 2 == 1; r.font.color.rgb = ORANGE if k % 2 == 1 else DARK
    nt = s.notes_slide.notes_text_frame
    if 'do not unpack' not in nt.text:
        nt.text = nt.text.rstrip() + '\n' + SAY
    p.save(DECK); print('Lab 1: notice on slide 1')


if __name__ == '__main__':
    run()
