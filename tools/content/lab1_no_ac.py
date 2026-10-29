"""Lab 1 (lecturer, 29/10/2026): the oscilloscope and function generator slides ("Overview of AC lab instruments" ...
"Step 5. Rotate the Multipurpose knob") are taught in Lab 2 now (lab2_instruments.py copied them), so they are deleted
from Lab 1. The course summary (labs 2 and 3 changed: lab3_rlc.py) and the devices slide no longer list them for Lab 1.
Run lab2_instruments.py before this; after this it keeps the Lab 2 copies. Safe to run again."""
from pptx import Presentation
from lab2_instruments import SRC, FIRST, LAST, text
from lab4_review import set_para, title

ROWS = [['**Lab 1** – Lab instruments & DC circuits', '3',
         'Lab rules, multimeter, DC power supply, breadboard; voltage & current dividers; equivalent resistance exercise'],
        ['**Lab 2** – Instruments & RC filters', '3',
         'Oscilloscope and function generator; 1st-order RC low-pass / high-pass filters; cut-off frequency; capacitors'],
        ['**Lab 3** – RLC filters & operational amplifiers', '3',
         'Inductors; series RLC band-pass filter (resonance, bandwidth); op-amp basics, LM358, ±6 V supply; inverting amplifier: gain and 180° phase shift']]
SUMMARY_SAY = ('This is our plan. Lab 1: instruments and DC circuits. Lab 2: the oscilloscope, the function generator and '
               'RC filters. Lab 3: RLC filters and op-amps. Lab 4: microphone to speaker, and exam review.\n'
               'After Lab 4, we have the final exam. It is also 3 hours.')
SCOPE_LINE = '1 oscilloscope, 1 function generator, 1X BNC cable, 10X probe'


def run():
    p = Presentation(SRC); lst = p.slides._sldIdLst; slides = list(p.slides)
    a = next((i for i, s in enumerate(slides) if FIRST in text(s)), None)
    gone = 0
    if a is not None:
        b = next(i for i, s in enumerate(slides) if LAST in text(s))
        for i in range(b, a - 1, -1):
            sid = lst[i]; p.part.drop_rel(sid.rId); lst.remove(sid); gone += 1
        p.part.rename_slide_parts([x.rId for x in lst])
    summary = next(s for s in p.slides if title(s) == 'Course Summary')
    tbl = next(sh.table for sh in summary.shapes if sh.has_table)
    for row, new in zip(list(tbl.rows)[1:4], ROWS):
        for cell, v in zip(row.cells, new):
            set_para(cell.text_frame.paragraphs[0], v)          # also fixes the bold of 'Lab N'
    summary.notes_slide.notes_text_frame.text = SUMMARY_SAY
    dev = next(s for s in p.slides if title(s) == 'Devices and components for this session')
    for sh in dev.shapes:
        if sh.has_text_frame:
            for para in list(sh.text_frame.paragraphs):
                if para.text == SCOPE_LINE: para._p.getparent().remove(para._p)
    nt = dev.notes_slide.notes_text_frame
    nt.text = nt.text.replace('You also have a multimeter, an oscilloscope and a function generator.',
                              'You also have a DC power supply and a multimeter.')
    p.save(SRC)
    print(f'Lab 1: {gone} oscilloscope / function generator slides deleted | {len(p.slides)} slides')


if __name__ == '__main__':
    run()
