"""Labs 2-4: write circuit variables with a real subscript on the slides (lecturer, 24/10/2026): fc -> f_c, f0 -> f_0,
Vin -> V_in, Vout -> V_out (also Vout,pp), VR/VC/VL, Vcc/Vee, XC/XL, RL, Rtotal, Z0, fL/fH, fmax. A run that contains one of these
words is split, and the subscript part gets baseline -25000 (same as rlc_slides.py). Text in pictures is not changed.
Speaker notes are spoken text: not changed. Run last (after the other slide scripts). Safe to run many times."""
import copy, re
from pptx import Presentation

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
WORD = re.compile(r'\b(?:(f)(c|0|L|H|max)|(V)(in|out)(,pp)?|(V)(cc|ee|CC|EE)|(V)(R|C|L)|(X)(C|L)|(R)(L|total|eq)(?! circuit)|(Z)(0))\b(?! circuit)')


def pieces(text):
    """'At fc, Vout = …' -> [('At ', False), ('f', False), ('c', True), (', ', False), ('V', False), ('out', True), …]"""
    out, pos = [], 0
    for m in WORD.finditer(text):
        g = [x for x in m.groups() if x is not None]
        out.append((text[pos:m.start()], False)); out.append((g[0], False)); out.append((''.join(g[1:]), True))
        pos = m.end()
    out.append((text[pos:], False))
    return [(t, sub) for t, sub in out if t]


def fix_run(r):
    rPr = r._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
    if rPr is not None and rPr.get('baseline') not in (None, '0'):
        return 0                                           # already a subscript / superscript
    parts = pieces(r.text)
    if not any(sub for _, sub in parts):
        return 0
    el = r._r; parent = el.getparent(); at = parent.index(el)
    for k, (t, sub) in enumerate(parts):
        new = copy.deepcopy(el)
        new.find('{http://schemas.openxmlformats.org/drawingml/2006/main}t').text = t
        nrPr = new.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        if nrPr is None:
            from pptx.oxml.ns import qn
            nrPr = new.makeelement(qn('a:rPr'), {}); new.insert(0, nrPr)
        if sub:
            nrPr.set('baseline', '-25000')
        parent.insert(at + k, new)
    parent.remove(el)
    return sum(1 for _, sub in parts if sub)


def frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from frames(sh.shapes)
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, 'has_table', False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame


def fix_deck(p):
    k = 0
    for s in p.slides:
        for tf in frames(s.shapes):
            for para in tf.paragraphs:
                for r in list(para.runs):
                    k += fix_run(r)
    return k


def run():
    for n in (2, 3, 4):
        path = f'{HCM}/Lab{n}/HCM Electronic Circuit Lab {n}.pptx'
        p = Presentation(path); k = fix_deck(p); p.save(path)
        print(f'Lab {n}: {k} subscripts')


if __name__ == '__main__':
    run()
