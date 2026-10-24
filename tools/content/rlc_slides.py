"""Lab 2: series RLC. Low-pass and high-pass are explained in detail in the RC part, so here they are one line each on
"Series RLC – one circuit, three filters" (lecturer, 24/10/2026); band-pass is the main topic: three slides (why only a
band passes, bandwidth, measuring BW), inserted after that slide.
Run after align_web.py; running again replaces the slides.   Figures: fig_rlc_theory.py."""
import copy, io, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/HCM Electronic Circuit Lab 2.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig'
AFTER = 'Series RLC – one circuit, three filters'
GREEN, DARK, BOXFILL, BOXLINE = RGBColor(0x3F, 0x7A, 0x2A), RGBColor(0x22, 0x22, 0x22), RGBColor(0xF3, 0xF7, 0xF0), RGBColor(0xB7, 0xCF, 0xA8)


def add_runs(par, text, size):
    """'**bold**' and '_{sub}' markup -> runs (subscripts as real subscripts)."""
    for chunk in re.split(r'(\*\*.*?\*\*)', text):                 # bold parts may contain subscripts
        bold = chunk.startswith('**'); chunk = chunk[2:-2] if bold else chunk
        for part in re.split(r'(_\{.*?\})', chunk):
            if not part:
                continue
            r = par.add_run(); r.font.size = Pt(size); r.font.name = 'Arial'; r.font.color.rgb = DARK; r.font.bold = bold
            if part.startswith('_{'):
                r.text = part[2:-1]; r._r.get_or_add_rPr().set('baseline', '-25000')
            else:
                r.text = part


def text_box(slide, x, y, w, h, paras, size=16, box=False, title=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    if box:
        tb.fill.solid(); tb.fill.fore_color.rgb = BOXFILL; tb.line.color.rgb = BOXLINE
        tf.margin_left = tf.margin_right = Inches(0.15); tf.margin_top = Inches(0.1)
    first = True
    if title:
        p = tf.paragraphs[0]; first = False; add_runs(p, f'**{title}**', size + 1); p.space_after = Pt(6)
    for t in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        add_runs(p, t, size); p.space_after = Pt(8)
    return tb


def new_slide(p, model, title):
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element))
            bar = s.shapes[-1]; runs = bar.text_frame.paragraphs[0].runs
            runs[0].text = title
            for r in runs[1:]: r._r.getparent().remove(r._r)
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    return s


def pic(slide, name, x, y, w):
    slide.shapes.add_picture(f'{FIG}/{name}', Inches(x), Inches(y), Inches(w))


SLIDES = [
    ('Band-pass RLC – why only a band passes', [
        'The output is across **R**: V_{out} = I × R. So V_{out} follows the **current**.',
        '**Below f_{0}:** X_{C} = 1/(2πfC) is big → C blocks → small current.',
        '**Above f_{0}:** X_{L} = 2πfL is big → L blocks → small current.',
        '**At f_{0}:** X_{L} = X_{C} (= Z_{0} ≈ 203 Ω). They are opposite, so they **cancel**. Only R_{total} limits the current → current and V_{out} are **maximum**.',
        'Highest V_{out}/V_{in} = R / (R + R_{L}): a little **below 1**, because part of the voltage drops on the coil wire R_{L}.',
        'So the circuit **passes a band** of frequencies around f_{0} and blocks the frequencies below and above it.'],
     [('rlc_xlxc.png', 7.25, 1.2, 5.7), ('d2_bp.png', 8.6, 5.05, 3.0)],
     'Band-pass is new for you, so we go slowly.\n'
     'The output is across the resistor. The voltage on R is current times R. So the output follows the current.\n'
     'Below f0, the capacitor has a big reactance. It blocks. The current is small.\n'
     'Above f0, the inductor has a big reactance. It blocks. The current is small.\n'
     'Look at the graph: the two lines cross at f0. There X L equals X C, about 203 ohms. They are opposite, so they cancel. '
     'Only R total limits the current, so the current is maximum and the output is maximum.\n'
     'The highest ratio is a little below one, because part of the voltage drops on the wire of the coil.\n'
     'So the circuit passes a band of frequencies around f0. That is why we call it band-pass.'),
    ('Band-pass – bandwidth (BW)', [
        'The edges of the band are **f_{L}** and **f_{H}**: there V_{out}/V_{in} = **0.707 × the highest value**.',
        'Why 0.707? 0.707 = 1/√2. At f_{L} and f_{H} the power in R is **half** of the maximum (the "−3 dB" points).',
        '**BW = f_{H} − f_{L}**   (in kHz).',
        'Theory: **BW = f_{0} / Q = R_{total} / (2πL)**. Bigger R → wider band. Smaller R → narrow, sharp band.',
        'f_{0} is not the average of f_{L} and f_{H}: **f_{0} = √(f_{L} × f_{H})**.'],
     [('rlc_bp_bw.png', 8.3, 1.12, 4.45), ('rlc_bp_q.png', 8.3, 4.08, 4.45)],
     'How wide is the band? We need its two edges, f L and f H.\n'
     'At the edges the ratio is 0.707 times the highest value. 0.707 is one over root two. There the power in the resistor is half of the maximum. '
     'Engineers call these the minus three dB points.\n'
     'The bandwidth is f H minus f L, in kilohertz.\n'
     'The theory says BW = f0 over Q, or R total over two pi L. A bigger R gives a wider band. Look at the lower graph: a bigger Q gives a narrower band.\n'
     'In the example, Q = 4: f L is about 28.6 and f H about 36.6 kilohertz, so BW is about 8.1 kilohertz.\n'
     'Be careful: f0 is not the average of f L and f H. It is the square root of f L times f H.'),
    ('Band-pass – measuring BW step by step', [
        '1.  Set the generator near **f_{0}** (≈ 32.4 kHz). Turn the frequency slowly and find the **highest V_{out}/V_{in}**. Write it down (example: 0.80).',
        '2.  Calculate the **target = 0.707 × highest** (example: 0.707 × 0.80 = 0.57).',
        '3.  Turn the frequency **down** until V_{out}/V_{in} = target. This frequency is **f_{L}**.',
        '4.  Turn the frequency **up**, past f_{0}, until V_{out}/V_{in} = target again. This is **f_{H}**.',
        '5.  **BW = f_{H} − f_{L}**. Check: √(f_{L} × f_{H}) should be close to f_{0}.'],
     [],
     'Now how to measure the bandwidth at your table. Five steps.\n'
     'One: go near f0 and turn the frequency slowly. Find the highest Vout over Vin. Write it down.\n'
     'Two: multiply it by 0.707. That is your target.\n'
     'Three: turn the frequency down until the ratio equals the target. That is f L.\n'
     'Four: turn it up, past f0, until the ratio equals the target again. That is f H.\n'
     'Five: BW is f H minus f L. Check that f0 is close to the square root of f L times f H.\n'
     'Three common mistakes. Do not use 0.707 times Vin: use 0.707 times the highest ratio. '
     'Read Vin again at f L and f H, because Vin changes near f0. '
     'And do not take f L and f H from the nine points on the card: they are too far apart. Turn the knob.'),
]
MISTAKES = ['Using 0.707 × V_{in}. Use 0.707 × the **highest V_{out}/V_{in}**.',
            'Reading V_{in} only once. Near f_{0}, V_{in} changes: read **CH1 and CH2** again at f_{L} and f_{H}.',
            'Taking f_{L} and f_{H} from the 9 points of your card: they are too far apart. **Turn the knob.**']


OLD = {'Low-pass RLC – output across C', 'High-pass RLC – output across L'}     # removed 24/10/2026
THREE = ['L = 1 mH, C = 22 nF ∥ 2.2 nF. The same current flows through R, L and C. Where you take the output decides the filter:',
         'Across **C** → **low-pass**, across **L** → **high-pass**: they work like the RC filters. New: near f_{0} the output '
         'can be **bigger than V_{in}** (a peak at f_{max}). Low-pass and high-pass teams measure f_{max} and the highest '
         'V_{out}/V_{in} in Task 2.',
         'Across **R** → **band-pass**: only a band of frequencies around f_{0} passes. This is new – the next slides explain it.',
         'Put the output component at the **ground side**, so both probe grounds are on the same node.']
THREE_SAY = ('One series RLC circuit can make three filters. It depends on where you take the output.\n'
             'Across C is a low-pass filter, across L is a high-pass filter. They work like the RC filters. '
             'New: near f0 the output can be bigger than the input: a peak at f max. '
             'Low-pass and high-pass teams measure f max and the highest Vout over Vin in Task 2.\n'
             'Across R is a band-pass filter: only a band of frequencies around f0 passes. This is new: the next slides explain it.\n'
             'Put the output part at the ground side, so both probe grounds are on the same point.')


def run():
    p = Presentation(DECK)
    titles = {t for t, *_ in SLIDES} | OLD
    lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):          # an earlier run: remove, then add again
        if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() in titles for sh in sl.shapes):
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    slides = list(p.slides)
    model = next(s for s in slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == AFTER for sh in s.shapes))
    at = slides.index(model) + 1
    tb = next(sh for sh in model.shapes if sh.has_text_frame and 'Where you take the output' in sh.text_frame.text)
    tf = tb.text_frame
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    for k, t in enumerate(THREE):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); add_runs(par, t, 16); par.space_after = Pt(8)
    model.notes_slide.notes_text_frame.text = THREE_SAY
    for k, (title, paras, figs, note) in enumerate(t for t in SLIDES if t[0] not in OLD):
        s = new_slide(p, model, title)
        if figs:
            text_box(s, 0.55, 1.2, 6.5 if figs[0][1] < 8 else 7.4, 5.7, paras, size=16)
        else:
            text_box(s, 0.55, 1.2, 7.3, 5.7, paras, size=17)
            text_box(s, 8.2, 1.3, 4.6, 4.2, ['•  ' + m for m in MISTAKES], size=15, box=True, title='Common mistakes')
        for name, x, y, w in figs:
            pic(s, name, x, y, w)
        s.notes_slide.notes_text_frame.text = note
        sid = lst[-1]; lst.remove(sid); lst.insert(at + k, sid)
    p.save(DECK)
    print('Lab 2:', len(p.slides), 'slides; RLC theory slides at', [i + 1 for i, s in enumerate(p.slides)
          if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() in titles for sh in s.shapes)])


if __name__ == '__main__':
    run()
