import json, io, contextlib, runpy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
with contextlib.redirect_stdout(io.StringIO()):
    runpy.run_path('bom.py'); g=runpy.run_path('design.py')
rows=json.load(open('bom.json'))
wb=Workbook(); ws=wb.active; ws.title='Danh sách mua'
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); thin=Side(style='thin',color='BBBBBB')
ws.append(['Linh kiện điện tử – Electric Circuit Lab HCM (6 cụm × 3 SV = 18 SV)'])
ws['A1'].font=Font(bold=True,size=14)
ws.append(['Nguyên tắc: linh kiện tiêu hao – mỗi buổi dùng linh kiện mới, không tái sử dụng sang buổi sau. Số "Cần" là tính đúng theo bài; "Mua" = Cần + % dư, làm tròn lên (điện trở theo bội 10, còn lại bội 5).'])
ws.append([])
hdr=['Linh kiện','Ngày 1','Ngày 2','Ngày 3','Ngày 4','Thi','Tổng cần','% dư','SỐ LƯỢNG MUA','Đơn giá','Thành tiền','Ghi chú']
ws.append(hdr)
for c in ws[4]: c.font=H; c.fill=HF; c.alignment=Alignment(horizontal='center',wrap_text=True)
notes={'Điện trở':'1/4W, màng kim loại 1% (vỏ xanh)','Tụ film':'Polyester/mylar 5%, 63–100V','NP0':'Ceramic NP0/C0G 5% (ổn định, sai số nhỏ)','Tụ gốm 100nF':'Ceramic 104 (X7R được) – dùng decoupling','Tụ hóa':'Điện phân, có cực, chân dài','Cuộn cảm':'Cuộn cảm đứng (radial) 10mH, ghi lại điện trở dây DCR','LM358':'DIP-8, nên mua hãng TI/ST','Electret':'2 chân, loại chân dài cắm breadboard','LED':'LED 5mm để kiểm tra chế độ đo diode'}
for i,r in enumerate(rows,5):
    n=next((v for k,v in notes.items() if k in r[0]),'')
    ws.append(list(r)+[None,f'=I{i}*J{i}',n])
last=4+len(rows)
ws.append(['TỔNG']+['']*9+[f'=SUM(K5:K{last})'])
ws.cell(last+1,1).font=Font(bold=True); ws.cell(last+1,11).font=Font(bold=True)
for row in ws.iter_rows(min_row=4,max_row=last+1):
    for c in row: c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
for i in range(5,last+1): ws.cell(i,9).font=Font(bold=True); ws.cell(i,9).fill=PatternFill('solid',fgColor='FFF2CC')
for col,w in zip('ABCDEFGHIJKL',[34,8,8,8,8,8,10,7,14,10,12,48]): ws.column_dimensions[col].width=w
ws.freeze_panes='B5'
o=last+3
extra=[('Không tiêu hao – kiểm tra lab đã có chưa',''),
('Dây cắm breadboard đực-đực (bộ 65 sợi)','6 bộ + 1 dự phòng'),
('Loa 8Ω 0.5W (nếu breadboard không có loa gắn sẵn)','6 + 2 dự phòng'),
('Kìm cắt/tuốt dây nhỏ','6'),
('Hộp chia ngăn đựng linh kiện (mỗi cụm 1 hộp/buổi)','6–8 hộp'),
('Băng keo giấy + bút lông (dán nhãn mã đề/khay)','1')]
for a,b in extra:
    ws.cell(o,1,a); ws.cell(o,9,b); o+=1
ws.cell(last+3,1).font=Font(bold=True)

w2=wb.create_sheet('Cách tính theo ngày')
lines=[('Ngày 1 (mỗi SV)','6 điện trở theo mã đề (3 điện trở "lạ", mỗi con ghép từ 2 điện trở chuẩn) + 1×2.2k + 1×4.7k cho cầu chia áp/dòng + 1 LED'),
('Ngày 2 (mỗi SV)','RC khởi động: 1R (1k…6.8k) + 1C (10/22/47nF). RLC: 1 cuộn 10mH + 2.2nF + 270pF (ghép song song ≈ 2.47nF → f0 ≈ 32 kHz) + R theo Q (Q=1:1k+1k, 1.5:1k+330, 2:1k, 3:680, 4:470, 5:390) + 47Ω, 100Ω để tinh chỉnh'),
('Ngày 3 (mỗi SV)','1 LM358 + cặp R theo mã đề (220/1.2k, 1.2k/47k, 3.6k/4.7k – mỗi cặp 6 SV) + 2×100nF decoupling. Mỗi cụm 2×470µF cho đường nguồn ±'),
('Ngày 4 (mỗi cụm)','Mạch mic→loa bản v2 của HN: 1 LM358, 1 electret, 5×10k (2 phân áp + 3 ghép 30k), 2×2.2k, 1×4.7k, 1×47k (bản v1), 1µF, 10µF, 1nF, 4×100nF gốm, 4×470µF (2 nguồn + 2 thử loa), 2×100nF film (thử loa)'),
('Thi (mỗi đội)','Khay chọn tự do để làm Section A (440Ω, 5k, 6.9k) và Section B (RC 100 Hz / 5 kHz): R 100…15k, C 10n…1µF – xem cột "Thi"'),
('Nguồn ±','Dùng máy nguồn có ngõ ra ± → bỏ phần cầu chia 2×1MΩ trong slide HN, không cần mua 1MΩ')]
w2.append(['Phần','Linh kiện dùng'])
for c in w2[1]: c.font=H; c.fill=HF
for l in lines: w2.append(list(l))
w2.column_dimensions['A'].width=20; w2.column_dimensions['B'].width=130
for r in w2.iter_rows(min_row=2):
    for c in r: c.alignment=Alignment(wrap_text=True,vertical='top')

w3=wb.create_sheet('Mã đề Ngày 1 (nháp)')
w3.append(['Mạch: (RA // RB) nối tiếp RC, nguồn 5V. X = nút giữa khối song song và RC.'])
w3.append(['Mã đề','RA (Ω)','Ghép RA','RB (Ω)','Ghép RB','RC (Ω)','Ghép RC','Rtđ (Ω)','I tổng (mA)','I qua RA (mA)','VX (V)'])
for c in w3[2]: c.font=H; c.fill=HF
f=lambda r: f'{r/1000:g}k' if r>=1000 else f'{r:g}'
for i,d in enumerate(g['D1'],1):
    row=[f'N1-{i:02d}']
    for k in 'ABC':
        v,t,a,b=d[k]; row+=[round(v),f'{f(a)} {"nối tiếp" if t=="+" else "//"} {f(b)}']
    row+=[round(d['Req']),round(d['I']*1e3,3),round(d['IA']*1e3,3),round(d['VX'],3)]
    w3.append(row)
for col,w in zip('ABCDEFGHIJK',[9,9,16,9,16,9,16,10,11,13,9]): w3.column_dimensions[col].width=w
wb.save('/Users/mac/Tra/SE9_ElecCircuitLab/HCM/docs/HCM_Linh_kien_can_mua.xlsx')
