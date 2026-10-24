"""Photos of the HCM instruments (taken from the Lab 1 deck) with numbered red marks, for the Lab 2 how-to slides:
the Tektronix AFG1022 function generator and the TBS1052C oscilloscope. Also three oscilloscope screens of a band-pass
filter below / at / above the peak. Writes fig/inst_gen.png, fig/inst_scope.png, fig/p2_bp_scope.png."""
import io, math
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

LAB1 = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab1/HCM Electronic Circuit Lab 1.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/'
RED = (220, 30, 30)


def photo(slide_no, min_w=2048 * 0.5):
    """The big photo of a Lab 1 slide (the slides of the hand-edited Lab 1 deck)."""
    s = Presentation(LAB1).slides[slide_no - 1]
    pics = [sh for sh in s.shapes if sh.shape_type == 13 and sh.image.size[0] >= min_w]
    return Image.open(io.BytesIO(max(pics, key=lambda sh: sh.width * sh.height).image.blob)).convert('RGB')


def font(size):
    for f in ('/System/Library/Fonts/Supplemental/Arial Bold.ttf', '/Library/Fonts/Arial Bold.ttf', '/System/Library/Fonts/Helvetica.ttc'):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            pass
    return ImageFont.load_default()


def mark(im, marks):
    """marks: (number, shape, box) with shape 'rect' or 'circle'; box in pixels of `im`. A numbered red tag next to it."""
    d = ImageDraw.Draw(im); f = font(30)
    for n, kind, (x0, y0, x1, y1), (tx, ty) in marks:
        (d.rectangle if kind == 'rect' else d.ellipse)((x0, y0, x1, y1), outline=RED, width=6)
        d.ellipse((tx - 22, ty - 22, tx + 22, ty + 22), fill=RED)
        w = d.textlength(str(n), font=f); d.text((tx - w / 2, ty - 18), str(n), font=f, fill='white')
    return im


def generator():
    im = photo(65).crop((300, 330, 2000, 1080)).resize((1360, 600))      # the AFG1022 front panel
    mark(im, [(1, 'rect', (278, 555, 336, 598), (250, 540)),             # Sine
              (2, 'rect', (683, 298, 730, 335), (645, 300)),              # button next to Freq/Period
              (3, 'rect', (762, 240, 962, 410), (742, 225)),              # number pad
              (4, 'rect', (683, 340, 730, 515), (645, 420)),              # buttons beside the screen: units (kHz, Vpp ...)
              (5, 'circle', (1000, 238, 1116, 352), (1150, 230)),         # round knob
              (6, 'rect', (990, 358, 1095, 400), (1130, 400)),            # ◄ ► keys
              (7, 'rect', (928, 435, 998, 480), (900, 500))])             # On/Off Out1
    im.crop((60, 170, 1330, 600)).save(FIG + 'inst_gen.png')           # without the blurred room above


def scope():
    im = photo(53, 1000).crop((0, 180, 1280, 800))                       # the TBS1052C front panel
    mark(im, [(1, 'rect', (998, 113, 1070, 152), (1100, 100)),            # Measure
              (2, 'circle', (800, 128, 905, 233), (780, 110)),            # Multipurpose knob
              (3, 'rect', (712, 188, 772, 462), (680, 320)),              # side buttons
              (4, 'circle', (835, 350, 890, 398), (815, 410)),            # CH1 menu
              (5, 'circle', (948, 350, 1003, 398), (1030, 410))])         # CH2 menu
    im.save(FIG + 'inst_scope.png')


def bp_scope():
    """Band-pass, three screens: below the peak, at the peak, above the peak. CH1 yellow = Vin, CH2 blue = Vout."""
    Q = 3.0
    fig, axs = plt.subplots(3, 1, figsize=(5.6, 6.6), dpi=170)
    for ax, (name, x) in zip(axs, (('below the peak', 1 / 1.35), ('at the peak', 1.0), ('above the peak', 1.35))):
        g = 0.92 / math.sqrt(1 + (Q * (x - 1 / x)) ** 2); ph = -math.atan(Q * (x - 1 / x))
        t = np.linspace(0, 2.5, 500)
        ax.set_facecolor('black')
        for k in range(1, 10): ax.axvline(k / 10 * t[-1], color='#333', lw=0.6)
        for k in (-0.75, -0.5, -0.25, 0.25, 0.5, 0.75): ax.axhline(k * 1.6, color='#333', lw=0.6)
        ax.axhline(0, color='#555', lw=0.8)
        ax.plot(t, np.sin(2 * np.pi * t), color='#f5d000', lw=2)
        ax.plot(t, g * np.sin(2 * np.pi * t + ph), color='#38b6ff', lw=2)
        ax.set_xlim(0, t[-1]); ax.set_ylim(-1.6, 1.6); ax.set_xticks([]); ax.set_yticks([])
        ax.text(0.01, 0.97, name, transform=ax.transAxes, color='white', fontsize=12, weight='bold', va='top')
        ax.text(0.99, 0.97, f'CH1 2.00 Vpp\nCH2 {2 * g:.2f} Vpp', transform=ax.transAxes, color='white', fontsize=9,
                ha='right', va='top', family='DejaVu Sans Mono')
    fig.tight_layout(h_pad=0.6); fig.savefig(FIG + 'p2_bp_scope.png'); plt.close(fig)


if __name__ == '__main__':
    generator(); scope(); bp_scope(); print('ok')
