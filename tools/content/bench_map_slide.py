"""Lab 1: slide "Find your table" with the room map (6 tables in a U, numbered clockwise, table 1 on the lecturer's
left; the lecturer's desk and the TV at the open end), right after "Your team and the student page". Same layout as public/js/core/benchmap.js.
Run after align_web.py; running again replaces the slide."""
import copy, io, re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle, FancyArrowPatch
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab1/HCM Electronic Circuit Lab 1.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/bench_map.png'
TITLE = 'Find your table'
AFTER = ('Your team and the student page (10 minutes)', 'Make your team (10 minutes)')
BENCHES = {1: (4, 50, 13, 19), 2: (4, 28, 13, 19), 3: (24, 6, 22, 13), 4: (54, 6, 22, 13), 5: (83, 28, 13, 19), 6: (83, 50, 13, 19)}
ORANGE, NAVY = '#C2410C', '#141E2B'


def draw():
    fig, ax = plt.subplots(figsize=(6.4, 5.9), dpi=170)
    ax.set_xlim(0, 100); ax.set_ylim(92, 0); ax.axis('off')                       # y down, like the SVG on the web
    ax.add_patch(Rectangle((1, 1), 98, 90, fill=False, ec='#999', lw=1.2))
    for n, (x, y, w, h) in BENCHES.items():
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=2', fc='white', ec=ORANGE, lw=2.2))
        ax.text(x + w / 2, y + h / 2, str(n), ha='center', va='center', fontsize=26, weight='bold', color=NAVY)
    ax.add_patch(FancyBboxPatch((36, 74), 28, 8, boxstyle='round,pad=0,rounding_size=1.5', fc='#eee', ec=NAVY, lw=1))
    ax.text(50, 78, 'Lecturer', ha='center', va='center', fontsize=12, color=NAVY)
    ax.add_patch(FancyBboxPatch((67, 75), 10, 6, boxstyle='round,pad=0,rounding_size=0.8', fc=NAVY, ec=NAVY))
    ax.text(72, 78, 'TV', ha='center', va='center', fontsize=11, color='white')
    ax.add_patch(FancyArrowPatch((23, 62), (77, 60), connectionstyle='arc3,rad=-0.9', arrowstyle='-|>', mutation_scale=14,
                                 color='#aaa', lw=1.3, ls=(0, (4, 3))))
    ax.text(50, 47, 'numbers go clockwise', ha='center', va='center', fontsize=11, color='#888')
    fig.tight_layout(pad=0.2); fig.savefig(FIG); plt.close(fig)


def run():
    import sys
    if '--force' not in sys.argv:           # Lab 1 was edited by hand on 24/10/2026: running this would undo it
        raise SystemExit('Lab 1 is edited by hand since 24/10/2026 - not changed (add --force to run anyway)')
    draw()
    p = Presentation(DECK); lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == TITLE for sh in sl.shapes):
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    slides = list(p.slides)
    model = next(s for s in slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() in AFTER for sh in s.shapes))
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element)); bar = s.shapes[-1]
            bar.text_frame.paragraphs[0].runs[0].text = TITLE
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    s.shapes.add_picture(FIG, Inches(6.55), Inches(1.15), height=Inches(5.75))
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(5.6), Inches(5)); tf = tb.text_frame; tf.word_wrap = True
    lines = ['Your **team number** = your **table number**.',
             '**Table 1** is on the lecturer\'s left. The numbers go **clockwise**: 1, 2 on the left, 3, 4 at the back, 5, 6 on the right.',
             'You keep the **same table** from Lab 1 to Lab 4.',
             'The student page shows your table in **orange** on this map.']
    for i, text in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.space_after = Pt(14)
        for part in re.split(r'(\*\*.*?\*\*)', text):
            if part:
                r = para.add_run(); r.text = part.strip('*'); r.font.size = Pt(20); r.font.name = 'Arial'
                r.font.bold = part.startswith('**'); r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    s.notes_slide.notes_text_frame.text = (
        'This is our room. Your team number is your table number.\n'
        'Table 1 is on my left. The numbers go clockwise: one and two on the left, three and four at the back, five and six on the right.\n'
        'You keep the same table from Lab 1 to Lab 4.\n'
        'On the student page you can see your table in orange on this map. Now go to your table.')
    sid = lst[-1]                                      # the slide just added: move it right after the team slide
    lst.remove(sid); lst.insert(slides.index(model) + 1, sid)
    p.save(DECK)
    print('Lab 1:', len(p.slides), 'slides; map slide at', [i + 1 for i, x in enumerate(p.slides)
          if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == TITLE for sh in x.shapes)])


if __name__ == '__main__':
    run()
