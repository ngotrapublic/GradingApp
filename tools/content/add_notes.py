import sys, runpy
sys.path.insert(0, "/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content")
from spoken import spoken
from pptx import Presentation
k = int(sys.argv[1]); src = sys.argv[2]; dst = sys.argv[3]
N = runpy.run_path(f'/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/notes{k}.py')['N']
RQ = lambda lab: f"""SAY: Let's review the new words from Lab {lab}.
SAY: I point to a word. You tell me what it means, and give me an example.
SAY: First word: {KWS[lab][0]}. What does it mean?"""
RA = lambda lab: f"""SAY: Here are the answers. Let's read them together.
SAY: We will use these words again today."""
KWS = {k: [w[0] for w in v] for k, v in runpy.run_path("/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/deck.py")["KEYWORDS"].items()}
seq = []                      # list of note texts in final order
if k == 1:
    seq = [N[i] for i in sorted(N)]
else:
    seq = [N[1], RQ(k - 1), RA(k - 1)]
    for i in sorted(N):
        if i == 1: continue
        seq.append(N[i])
        if k == 4 and i == 84:            # exam review section header
            seq += [RQ(4), RA(4)]
p = Presentation(src)
assert len(p.slides) == len(seq), (len(p.slides), len(seq))
nw = 0
for s, txt in zip(p.slides, seq):
    for sh in s.shapes:
        if sh.name == 'Key word callout':
            txt += f"\nWORD: {sh.text_frame.text.replace('New word: ', '')}\nTIP: Say the word slowly. Students repeat it once. Ask one student to explain it in their own words."
            nw += 1
    s.notes_slide.notes_text_frame.text = spoken(txt)
p.save(dst); print(k, len(seq), 'slides,', nw, 'word callouts')
