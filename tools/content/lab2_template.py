"""Excel template for Lab 2 Task 2 (band-pass, the team's own C and R): the student types the 9 frequencies of the card
(rows 6-14) and up to 11 extra points; Vout/Vin and the graph (log scale, 2-250 kHz) fill in by themselves."""
import shutil
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb=Workbook(); ws=wb.active; ws.title='Lab 2 data'
ws['A1']='Lab 2 – RLC frequency response'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']='Variant:'; ws['B2'].fill=PatternFill('solid',fgColor='FFF2CC')
ws['A3']='C and R (Task 2):'; ws['B3'].fill=PatternFill('solid',fgColor='FFF2CC')
hdr=['f (kHz)','Vin p-p (V) – CH1','Vout p-p (V) – CH2','Vout / Vin']
for i,h in enumerate(hdr,1):
    c=ws.cell(5,i,h); c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(wrap_text=True,horizontal='center')
thin=Side(style='thin',color='999999')
for r in range(6,26):
    for c in (1,2,3): ws.cell(r,c).fill=PatternFill('solid',fgColor='FFF2CC')
    ws.cell(r,4,f'=IF(OR(B{r}="",C{r}=""),NA(),C{r}/B{r})').number_format='0.000'
    for c in range(1,5): ws.cell(r,c).border=Border(top=thin,bottom=thin,left=thin,right=thin)
ws['A27']='Rows 6–14: the 9 frequencies of your card. Rows 15–25: extra points near the peak (any order is fine).'
ws['A29']='Conclusion (1–3 sentences):'; ws['A29'].font=Font(bold=True)
ws.merge_cells('A30:D34'); ws['A30'].fill=PatternFill('solid',fgColor='FFF2CC'); ws['A30'].alignment=Alignment(wrap_text=True,vertical='top')
for col,w in zip('ABCD',[12,16,16,12]): ws.column_dimensions[col].width=w
ch=ScatterChart(); ch.title='Frequency response'; ch.style=13
ch.x_axis.title='f (kHz, log scale)'; ch.y_axis.title='Vout / Vin'
ch.x_axis.scaling.logBase=10; ch.x_axis.scaling.min=2; ch.x_axis.scaling.max=250
ch.x_axis.delete=False; ch.y_axis.delete=False
s=Series(Reference(ws,min_col=4,min_row=6,max_row=25),Reference(ws,min_col=1,min_row=6,max_row=25),title='Vout/Vin')
s.marker.symbol='circle'; s.marker.size=7; s.graphicalProperties.line.noFill=True
ch.series.append(s); ch.legend=None; ch.width=18; ch.height=10
ws.add_chart(ch,'F5')
out = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/Lab2_RLC_data_template.xlsx'
wb.save(out)
shutil.copy(out, '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/public/courses/SE009/files/Lab2_RLC_data_template.xlsx'); print('ok')
