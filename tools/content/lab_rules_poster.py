"""Bench sign for the shared electronics lab (A4 landscape, laminated, one per bench), in the style of the Hanoi
lab poster: numbered rows with an icon and one short sentence. Left: RULES (10). Right: BEFORE YOU LEAVE (6) + photo.
Icons: poster_icons.py (Google Material Symbols). Writes HCM/docs/Lab_rules_poster.pptx; PDF via LibreOffice (end)."""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/docs/Lab_rules_poster.pptx'
PHOTO = f'{HERE}/fig/bench_tidy.jpg'
ICON = lambda n: f'{HERE}/fig/icons/{n}.png'
NAVY, GREEN, ORANGE = RGBColor(0x1B, 0x2A, 0x3A), RGBColor(0x2E, 0x7D, 0x32), RGBColor(0xF2, 0xA9, 0x00)
BOX, ROW, PAGE, WHITE = RGBColor(0xDC, 0xEF, 0xC8), RGBColor(0xE6, 0xE8, 0xE6), RGBColor(0xF2, 0xF3, 0xF2), RGBColor(0xFF, 0xFF, 0xFF)
W, H = 11.69, 8.27
FONT = 'Arial'

RULES = [
    ('power_settings_new', 'Switch the power **off** before you change a circuit.'),
    ('electric_bolt', 'Set the **current limit** and check **polarity** before you switch on.'),
    ('warning', 'Never leave a powered circuit alone. Smell, smoke or heat? **Switch off** and call the lecturer.'),
    ('electric_meter', 'Choose the right **multimeter mode and range**. Never measure current or resistance across a power source.'),
    ('electrical_services', 'Never open an instrument or touch mains plugs and sockets.'),
    ('fire_extinguisher', 'Know where the **main power switch** and the **fire extinguisher** are.'),
    ('no_food', 'No food or drinks. Keep bags off the benches and the walkways clear.'),
    ('desk', 'Use only your bench. Do not move instruments, cables or parts to another bench or out of the lab.'),
    ('report', 'Something missing or broken? **Tell the lecturer.** Do not repair it yourself.'),
    ('cleaning_services', 'Before you leave, **tidy your bench** (see "Before you leave").'),
]
LEAVE = [
    ('power_settings_new', 'Switch off **every instrument** and the trainer.'),
    ('developer_board', 'Clear the breadboard: **no parts, no wires**.'),
    ('inventory_2', 'Parts back in the parts box, in the right bag. Used or broken parts in the bin.'),
    ('cable', 'Test leads, probes and BNC cables into the box behind the trainer.'),
    ('check_circle', 'Instruments back in place, **like the photo**. Chairs in, rubbish in the bin.'),
    ('how_to_reg', '**Wait for the lecturer** to check your bench before you go.'),
]


def runs(p, text, size, color=NAVY, bold_all=False):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if part:
            r = p.add_run(); r.text = part.strip('*') if part.startswith('**') else part
            r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = color
            r.font.bold = bold_all or part.startswith('**')


def shape(s, kind, x, y, w, h, fill, radius=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
    style = sh._element.find(qn('p:style'))             # the theme style adds a shadow in some viewers: drop it
    if style is not None:
        sh._element.remove(style)
    if radius is not None:
        sh.adjustments[0] = radius
    return sh


def text(sh, t, size, color=NAVY, align=PP_ALIGN.LEFT, bold_all=False, margin=0.12):
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align; runs(p, t, size, color, bold_all)


def row(s, x, y, w, h, num, icon, t, size):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.25, y, w - 0.25, h, ROW, 0.25)          # grey band behind number + icon
    pill = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.04, y + (h - 0.32) / 2, 0.5, 0.32, NAVY, 0.5)
    text(pill, f'{num:02d}', 13, ORANGE, PP_ALIGN.CENTER, bold_all=True, margin=0.0)
    ic = min(h - 0.08, 0.44)
    s.shapes.add_picture(ICON(icon), Inches(x + 0.72), Inches(y + (h - ic) / 2), Inches(ic), Inches(ic))
    box = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 1.25, y, w - 1.25, h, BOX, 0.22)
    text(box, t, size)


def heading(s, x, y, w, t):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    p = tb.text_frame.paragraphs[0]; runs(p, t, 17, GREEN, bold_all=True)
    shape(s, MSO_SHAPE.RECTANGLE, x + 0.05, y + 0.42, 1.6, 0.05, GREEN)


prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
s = prs.slides.add_slide(prs.slide_layouts[6])
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, PAGE)
title = s.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(W - 0.9), Inches(0.6))
p = title.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; runs(p, 'ELECTRONICS LAB RULES', 32, NAVY, bold_all=True)
sub = s.shapes.add_textbox(Inches(0.45), Inches(0.82), Inches(W - 0.9), Inches(0.35))
p = sub.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
runs(p, 'This lab is shared by many classes. Leave your bench the way you want to find it.', 14, GREEN)

LX, LW, RX, RW, TOP = 0.4, 6.55, 7.2, 4.1, 1.42
heading(s, LX, TOP, LW, 'RULES')
rh, gap = 0.5, 0.085
for i, (icon, t) in enumerate(RULES):
    row(s, LX, TOP + 0.6 + i * (rh + gap), LW, rh, i + 1, icon, t, 11.5)
heading(s, RX, TOP, RW, 'BEFORE YOU LEAVE')
lh = 0.52
for i, (icon, t) in enumerate(LEAVE):
    row(s, RX, TOP + 0.6 + i * (lh + gap), RW, lh, i + 1, icon, t, 10.5)
py = TOP + 0.6 + 6 * (lh + gap) + 0.05
pic = s.shapes.add_picture(PHOTO, Inches(RX), Inches(py), Inches(RW))
band = shape(s, MSO_SHAPE.RECTANGLE, 0, H - 0.42, W, 0.42, GREEN)
text(band, '+    ' * 30, 12, WHITE, PP_ALIGN.CENTER, bold_all=True, margin=0)
prs.save(OUT); print('saved', OUT, '| photo bottom at', round(py + pic.height / 914400, 2), 'in; band at', H - 0.42)
# PDF:  soffice --headless --convert-to pdf --outdir HCM/docs HCM/docs/Lab_rules_poster.pptx
