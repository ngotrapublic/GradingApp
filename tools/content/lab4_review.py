"""Lab 4 (lecturer, 29/10/2026: students are slower than planned; the exam has no decoupling).
1. No decoupling lesson: the slides "Capacitors" (coupling and decoupling) and "Decoupling capacitor" go; the recap line
   "Decouple both supply rails" goes; the overview says "Coupling capacitor". The 0.1 uF capacitors at the op-amp pins in
   "Final circuit version 2" stay (they are in the HN schematics), without the word "decoupling".
2. Exam practice after "Skills checklist for the exam" (not marked, after the Lab 4 marks): Exercise 6 combine
   resistors, Exercise 7 set up the instruments from the default state (as the exam starts), Exercise 8 design an RC
   low-pass filter with the parts on the table. Each has its answer slide (practice_slides.add / answer_slide).
3. One final circuit instead of version 1 + version 2 (see one_final).
Run after practice_slides.py (`practice_slides.py 4` removes the practice slides), before subscripts.py. Safe to run again."""
import copy, re
import schemdraw, schemdraw.elements as elm
from pptx.oxml.ns import qn
from pptx import Presentation
import practice_slides as ps

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab4/HCM Electronic Circuit Lab 4.pptx'
NAME = 'Lab 4 exam practice'
AFTER = 'Skills checklist for the exam'
DROP = ('Depending on how you connect it, we have:', 'When you only wants DC signal to go through',
        'Decoupling capacitors in microphone circuit', 'We build a similar circuit today.')
RECAP = 'Decouple both supply rails'

EX = [
    dict(bar='Exam practice', fig='p4_combine.png',
         head='Exercise 6: Combine resistors to get a target value (5 minutes)',
         steps=['Your marks are done. You can take the resistors out of your amplifier.',
                'Make 5 kΩ from two 10 kΩ resistors.',
                'Make 12.2 kΩ from the resistors on your table.',
                'Make about 3.2 kΩ from one 10 kΩ and one 4.7 kΩ resistor.',
                'Calculate each value. Then measure it with the multimeter. Write down both.'],
         ask='How did you connect them? How close is each measured value to the target?',
         answer=['**5 kΩ:** two 10 kΩ in **parallel**: 10 × 10 / (10 + 10) = 5 kΩ.',
                 '**12.2 kΩ:** 10 kΩ and 2.2 kΩ in **series**: 10 + 2.2 = 12.2 kΩ.',
                 '**3.2 kΩ:** 10 kΩ and 4.7 kΩ in **parallel**: 10 × 4.7 / (10 + 4.7) ≈ 3.2 kΩ.',
                 'The measured values are usually within 1–2% of the target. In the exam, write the measured values on your schematic.']),
    dict(bar='Exam practice', fig='inst_gen.png',
         head='Exercise 7: Set up the instruments from the start (5 minutes)',
         steps=['Press Default Setup on the oscilloscope. Set the generator load to 50 Ω (Utility → Output Setup).',
                'Now another member sets everything again, without help.',
                'Generator: High-Z, sine, 3 Vpp, 1 kHz, offset 0 V. Connect it to CH1 with the 1X BNC cable.',
                'Oscilloscope: set CH1 to 1X. Measure the frequency and the Pk-Pk of CH1.',
                'The screen must show 1 kHz and 3 Vpp.'],
         ask='What do you see if you forget High-Z? And if you forget the 1X setting?',
         answer=['The exam starts like step 1: the oscilloscope in its default state and the generator at 50 Ω. Check both first.',
                 '**Generator at 50 Ω:** the real output is **2 times** the number you set: 6 Vpp instead of 3 Vpp.',
                 '**Oscilloscope at 10X (the default) with a 1X cable:** the reading is **10 times too big**: 30 Vpp instead of 3 Vpp.']),
    dict(bar='Exam practice', fig='p4_rc_design.png',
         head='Exercise 8: Design a low-pass filter for fc = 1 kHz (10 minutes)',
         steps=['Use one 100 nF capacitor and the resistors on your table. You can combine resistors.',
                'Choose R so that fc = 1 / (2πRC) is within ±20% of 1 kHz (800 Hz – 1.2 kHz).',
                'Build the filter. Generator: sine, 2 Vpp, High-Z. CH1 = Vin, CH2 = Vout (across C).',
                'Measure Vout at 100 Hz, at your fc and at 10 kHz.'],
         ask='Which R did you choose? Is Vout at your fc close to 0.707 × Vin?',
         answer=['You need R = 1 / (2π × 1 kHz × 100 nF) ≈ **1.6 kΩ**. There is no 1.6 kΩ resistor, so combine two.',
                 '**2.2 kΩ in parallel with 4.7 kΩ ≈ 1.5 kΩ:** fc ≈ 1.06 kHz. **10 kΩ in parallel with 2.2 kΩ ≈ 1.8 kΩ:** fc ≈ 880 Hz. Both are within ±20%.',
                 'At your fc, Vout ≈ 0.707 × Vin (about 1.4 Vpp). At 100 Hz, Vout ≈ Vin. At 10 kHz, Vout is about 10 times smaller.']),
]


def figures():
    schemdraw.config(fontsize=16, lw=2)
    with schemdraw.Drawing(file=ps.FIG + 'p4_combine.png', dpi=200, show=False) as d:
        elm.Label().at((2.25, 1.3)).label('Series: R = R1 + R2')
        elm.Dot(open=True).at((0, 0))
        elm.Resistor().right().length(2.25).label('R1')
        elm.Resistor().right().length(2.25).label('R2')
        elm.Dot(open=True)
        elm.Label().at((2.25, -2.2)).label('Parallel: R = R1 × R2 / (R1 + R2)')
        elm.Dot(open=True).at((0, -4.5))
        elm.Line().right(1); a = d.here
        elm.Line().up(1); elm.Resistor().right().length(2.5).label('R1'); elm.Line().down(1); b = d.here
        elm.Line().at(a).down(1); elm.Resistor().right().length(2.5).label('R2', loc='bottom'); elm.Line().up(1)
        elm.Line().at(b).right(1); elm.Dot(open=True)
    with schemdraw.Drawing(file=ps.FIG + 'p4_rc_design.png', dpi=200, show=False) as d:
        elm.SourceSin().up().length(3).label('Function\ngenerator', loc='bottom', ofst=0.3)
        elm.Dot().label('CH1', loc='left', ofst=0.15)
        elm.Resistor().right().length(3.5).label('R = ?')
        top = d.here; elm.Dot()
        elm.Capacitor().down().length(3).label('C = 100 nF\n(104)', loc='bottom')
        elm.Line().left(3.5)
        elm.Line().at(top).right(2); elm.Dot(open=True).label('CH2\n(across C)', loc='right')
        elm.Line().at((3.5, 0)).right(2); elm.Dot(open=True)
        elm.Ground().at((0, 0))


# 3. One final circuit (lecturer, 29/10/2026): version 1 "dies out after about 1 minute" only with the HN home-made
#    supply (12 V + 2 x 1 MOhm divider as 0 V); on the PB-503A it works. The marks use the version 2 values
#    (4.7 kOhm, 1 nF), so version 1 goes and version 2 becomes "Final circuit". Slides are found by old or new text.
V1 = ('Final circuit (version 1)', 'Final circuit version 1')
FINAL = {   # old first line -> new lines of the text box ("**x**" = bold)
    'So how do we solve the echo?': [
        '**Before you connect the microphone**',
        'Your voice gives tens to hundreds of mVpp (Exercise 4).',
        'With a gain of 291, the output would need tens of volts. It clips (Lab 3).',
        'The speaker is also very close to the microphone. A loud speaker gives an echo (acoustic feedback).',
        'So we make three changes.'],
    'The resulting circuit is this': [
        '**The final circuit is this**',
        'Disconnect the function generator.',
        'Connect the microphone circuit to stage 1 through the 1 µF capacitor.',
        'Connect the final output to the speaker.'],
    'You will see that the echo is much less': [
        'Speak into the microphone through the rolled paper.',
        'You hear your voice in the speaker, with little echo.',
        'The output signal looks something like this'],
}
FINAL_NOTES = {
    'Before you connect the microphone': 'Before you connect the microphone: your voice gives tens to hundreds of millivolts. '
        'With a gain of 291, the output would need tens of volts, so it clips, like in Lab 3. The speaker is also very '
        'close to the microphone. A loud speaker gives an echo. So we make three changes.',
    'The final circuit is this': 'This is the final circuit. Disconnect the function generator. Connect the microphone '
        'circuit to stage 1 through the 1 microfarad capacitor. Connect the final output to the speaker.',
    'Speak into the microphone through': 'Speak into the microphone through the rolled paper. You hear your voice in the '
        'speaker, with little echo. The output signal looks like this.',
}


def set_lines(sh, lines):
    """Replace the text of an HN text box: one paragraph per line with an empty paragraph between (as in HN)."""
    tf = sh.text_frame; paras = tf.paragraphs
    tpl = copy.deepcopy(paras[0]._p)
    for r in tpl.findall(qn('a:r')) + tpl.findall(qn('a:br')): tpl.remove(r)
    body = paras[0]._p.getparent()
    for para in paras: body.remove(para._p)
    for k, line in enumerate(lines):
        if k:
            body.append(copy.deepcopy(tpl))
        p_ = copy.deepcopy(tpl); body.append(p_)
        from pptx.text.text import _Paragraph
        para = _Paragraph(p_, tf)
        for part in re.split(r'(\*\*.*?\*\*)', line):
            if part:
                r = para.add_run(); r.text = part.strip('*'); r.font.bold = part.startswith('**')


def one_final(p):
    gone = drop(p, lambda s: any(v in ps.text_of(s) for v in V1))
    for s in p.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for r in (r for para in sh.text_frame.paragraphs for r in para.runs):
                if r.text in ('Final circuit version 2', 'Final circuit (version 2)'):
                    r.text = 'Final circuit'
            first = sh.text_frame.paragraphs[0].text
            if first in FINAL:
                set_lines(sh, FINAL[first])
        for key, say in FINAL_NOTES.items():
            if key in ps.text_of(s):
                s.notes_slide.notes_text_frame.text = say
        nt = s.notes_slide.notes_text_frame if s.has_notes_slide else None
        if nt is not None and nt.text in ('Final circuit, version 2.', 'This is version 2 of the schematic.'):
            nt.text = {'Final circuit, version 2.': 'The final circuit.',
                       'This is version 2 of the schematic.': 'This is the schematic of the final circuit.'}[nt.text]
    return gone


def title(s):
    return next((sh.text_frame.text.strip() for sh in s.shapes if sh.name == 'Title Bar'), '')


def drop(p, test):
    lst = p.slides._sldIdLst; n = 0
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if test(sl):
            p.part.drop_rel(sid.rId); lst.remove(sid); n += 1
    p.part.rename_slide_parts([x.rId for x in lst])
    return n


def run():
    figures()
    p = Presentation(DECK); lst = p.slides._sldIdLst
    drop(p, lambda s: s._element.cSld.get('name') == NAME)
    gone = drop(p, lambda s: any(d in ps.text_of(s) for d in DROP))
    v1 = one_final(p)
    # recap: no decoupling line on the slide or in the notes
    rec = next(s for s in p.slides if title(s).startswith('Recap – the inverting amplifier'))
    for sh in rec.shapes:
        if sh.has_text_frame:
            for para in list(sh.text_frame.paragraphs):
                if para.text.startswith(RECAP):
                    para._p.getparent().remove(para._p)
    nt = rec.notes_slide.notes_text_frame
    nt.text = nt.text.replace(', decoupling on both rails.', '.')
    # overview: step 1 is the coupling capacitor only; step 4 has the exam practice
    ov = next(s for s in p.slides if title(s) == 'Lab 4 overview')
    for sh in ov.shapes:
        if sh.has_text_frame:
            for r in (r for para in sh.text_frame.paragraphs for r in para.runs):
                r.text = r.text.replace('Coupling and decoupling capacitors', 'Coupling capacitor') \
                               .replace('Break & fix practice', 'Exam practice: Exercises 6–8')
    # the exam practice
    model = next(s for s in p.slides if title(s) == ps.MODEL)
    ref = next(s for s in p.slides if title(s) == AFTER)
    new = []
    for ex in EX:
        new += [ps.add(p, model, ex), ps.answer_slide(p, model, ex)]
    for s in new:
        s._element.cSld.set('name', NAME)
    sids = [lst[list(p.slides).index(x)] for x in new]
    for sid in sids: lst.remove(sid)
    at = list(p.slides).index(ref) + 1
    for k, sid in enumerate(sids): lst.insert(at + k, sid)
    p.save(DECK)
    left = [i + 1 for i, s in enumerate(p.slides) if 'ecoupl' in ps.text_of(s) and s._element.get('show') != '0']
    print(f'Lab 4: {gone} decoupling slides removed | {v1} version-1 slides removed | exam practice at {at + 1}-{at + len(new)} | {len(p.slides)} slides |'
          f' "decoupl" still on visible slides: {left}')


if __name__ == '__main__':
    run()
