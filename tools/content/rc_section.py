"""Lab 2, part 1 (first-order RC) and the order of the lab (lecturer, 24/10/2026):
- one circuit for the whole part: R = 10 kOhm + C = 100 nF (104), 2 Vpp, fc = 1/(2 pi R C) = 159 Hz (not the Hanoi
  47k + 47k + 0.1 uF circuit);
- simple theory (no impedance calculation): oscilloscope view, voltage divider, high-pass, the two response graphs +
  "Reading the graph – the -3 dB point", cut-off formula;
- Task 1 moves down to just before Task 2 ("Lab 2 exercise"): the marked tasks are at the end of the lab; the practice exercises stay in the lesson
  (Exercise 1 -> table -> "What do you notice?" -> why; Exercise 3 after the cut-off formula);
- every idea in the speaker notes is also on the slide.
Removed: the Hanoi photo "For the first circuit", "analyse these circuits" (impedance), the earlier "Why 0.707?" slide.
Run after practice_slides.py, then subscripts.py. Running again gives the same slides."""
import copy, io, math, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/HCM Electronic Circuit Lab 2.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/'
R, C = 10e3, 100e-9
FC = 1 / (2 * math.pi * R * C)
ROWS = [10, 100, 200, 1000, 10000]
MARK = 'RC Section'                 # shapes added by this script (removed and added again on every run)
PRACTICE = 'Practice Exercise'      # marker of the practice slides (practice_slides.py)
MODEL = 'Devices and components for this session'
DARK = RGBColor(0x22, 0x22, 0x22)
fmt = lambda f: f'{f / 1000:g} kHz' if f >= 1000 else f'{f:g} Hz'
gain = {'lp': lambda f: 1 / np.sqrt(1 + (f / FC) ** 2), 'hp': lambda f: (f / FC) / np.sqrt(1 + (f / FC) ** 2)}
xc = lambda f: 1 / (2 * math.pi * f * C)
ohm = lambda r: f'{r / 1000:.3g} kΩ' if r >= 1000 else f'{r:.3g} Ω'

# ---------------------------------------------------------------- slide text (notes = the same ideas, spoken)
PLAN = ['In this part we use one circuit: **R = 10 kΩ, C = 100 nF (104)**. Function generator → R → C.',
        '1. Exercise 1: R then C – change the frequency, watch Vout.',
        '2. Why does Vout change? A voltage divider.',
        '3. Exercise 2: swap R and C.',
        '4. Reading the graph: the −3 dB point.',
        '5. Exercise 3: measure the cut-off frequency fc.',
        'Task 1 (marked, on the student page) is at the end of the lab.']
SCOPE = ['**CH1** (yellow) = Vin. **CH2** (blue) = Vout across C. Vin is 2 Vpp at every frequency.',
         '**10 Hz**: Vout is as big as Vin.',
         f'**{FC:.0f} Hz (fc)**: Vout is smaller, and it comes a little later than Vin.',
         '**1 kHz**: Vout is very small.',
         'At each frequency, turn the time base (s/div) so that you see 2–3 waves.']
DIVIDER = ['**Lab 1:** in a voltage divider, the bigger resistor takes the bigger part of the voltage.',
           'The RC filter is a divider too: **R on top, C at the bottom**. Vout is the voltage on C.',
           'C acts like a resistor that **changes with frequency**: XC = 1 / (2πfC). The table shows the values.',
           '**Low frequency:** XC is much bigger than R → C takes almost all of Vin → **Vout ≈ Vin**.',
           '**High frequency:** XC is much smaller than R → C takes almost nothing → **Vout ≈ 0**.']
SWAP = ['R and C are swapped. The output (CH2) is now across R.']
READ = ['**Top: gain in dB** = 20 log (Vout / Vin). **0 dB** means Vout = Vin.',
        '**Pass band** (left): gain ≈ 0 dB. The signal passes.',
        '**Corner frequency fc**: the gain is **-3 dB**. −3 dB means Vout = **0.707 × Vin** (half of the power).',
        '**Why there?** At −3 dB the output is still about 70% of the input: it has not dropped much yet. '
        'After fc it drops fast: **-20 dB per decade** (10 × the frequency → 10 × smaller Vout). '
        'So fc is the border between "passes" and "blocked".',
        '**Stop band** (right): the signal is blocked.',
        '**Bottom: phase.** Vout comes later than Vin. At fc the shift is −45°.',
        'High-pass: the same graph, mirrored.']
CUTOFF_LP = ['The cut-off frequency – the same formula for low-pass and high-pass:',
             '**fc = 1 / (2πRC)**',
             'At fc: Vout = 0.707 × Vin, the **-3 dB point**.',
             f'Our circuit: R = 10 kΩ, C = 100 nF → fc ≈ {FC:.0f} Hz.']
CUTOFF_HP = ['The **same formula** for the high-pass filter: fc = 1 / (2πRC).',
             f'Same R and C → **same fc** ≈ {FC:.0f} Hz.',
             'At fc: Vout = 0.707 × Vin (−3 dB), as for the low-pass filter.']
NOTICE = {
 'lp': ['**Low frequency** (10 Hz): Vout ≈ Vin. The signal **passes**.',
        '**High frequency** (10 kHz): Vout is almost 0. The signal is **blocked**.',
        f'**In between**, Vout falls. It falls around one frequency: the **cut-off frequency fc** (here about {FC:.0f} Hz).',
        '→ This circuit is a **low-pass filter**: it passes low frequencies.',
        'Why? See the next slides.'],
 'hp': ['**Low frequency** (10 Hz): Vout is almost 0. The signal is **blocked**.',
        '**High frequency** (10 kHz): Vout ≈ Vin. The signal **passes**.',
        f'**Same R and C → same fc** (about {FC:.0f} Hz).',
        '**Why?** The same divider, but now Vout is on R. Low f: XC is big, C takes most of Vin, little is left for R. '
        'High f: XC is small, R takes almost all of Vin.',
        '→ This circuit is a **high-pass filter**: it passes high frequencies.'],
}
OVERVIEW = {'RC filters': ['RC filters', '1st-order low-pass and high-pass', 'Cut-off frequency fc', 'Exercises 1–3, −3 dB point'],
            'box4': ['Marked tasks (end of the lab)', 'Task 1: RC cut-off (individual)', 'Task 2: series RLC (team)', 'Measure, plot and conclude']}
BOX4 = ('Lab 2 exercise', 'Practice, then marks', 'Marked tasks (end of the lab)')     # first line of box 4, old or new
OVERVIEW_SAY = ('Today has 4 parts: RC filters with Exercises 1 to 3 and the minus 3 dB point, capacitors, series RLC, '
                'and at the end of the lab the marked tasks: Task 1, RC cut-off, individual, and Task 2, series RLC, as a team: '
                'measure, plot and conclude.\n'
                'Reminder: at the start, the instruments are in their default state. Before you measure: check 1X or 10X on the '
                'oscilloscope, and set the function generator to High-Z.')


def spoken(lines):
    return '\n'.join(re.sub(r'\*\*', '', l) for l in lines)


# ---------------------------------------------------------------- figures
def notice_graph(kind):
    f = np.logspace(0, 5, 400); g = gain[kind]
    fig, ax = plt.subplots(figsize=(6.0, 4.4), dpi=170)
    ax.semilogx(f, g(f), color='#1f77b4', lw=2.5)
    ax.axvline(FC, color='#c0392b', ls=':', lw=1.5)
    ax.text(FC * 1.15, 0.05 if kind == 'lp' else 0.12, f'$f_c$ ≈ {FC:.0f} Hz', color='#c0392b', fontsize=11)
    for x in ROWS:
        ax.plot(x, g(x), 'o', color='#C2410C', ms=7)
        ax.annotate(f'{g(x):.2f}', (x, g(x)), textcoords='offset points', xytext=(6, 6), fontsize=9, color='#444')
    ax.set_xlim(1, 1e5); ax.set_ylim(0, 1.1)
    ax.set_xticks([1, 10, 100, 1000, 10000, 100000]); ax.set_xticklabels(['1 Hz', '10 Hz', '100 Hz', '1 kHz', '10 kHz', '100 kHz'])
    ax.set_ylabel('$V_{out}$ / $V_{in}$'); ax.set_xlabel('frequency (log scale)')
    ax.set_title(('Low-pass' if kind == 'lp' else 'High-pass') + ': R = 10 kΩ, C = 100 nF', fontsize=12)
    ax.grid(True, which='major', alpha=0.3)
    for s in ('top', 'right'): ax.spines[s].set_visible(False)
    fig.tight_layout(); path = FIG + f'p2_rc_{kind}_graph.png'; fig.savefig(path); plt.close(fig)
    return path


def scope_figure():
    """Three oscilloscope screens of the low-pass filter: 10 Hz, fc, 1 kHz. CH1 yellow = Vin, CH2 blue = Vout."""
    fig, axs = plt.subplots(3, 1, figsize=(5.6, 6.6), dpi=170)
    for ax, f in zip(axs, (10, FC, 1000)):
        g = 1 / math.sqrt(1 + (f / FC) ** 2); ph = -math.atan(f / FC)
        t = np.linspace(0, 2.5 / f, 500)
        ax.set_facecolor('black')
        for k in range(1, 10): ax.axvline(k / 10 * t[-1], color='#333', lw=0.6)
        for k in (-0.75, -0.5, -0.25, 0.25, 0.5, 0.75): ax.axhline(k * 1.6, color='#333', lw=0.6)
        ax.axhline(0, color='#555', lw=0.8)
        ax.plot(t, np.sin(2 * np.pi * f * t), color='#f5d000', lw=2)
        ax.plot(t, g * np.sin(2 * np.pi * f * t + ph), color='#38b6ff', lw=2)
        ax.set_xlim(0, t[-1]); ax.set_ylim(-1.6, 1.6); ax.set_xticks([]); ax.set_yticks([])
        name = f'{f:.0f} Hz ($f_c$)' if f == FC else fmt(f)
        ax.text(0.01, 0.97, name, transform=ax.transAxes, color='white', fontsize=12, weight='bold', va='top')
        ax.text(0.99, 0.97, f'CH1 2.00 Vpp\nCH2 {2 * g:.2f} Vpp', transform=ax.transAxes, color='white', fontsize=9,
                ha='right', va='top', family='DejaVu Sans Mono')
    fig.tight_layout(h_pad=0.6); path = FIG + 'p2_rc_scope.png'; fig.savefig(path); plt.close(fig)
    return path


# ---------------------------------------------------------------- slide helpers
def text_of(s):
    return '\n'.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)


def find(p, snippet, required=True):
    hits = [s for s in p.slides if snippet in text_of(s)]
    if len(hits) > 1:
        hits = [s for s in hits if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == snippet for sh in s.shapes)]
    if not required and not hits:
        return None
    assert len(hits) == 1, (snippet, len(hits))
    return hits[0]


def is_practice(s):
    return any(sh.name == PRACTICE for sh in s.shapes)


def add_runs(para, text, size=16, color=DARK):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if part:
            r = para.add_run(); r.text = part.strip('*'); r.font.size = Pt(size); r.font.name = 'Rockwell'
            r.font.bold = part.startswith('**'); r.font.color.rgb = color


def textbox(s, x, y, w, h, lines, size=16, gap=10):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.name = MARK
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); add_runs(para, line, size); para.space_after = Pt(gap)
    return tb


def picture(s, path, x, y, w, h, top=False):
    if isinstance(path, (bytes, bytearray)):
        src = io.BytesIO(path); iw, ih = Image.open(io.BytesIO(path)).size
    else:
        src = path; iw, ih = Image.open(path).size
    k = min(w / iw, h / ih); pw, ph = iw * k, ih * k
    pic = s.shapes.add_picture(src, Inches(x + (w - pw) / 2), Inches(y if top else y + (h - ph) / 2), Inches(pw), Inches(ph))
    pic.name = MARK
    return pic


def clear(s):
    for sh in list(s.shapes):
        if (sh.name.startswith(MARK) and sh.name != MARK + ' fig') or sh.name.startswith('RC Notice'):   # RC Notice: older version
            sh._element.getparent().remove(sh._element)


def big_pictures(s):
    return sorted([sh for sh in s.shapes if sh.shape_type == 13 and sh.width > Inches(2.2) and not sh.name.startswith('Logo')
                   and (not sh.name.startswith(MARK) or sh.name == MARK + ' fig')], key=lambda sh: sh.left)


def swap_pictures(s, paths):
    """Replace the big pictures of a slide (left to right) by our figures, fitted into the same boxes."""
    old = big_pictures(s)
    assert len(old) == len(paths), (len(old), paths)
    for sh, path in zip(old, paths):
        box = [Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height)]
        sh._element.getparent().remove(sh._element)
        picture(s, path, *box).name = MARK + ' fig'


def new_slide(p, model, title):
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element)); bar = s.shapes[-1]
            bar.text_frame.paragraphs[0].runs[0].text = title
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
            bar.name = MARK + ' bar'
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height).name = MARK + ' logo'
    return s


def move_after(p, slides, ref):
    """Put `slides` (in this order) right after `ref`."""
    lst = p.slides._sldIdLst
    sids = [lst[list(p.slides).index(s)] for s in slides]
    for sid in sids: lst.remove(sid)
    at = list(p.slides).index(ref) + 1
    for k, sid in enumerate(sids): lst.insert(at + k, sid)


def move_before(p, slides, ref):
    lst = p.slides._sldIdLst
    sids = [lst[list(p.slides).index(s)] for s in slides]
    for sid in sids: lst.remove(sid)
    at = list(p.slides).index(ref)
    for k, sid in enumerate(sids): lst.insert(at + k, sid)


def set_notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def put(cell, runs):
    """Write the cell text into its existing runs (keeps the table style): first runs get the text, the rest are emptied."""
    rs = [r for para in cell.text_frame.paragraphs for r in para.runs]
    for r, t in zip(rs, runs + [''] * (len(rs) - len(runs))):
        r.text = t


# ---------------------------------------------------------------- build
def run():
    lp, hp = FIG + 'p2_lp.png', FIG + 'p2_hp.png'
    p = Presentation(DECK); lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        t = text_of(sl)
        if any(sh.name == MARK + ' bar' for sh in sl.shapes) or 'For the first circuit' in t or 'analyse these circuits' in t:
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    model = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == MODEL for sh in s.shapes))

    # overview: the new order of the lab
    s = find(p, 'Lab 2 overview')
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        paras = [para for para in sh.text_frame.paragraphs if para.runs]
        first = paras[0].runs[0].text if paras else ''
        lines = OVERVIEW['RC filters'] if first == 'RC filters' else OVERVIEW['box4'] if first in BOX4 else None
        if lines:
            assert len(paras) == len(lines), (first, len(paras))
            for para, t in zip(paras, lines):
                para.runs[0].text = t
                for r in para.runs[1:]: r._r.getparent().remove(r._r)
    set_notes(s, OVERVIEW_SAY)

    # plan (the old Hanoi step slide)
    s = (find(p, 'We now construct RC circuits', False) or find(p, 'We build RC circuits', False)
         or find(p, 'In this part we use one circuit'))
    for sh in list(s.shapes):
        if sh.has_text_frame and 'We now construct RC circuits' in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)
    clear(s); textbox(s, 0.47, 1.5, 5.0, 5.5, PLAN, gap=12)
    swap_pictures(s, [lp])
    set_notes(s, spoken(PLAN))
    plan = s

    # breadboard photo
    s = find(p, 'Example of breadboard circuit')
    clear(s); textbox(s, 0.47, 2.3, 3.2, 1.5, ['The output (CH2) is across the capacitor.'])
    set_notes(s, 'Here is an example on the breadboard. The output (CH2) is across the capacitor.')
    photo = s

    # new: oscilloscope screens and the voltage divider
    s1 = new_slide(p, model, 'What you see on the oscilloscope')
    textbox(s1, 0.5, 1.4, 6.6, 5.6, SCOPE, gap=12); picture(s1, scope_figure(), 7.3, 1.1, 5.7, 6.0, top=True)
    set_notes(s1, spoken(SCOPE))
    s2 = new_slide(p, model, 'Why the output falls – a voltage divider')
    textbox(s2, 0.5, 1.4, 6.4, 5.6, DIVIDER, gap=12)
    rows = [('f', 'XC (C = 100 nF)', 'R', 'Who takes most of Vin?')] + [
        (fmt(f), ohm(xc(f)), '10 kΩ', {10: 'C → Vout ≈ Vin', 100: 'C → Vout big', 1000: 'R → Vout small', 10000: 'R → Vout ≈ 0'}[f])
        for f in (10, 100, 1000, 10000)]
    tb = s2.shapes.add_table(len(rows), 4, Inches(7.1), Inches(1.4), Inches(5.9), Inches(2.4)); tb.name = MARK
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tb.table.cell(i, j); cell.text = ''
            add_runs(cell.text_frame.paragraphs[0], f'**{v}**' if i == 0 else v, 13, RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else DARK)
    picture(s2, lp, 7.6, 4.2, 4.9, 2.7)
    set_notes(s2, spoken(DIVIDER))

    # high-pass (old Hanoi "Now do the same thing" slide)
    s = find(p, 'Now do the same thing as before', False) or find(p, 'Swap R and C: a high-pass filter')
    for sh in list(s.shapes):
        if sh.shape_type == 9 or (sh.has_text_frame and sh.text_frame.text.strip() == 'Output signal'):
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and sh.text_frame.text.startswith('Now do the same thing as before'):
            rs = sh.text_frame.paragraphs[0].runs; rs[0].text = 'Swap R and C: a high-pass filter.'
            for r in rs[1:]: r.text = ''
    clear(s); textbox(s, 0.47, 2.4, 5.2, 4.0, SWAP, gap=12)
    swap_pictures(s, [hp])
    set_notes(s, 'Swap R and C: a high-pass filter.\n' + spoken(SWAP))
    swap = s

    # theory slides with circuits: our circuits; the response graphs + "Reading the graph"
    graphs = None
    for s in p.slides:
        t = text_of(s)
        if 'Low pass filter:' in t and any(sh.name == 'Key word callout' for sh in s.shapes):
            swap_pictures(s, [lp, hp])
        elif 'Low pass filter:' in t:
            graphs = s
            clear(s); line = 'Both curves bend at fc (next slide).'
            textbox(s, 7.2, 6.88, 5.6, 0.45, [line], size=14)
            set_notes(s, 'Low-pass: lower frequency gives a higher output. High-pass: higher frequency gives a higher output.\n' + line)
        elif 'You already see how capacitors are used in filters' in t:
            swap_pictures(s, [lp, hp])
    s3 = new_slide(p, model, 'Reading the graph – the −3 dB point')
    textbox(s3, 0.5, 1.3, 6.6, 5.8, READ, size=15, gap=8)
    lp_graph = big_pictures(graphs)[0]                     # the low-pass response picture of the graphs slide
    picture(s3, lp_graph.image.blob, 7.3, 1.2, 5.7, 5.9, top=True)
    set_notes(s3, spoken(READ))
    move_after(p, [s3], graphs)

    # cut-off slides: the Hanoi text box (formula object + 47k + 47k) -> plain text with our circuit
    cut = [s for s in p.slides if any(sh.name.startswith(MARK + ' cutoff') for sh in s.shapes) or
           b'You should also know the formula' in __import__('lxml.etree', fromlist=['x']).tostring(s._element)]
    assert len(cut) == 2, len(cut)
    callout = 'New word: Cut-off frequency (fc). Cut-off frequency (fc) means: where Vout = 0.707 × Vin. Say it with me: Cut-off frequency (fc).'
    for s, kind in zip(cut, ('lp', 'hp')):
        for ac in s._element.findall('.//{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'):
            ac.getparent().remove(ac)
        clear(s)
        lines = CUTOFF_LP if kind == 'lp' else CUTOFF_HP
        textbox(s, 0.47, 1.5, 6.2, 4.4, lines, gap=14).name = MARK + ' cutoff'
        swap_pictures(s, [lp if kind == 'lp' else hp])
        set_notes(s, spoken(lines) + ('\n' + callout if kind == 'lp' else ''))
    # the marked Task 1 goes to the end of the lab, right before Task 2 ("Lab 2 exercise" section)
    task1 = find(p, 'Task 1 (individual) – RC cut-off frequency')
    task2 = find(p, 'Series RLC filter')
    move_before(p, [task1], task2)

    # tables and the hidden answer slides follow Exercise 1 and Exercise 2
    tables = [s for s in p.slides if any(sh.has_table and sh.table.cell(0, 0).text.startswith('Generated frequency') for sh in s.shapes)]
    notice = [s for s in p.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == 'What do you notice?' for sh in s.shapes)]
    assert len(tables) == 2 and len(notice) == 2
    for s, kind in zip(tables, ('lp', 'hp')):
        t = next(sh for sh in s.shapes if sh.has_table).table
        put(t.cell(0, 0), ['Generated frequency at function generator (2 V', 'peak-to-peak', ')'])
        put(t.cell(0, 1), ['Observation (oscilloscope waveform, the frequency and V', 'peak-to-peak',
                           ' of the signal across C)' if kind == 'lp' else ' of the signal across R)'])
        for i, f in enumerate(ROWS, 1):
            put(t.cell(i, 0), [fmt(f)])
        clear(s); line = f'Fill in what you see on CH2 (Vout across {"C" if kind == "lp" else "R"}) at each frequency.'
        textbox(s, 0.69, 1.2, 10, 0.5, [line])
        set_notes(s, line)
    for s, kind in zip(notice, ('lp', 'hp')):
        clear(s); textbox(s, 0.47, 2.1, 6.2, 4.8, NOTICE[kind])
        picture(s, notice_graph(kind), 6.9, 1.45, 6.1, 4.6, top=True)
        set_notes(s, 'What do you notice?\n' + spoken(NOTICE[kind]))
    # order of the RC part: plan, photo, Exercise 1, table, notice, oscilloscope, divider, high-pass, Exercise 2, table,
    # notice, ... response graphs, reading the graph, ..., cut-off (low-pass), Exercise 3, cut-off (high-pass)
    ex = {k: next(sl for sl in p.slides if is_practice(sl) and f'Exercise {k}:' in text_of(sl)) for k in (1, 2, 3)}
    move_after(p, [ex[1], tables[0], notice[0], s1, s2, swap, ex[2], tables[1], notice[1]], photo)
    move_after(p, [ex[3], cut[1]], cut[0])

    p.save(DECK)
    idx = lambda s: list(p.slides).index(s) + 1
    print('Lab 2:', ', '.join(f'{idx(s)} {n}' for s, n in [(plan, 'plan'), (s1, 'scope'), (s2, 'divider'), (swap, 'high-pass'),
          (graphs, 'graphs'), (s3, 'reading -3 dB'), (cut[0], 'cut-off'), (ex[3], 'Exercise 3'), (tables[0], 'table LP'),
          (tables[1], 'table HP'), (task1, 'Task 1'), (task2, 'Task 2')]), '|', len(p.slides), 'slides')


if __name__ == '__main__':
    run()
