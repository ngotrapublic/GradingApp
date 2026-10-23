"""Generic generator: variant cards, lecturer answer key, Canvas CSV scorer, Canvas QTI quiz."""
import os, zipfile, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E78')
thin = Side(style='thin', color='999999'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
YEL = PatternFill('solid', fgColor='FFF2CC')
wrap = Alignment(wrap_text=True, vertical='top')


def build(lab, total, V, Q, card_fn, key_cols, manual, out_dir, guide_extra=''):
    """
    V: {variant: data}
    Q: list of (match_token, question_text, unit, marks, keyfn or None). Q[0] must be the variant question.
    card_fn(code, d) -> list of (text, bold) lines for the printed card.
    key_cols: list of (header, fn(d)) shown on the lecturer answer sheet.
    manual: list of (header, max_marks) columns typed by the lecturer in the score sheet.
    """
    os.makedirs(out_dir, exist_ok=True)
    wb = Workbook()
    ws = wb.active; ws.title = 'Hướng dẫn'
    meas = [q for q in Q[1:] if q[4]]
    rule = ' + '.join(f'{q[0]} {q[3]}đ' for q in meas) + ''.join(f' + {m[0]} {m[1]}đ' for m in manual)
    guide = [
        (f'LAB {lab} – CÁCH DÙNG FILE NÀY', None),
        ('Trước buổi học', f'Đăng file PDF phiếu đề (HCM_Lab{lab}_Phieu_de_LopX.pdf) lên Canvas. Mỗi trang là một nhóm, SV tìm theo nhóm và số thành viên.\nChấm trên điện thoại: dùng trang web trong thư mục GradingApp (file Excel này là cách chấm dự phòng).'),
        ('Trong giờ', 'SV gọi → GV chọn 1 số bất kỳ, nhìn số trên máy đo của SV, so với đáp án (lệch ≤ 10% là được) → ✓ vào cột "Tick" trên tờ đáp án.' + guide_extra),
        ('Cuối giờ', f'SV nộp Canvas quiz "Lab {lab} – Data". Quiz đóng khi hết giờ.'),
        ('Sau giờ', f'1) Canvas → Quiz "Lab {lab} – Data" → Quiz Statistics → Student Analysis → tải CSV.\n2) Mở CSV, Ctrl+A, Ctrl+C → dán vào ô A1 của sheet "CSV Lớp x".\n3) Sheet "Điểm Lớp x": nhập 1 vào cột Tick cho SV có ✓' + (', nhập các cột chấm tay (nền vàng).' if manual else '.') + ' Điểm tự tính.'),
        ('Luật chấm', f'Mỗi số: lệch ≤ 10% so với đáp án → đủ điểm; ≤ 20% → nửa điểm; > 20% → 0. Không có Tick → 0 cho các số đo.\nLab {lab}: {rule} = {total} điểm.'),
        ('', ''),
        ('TẠO QUIZ TRÊN CANVAS', None),
        ('Import', f'Course → Settings → Import Course Content → Content Type: "QTI .zip file" → chọn Canvas_Lab{lab}_QTI.zip → Import. Quiz "Lab {lab} – Data" xuất hiện trong Quizzes (Classic Quiz).'),
        ('Cài đặt nên chỉnh', '• Allowed attempts: 1\n• Bỏ chọn "Let Students See Their Quiz Responses" (quiz 0 điểm; điểm thật tính bằng file này)\n• Available from/Until: giờ bắt đầu → giờ kết thúc buổi học\n• Nên đặt Access code đọc trong lớp để SV không nộp ở nhà'),
        ('Nội dung câu hỏi', '\n'.join(f'Q{i}. {q[1]}' for i, q in enumerate(Q, 1))),
    ]
    for r, (a, b) in enumerate(guide, 1):
        ws.cell(r, 1, a).font = Font(bold=True, size=13 if b is None else 11)
        if b: ws.cell(r, 2, b).alignment = wrap
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110

    classes = sorted({k // 100 for k in V})
    for cls in classes:
        vs = {k: v for k, v in V.items() if k // 100 == cls}
        # cards
        if card_fn is None:
            pass
        else:
          pc = wb.create_sheet(f'Phiếu đề Lớp {cls} (in)')
          pc.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
          pc.column_dimensions['A'].width = 46; pc.column_dimensions['B'].width = 2; pc.column_dimensions['C'].width = 46
          n_lines = max(len(card_fn(k, d)) for k, d in vs.items()) + 1
          for i, (code, d) in enumerate(sorted(vs.items())):
              col = 1 if i % 2 == 0 else 3; r0 = (i // 2) * n_lines + 1
              lines = card_fn(code, d) + [('Name: ______________________________', False)]
              for j, (t, b) in enumerate(lines):
                  c = pc.cell(r0 + j, col, t); c.font = Font(bold=b, size=12 if j == 0 else 11)
                  c.border = Border(left=thin, right=thin, top=thin if j == 0 else None, bottom=thin if j == len(lines) - 1 else None)
                  if j == 0: c.fill = YEL
        # lecturer answer key
        ak = wb.create_sheet(f'Đáp án Lớp {cls} (in)')
        ak.append([f'Lab {lab} – Đáp án Lớp {cls} – chấp nhận ±10% (đủ điểm), ±20% (nửa điểm)'])
        ak['A1'].font = Font(bold=True, size=13)
        hdr = ['Mã đề'] + [h for h, _ in key_cols] + ['Tên SV', 'Tick ✓'] + [m[0] for m in manual]
        ak.append(hdr)
        for c in ak[2]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', wrap_text=True)
        for code, d in sorted(vs.items()):
            ak.append([code] + [fn(d) for _, fn in key_cols] + [''] * (2 + len(manual)))
        for row in ak.iter_rows(min_row=2, max_row=ak.max_row):
            for c in row: c.border = B
        for r in range(3, ak.max_row + 1): ak.row_dimensions[r].height = 24
        for i in range(1, len(hdr) + 1):
            ak.column_dimensions[ak.cell(2, i).column_letter].width = 22 if hdr[i - 1] == 'Tên SV' else 11
        ak.page_setup.orientation = 'landscape'; ak.page_setup.fitToWidth = 1
        # hidden key
        key = wb.create_sheet(f'Key{cls}')
        key.append(['Variant'] + [q[0] for q in meas])
        for code, d in sorted(vs.items()):
            key.append([code] + [q[4](d) for q in meas])
        key.sheet_state = 'hidden'
        # CSV + score
        csv = wb.create_sheet(f'CSV Lớp {cls}')
        csv['A1'] = 'DÁN TOÀN BỘ FILE CSV "Student Analysis" CỦA CANVAS VÀO Ô A1 (ghi đè dòng này)'
        csv['A1'].font = Font(bold=True, color='C00000')
        sc = wb.create_sheet(f'Điểm Lớp {cls}')
        n = len(meas)
        cols = (['Tên SV', 'Mã đề', 'Tick (1 = có ✓)'] + [f'{q[0]} ({q[2]})' for q in meas]
                + [f'Điểm {q[0]}' for q in meas] + [f'{m[0]} (0–{m[1]})' for m in manual] + [f'TỔNG /{total}', 'Ghi chú'])
        sc.append(cols)
        for c in sc[1]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', wrap_text=True)
        CS = f"'CSV Lớp {cls}'"; KS = f'Key{cls}'; last_key_col = chr(65 + n)
        for r in range(2, 62):
            get = lambda pat: f'IFERROR(VALUE(INDEX({CS}!$A$1:$ZZ$200,ROW(),MATCH("*{pat}*",{CS}!$1:$1,0))),"")'
            sc.cell(r, 1, f'=IF({CS}!A{r}="","",{CS}!A{r})')
            sc.cell(r, 2, '=' + get(Q[0][0]))
            sc.cell(r, 3).fill = YEL
            for i, q in enumerate(meas):
                vc = 4 + i; scol = 4 + n + i
                sc.cell(r, vc, '=' + get(q[0]))
                vref = sc.cell(r, vc).coordinate
                k = f'VLOOKUP($B{r},{KS}!$A:${last_key_col},{i + 2},0)'
                sc.cell(r, scol, f'=IFERROR(IF($C{r}<>1,0,IF(ABS({vref}/{k}-1)<=0.1,{q[3]},IF(ABS({vref}/{k}-1)<=0.2,{q[3] / 2},0))),0)')
            for j in range(len(manual)):
                sc.cell(r, 4 + 2 * n + j).fill = YEL
            first = sc.cell(r, 4 + n).coordinate; lastc = sc.cell(r, 3 + 2 * n + len(manual)).coordinate
            tc = 4 + 2 * n + len(manual)
            sc.cell(r, tc, f'=IF(A{r}="","",SUM({first}:{lastc}))').font = Font(bold=True)
            sc.cell(r, tc + 1, f'=IF(A{r}="","",IF(ISERROR(VLOOKUP($B{r},{KS}!$A:$A,1,0)),"Sai mã đề / chưa nhập",IF(C{r}<>1,"Chưa có tick","")))')
        for i in range(1, len(cols) + 1):
            sc.column_dimensions[sc.cell(1, i).column_letter].width = 24 if i == 1 else 11
        sc.freeze_panes = 'D2'
    wb.save(f'{out_dir}/HCM_Lab{lab}_Phieu_de_va_Cham_diem.xlsx')
    qti(lab, Q, out_dir)


def qti(lab, Q, out_dir):
    def item(i, title, text):
        return f'''<item ident="lab{lab}_q{i}" title="{html.escape(title)}">
<itemmetadata><qtimetadata>
<qtimetadatafield><fieldlabel>question_type</fieldlabel><fieldentry>numerical_question</fieldentry></qtimetadatafield>
<qtimetadatafield><fieldlabel>points_possible</fieldlabel><fieldentry>0.0</fieldentry></qtimetadatafield>
</qtimetadata></itemmetadata>
<presentation><material><mattext texttype="text/html">{html.escape('<p>' + html.escape(text) + '</p>')}</mattext></material>
<response_str ident="response1" rcardinality="Single"><render_fib fibtype="Decimal"><response_label ident="answer1"/></render_fib></response_str>
</presentation>
<resprocessing><outcomes><decvar maxvalue="100" minvalue="0" varname="SCORE" vartype="Decimal"/></outcomes>
<respcondition continue="No"><conditionvar><vargte respident="response1">-1000000000</vargte><varlte respident="response1">1000000000</varlte></conditionvar>
<setvar action="Set" varname="SCORE">100</setvar></respcondition></resprocessing>
</item>'''
    ident = f'hcm_lab{lab}'
    items = '\n'.join(item(i, q[0], q[1]) for i, q in enumerate(Q, 1))
    assess = f'''<?xml version="1.0" encoding="UTF-8"?>
<questestinterop xmlns="http://www.imsglobal.org/xsd/ims_qtiasiv1p2" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:schemaLocation="http://www.imsglobal.org/xsd/ims_qtiasiv1p2 http://www.imsglobal.org/xsd/ims_qtiasiv1p2p1.xsd">
<assessment ident="{ident}" title="Lab {lab} – Data">
<qtimetadata><qtimetadatafield><fieldlabel>cc_maxattempts</fieldlabel><fieldentry>1</fieldentry></qtimetadatafield></qtimetadata>
<section ident="root_section">
{items}
</section></assessment></questestinterop>'''
    manifest = f'''<?xml version="1.0" encoding="UTF-8"?>
<manifest identifier="{ident}_manifest" xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<metadata><schema>IMS Content</schema><schemaversion>1.1.3</schemaversion></metadata>
<organizations/>
<resources><resource identifier="{ident}" type="imsqti_xmlv1p2"><file href="{ident}/{ident}.xml"/></resource></resources>
</manifest>'''
    with zipfile.ZipFile(f'{out_dir}/Canvas_Lab{lab}_QTI.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('imsmanifest.xml', manifest)
        z.writestr(f'{ident}/{ident}.xml', assess)
