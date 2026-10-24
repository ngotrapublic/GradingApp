"""Lab 2, Exercise 6 (band-pass RLC: L = 1 mH, C = 100 nF, R = 100 Ohm, f0 = 15.9 kHz), the same way as the RC part:
a table to fill in at 7 frequencies + an empty graph to draw, then "What do you notice?" with the theory curve and the
7 points (it also answers the question of Exercise 6). Placed right after Exercise 6.
Run after practice_slides.py (and rc_section.py), before subscripts.py. Running again replaces the two slides."""
import copy, io, math, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/HCM Electronic Circuit Lab 2.pptx'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/'
L, C, R = 1e-3, 100e-9, 100.0
F0 = 1 / (2 * math.pi * math.sqrt(L * C))
ROWS = [2e3, 5e3, 10e3, 16e3, 25e3, 50e3, 100e3]
MARK = 'RLC Table'
PRACTICE = 'Practice Exercise'
MODEL = 'Devices and components for this session'
BAR = 'Series RLC – band-pass'
DARK, WHITE = RGBColor(0x22, 0x22, 0x22), RGBColor(0xFF, 0xFF, 0xFF)
khz = lambda f: f'{f / 1e3:g} kHz'


def ratio(f):
    w = 2 * np.pi * np.asarray(f)
    return R / np.abs(R + 1j * (w * L - 1 / (w * C)))


TABLE_LINE = 'Write Vin and Vout at each frequency. Then calculate Vout / Vin and draw the graph in your notebook.'
NOTICE = ['At **low frequency** (2 kHz), Vout is small. At **high frequency** (100 kHz), Vout is small again.',
          f'The **peak** is near **f0 ≈ {F0 / 1e3:.1f} kHz**. Your peak can be a few % away, because L and C are not exact.',
          'Only the frequencies **in the middle** pass: this is a **band-pass** filter.',
          'Your highest Vout / Vin is a little **below 1**: the coil wire has a small resistance RL, and part of the voltage is lost there.',
          'Why does only the middle pass? The next slides explain it.']


def axes(ax):
    ax.set_xscale('log'); ax.set_xlim(1.5e3, 140e3); ax.set_ylim(0, 1.1)
    ax.set_xticks(ROWS); ax.set_xticklabels([f'{f / 1e3:g}' for f in ROWS]); ax.minorticks_off()
    ax.set_xlabel('frequency (kHz, log scale)'); ax.set_ylabel('$V_{out}$ / $V_{in}$')
    ax.grid(True, color='#ddd')
    for s in ('top', 'right'): ax.spines[s].set_visible(False)


def empty_graph():
    fig, ax = plt.subplots(figsize=(5.4, 4.2), dpi=170)
    axes(ax); ax.set_title('Draw your graph', fontsize=12)
    fig.tight_layout(); path = FIG + 'p2_rlc_empty.png'; fig.savefig(path); plt.close(fig)
    return path


def theory_graph():
    f = np.logspace(math.log10(1.5e3), math.log10(140e3), 400)
    fig, ax = plt.subplots(figsize=(6.0, 4.4), dpi=170)
    axes(ax)
    ax.plot(f, ratio(f), color='#3f7a2a', lw=2.5)
    ax.axvline(F0, color='#c0392b', ls=':', lw=1.5)
    ax.text(F0 * 1.08, 0.05, f'$f_0$ ≈ {F0 / 1e3:.1f} kHz', color='#c0392b', fontsize=11)
    for x in ROWS:
        ax.plot(x, ratio(x), 'o', color='#C2410C', ms=7)
        ax.annotate(f'{ratio(x):.2f}', (x, ratio(x)), textcoords='offset points', xytext=(6, 6), fontsize=9, color='#444')
    ax.set_title('Band-pass: L = 1 mH, C = 100 nF, R = 100 Ω (theory)', fontsize=11)
    fig.tight_layout(); path = FIG + 'p2_rlc_graph.png'; fig.savefig(path); plt.close(fig)
    return path


def add_runs(para, text, size=18, color=DARK):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if part:
            r = para.add_run(); r.text = part.strip('*'); r.font.size = Pt(size); r.font.name = 'Rockwell'
            r.font.bold = part.startswith('**'); r.font.color.rgb = color


def textbox(s, x, y, w, h, lines, size=18, gap=10):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tb.name = MARK
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); add_runs(para, line, size); para.space_after = Pt(gap)


def picture(s, path, x, y, w, h):
    iw, ih = Image.open(path).size; k = min(w / iw, h / ih)
    s.shapes.add_picture(path, Inches(x + (w - iw * k) / 2), Inches(y), Inches(iw * k), Inches(ih * k)).name = MARK


def new_slide(p, model):
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element)); bar = s.shapes[-1]
            bar.text_frame.paragraphs[0].runs[0].text = BAR
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
            bar.name = MARK + ' bar'
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    return s


def run():
    p = Presentation(DECK); lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if any(sh.name == MARK + ' bar' for sh in sl.shapes):
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    model = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == MODEL for sh in s.shapes))
    ex6 = next(s for s in p.slides if any(sh.name == PRACTICE for sh in s.shapes) and
               any(sh.has_text_frame and sh.text_frame.text.startswith('Exercise 6:') for sh in s.shapes))

    t = new_slide(p, model)
    textbox(t, 0.5, 1.2, 12.3, 0.5, [TABLE_LINE])
    rows = [('f', 'Vin (CH1, Vpp)', 'Vout (CH2, Vpp)', 'Vout / Vin')] + [(khz(f), '', '', '') for f in ROWS]
    tb = t.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.9), Inches(6.6), Inches(4.9)); tb.name = MARK
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tb.table.cell(i, j); cell.text = ''
            add_runs(cell.text_frame.paragraphs[0], f'**{v}**' if i == 0 or j == 0 else v, 15, WHITE if i == 0 else DARK)
    picture(t, empty_graph(), 7.3, 1.8, 5.8, 5.1)
    t.notes_slide.notes_text_frame.text = TABLE_LINE + '\nThe frequencies are 2, 5, 10, 16, 25, 50 and 100 kilohertz.'

    n = new_slide(p, model)
    textbox(n, 0.5, 1.2, 6.3, 0.5, ['What do you notice?'])
    textbox(n, 0.5, 1.9, 6.3, 5.0, NOTICE, gap=12)
    picture(n, theory_graph(), 6.9, 1.3, 6.1, 5.6)
    n.notes_slide.notes_text_frame.text = 'What do you notice?\n' + '\n'.join(x.replace('**', '') for x in NOTICE)

    for s in (n, t):                                       # Exercise 6, table, notice
        sid = lst[list(p.slides).index(s)]; lst.remove(sid)
        lst.insert(list(p.slides).index(ex6) + 1, sid)
    p.save(DECK)
    idx = lambda s: list(p.slides).index(s) + 1
    print(f'Lab 2: Exercise 6 at {idx(ex6)}, table at {idx(t)}, notice at {idx(n)} | {len(p.slides)} slides |',
          'theory ratios', [f'{ratio(f):.2f}' for f in ROWS])


if __name__ == '__main__':
    run()
