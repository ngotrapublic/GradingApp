"""Labs 2-4, "Review – words from Lab N": the answers appear on the same slide, one row per click (lecturer, 24/10/2026),
instead of a separate "(answers)" slide after all the students have answered.
The review table becomes the answers table (Word | Meaning | Example) with Meaning and Example empty; on top of every
row two text boxes (meaning, example) appear together on a click, row after row. The "(answers)" slide is removed.
A table row cannot be animated in PowerPoint, so the answers are text boxes placed exactly over the cells.
Running again does nothing once the "(answers)" slide is gone."""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
MARK = 'Reveal answer'


def title(s):
    return ' '.join(sh.text_frame.text.strip() for sh in s.shapes if sh.name == 'Title Bar')


def table_shape(s):
    return next(sh for sh in s.shapes if sh.has_table)


def timing_xml(groups):
    """Appear on click: one click per group of shape ids (the first id starts on the click, the others with it)."""
    n = [2]

    def nid():
        n[0] += 1
        return n[0]
    clicks = []
    for ids in groups:
        effects = []
        for k, sid in enumerate(ids):
            a, b = nid(), nid()
            effects.append(
                f'<p:par><p:cTn id="{a}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" '
                f'nodeType="{"clickEffect" if k == 0 else "withEffect"}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')
        o, i = nid(), nid()
        clicks.append(f'<p:par><p:cTn id="{o}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                      f'<p:par><p:cTn id="{i}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                      + ''.join(effects) + '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    bld = ''.join(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>' for ids in groups for sid in ids)
    return (f'<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            + ''.join(clicks) +
            '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
            f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{bld}</p:bldLst></p:timing>')


def set_timing(slide, xml):
    sld = slide._element
    for old in sld.findall(f'{{{NS}}}timing'):
        sld.remove(old)
    el = etree.fromstring(xml)
    anchor = None                                          # schema order: cSld, clrMapOvr, transition, timing, extLst
    for tag in ('transition', 'clrMapOvr', 'cSld'):
        anchor = sld.find(f'{{{NS}}}{tag}')
        if anchor is not None:
            break
    anchor.addnext(el)


def merge(p, review, answers):
    old, ans = table_shape(review), table_shape(answers)
    rows = [[c.text for c in r.cells] for r in ans.table.rows]
    # the answers table, with Meaning and Example empty, takes the place of the review table
    frame = copy.deepcopy(ans._element)
    old._element.addprevious(frame); old._element.getparent().remove(old._element)
    tb = table_shape(review)
    for r in list(tb.table.rows)[1:]:
        for c in list(r.cells)[1:]:
            for para in c.text_frame.paragraphs:
                for run in para.runs:
                    run.text = ''
    size = next((r.font.size for row in ans.table.rows for c in list(row.cells)[1:]
                 for para in c.text_frame.paragraphs for r in para.runs if r.font.size), Pt(14))
    x0, y0 = tb.left, tb.top
    widths = [c.width for c in tb.table.columns]; heights = [r.height for r in tb.table.rows]
    groups = []
    for i, row in enumerate(rows[1:], 1):
        y = y0 + sum(heights[:i]); ids = []
        for j in (1, 2):
            x = x0 + sum(widths[:j])
            box = review.shapes.add_textbox(x, y, widths[j], heights[i]); box.name = MARK
            tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Emu(91440); tf.margin_top = tf.margin_bottom = Emu(18288)
            r = tf.paragraphs[0].add_run(); r.text = row[j]; r.font.size = size; r.font.name = 'Arial'
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            ids.append(box.shape_id)
        groups.append(ids)
    set_timing(review, timing_xml(groups))
    review._element.attrib.pop('show', None)
    notes = review.notes_slide.notes_text_frame
    if 'click' not in notes.text:
        notes.text = notes.text.rstrip() + '\nAfter each answer, I click: the meaning and an example of that word appear.'
    # remove the "(answers)" slide
    lst = p.slides._sldIdLst
    sid = lst[list(p.slides).index(answers)]
    p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    return len(groups)


def run():
    for n in (2, 3, 4):
        path = f'{HCM}/Lab{n}/HCM Electronic Circuit Lab {n}.pptx'
        p = Presentation(path); done = []
        while True:
            slides = list(p.slides)
            pair = next(((s, a) for s in slides for a in slides
                         if title(s).startswith('Review – words from') and title(a) == title(s) + ' (answers)'), None)
            if pair is None:
                break
            done.append((title(pair[0]), merge(p, *pair)))
        if done:
            p.save(path)
        print(f'Lab {n}:', done or 'nothing to do', '|', len(p.slides), 'slides')


if __name__ == '__main__':
    run()
