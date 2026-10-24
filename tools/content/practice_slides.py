"""Labs 2-4: short hands-on practice slides in the style of the Lab 1 deck (edited by hand on 24/10/2026):
"Exercise N: <what to do> (X minutes)" in Rockwell 16, numbered steps on the left, a figure or photo on the right, one
question at the end. Practice only, not marked. Demo values are not on any card and use parts the boxes have spare
(RC 10 kOhm + 104, RLC 1 mH + 104 + 100 Ohm; Lab 3 and Lab 4 use the circuits of the lab itself).
All exercises of a lab form one block (a "Practice exercises" slide + the exercises), right before the slide in BLOCK
(the first marked task). A figure is a file in fig/ or the largest picture of the slide that contains a text.
Running again replaces the practice slides. Order: align_web.py, rlc_slides.py, this, rc_section.py, subscripts.py."""
import copy, io, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
FIG = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/'
MARK = 'Practice Exercise'
MODEL = 'Devices and components for this session'
DARK = RGBColor(0x22, 0x22, 0x22)

# one block of exercises per lab, at the end of the theory, before the first marked task (lecturer, 24/10/2026)
BLOCK = {2: ('Task 1 (individual) – RC cut-off frequency', 'Then: Task 1 (individual) and Task 2 (team) – these are marked.'),
         3: ('From an oscilloscope screen', 'Then: the Lab 3 exercise – this is marked.'),
         4: ('Final circuit (version 1)', 'Then: the final circuit – this is marked.')}

LAB = {
2: [
    dict(bar='1st order RC circuits', fig='p2_lp.png',
         head='Exercise 1: Build an RC low-pass filter and change the frequency (10 minutes)',
         steps=['Take R = 10 kΩ and C = 100 nF (code 104) from the parts box.',
                'Measure R with the multimeter. Write it down.',
                'Build the circuit. CH1 = Vin, CH2 = Vout (across C).',
                'Function generator: sine, 2 Vpp, High-Z.',
                'Set 10 Hz, 100 Hz, 200 Hz, 1 kHz, 10 kHz. Write what you see in the table (next slide).'],
         ask='What happens to Vout when the frequency goes up?'),
    dict(bar='1st order RC circuits', fig='p2_hp.png',
         head='Exercise 2: Swap R and C – a high-pass filter (5 minutes)',
         steps=['Swap R and C: the capacitor is now in series, the resistor goes to ground.',
                'CH2 = Vout across R.',
                'Use the same frequencies: 10 Hz, 100 Hz, 200 Hz, 1 kHz, 10 kHz.',
                'Fill in the second table (next slide). Compare it with the first one.'],
         ask='Which frequencies pass now? Which are blocked?'),
    dict(bar='1st order RC circuits', fig='p2_lp.png',
         head='Exercise 3: Find the cut-off frequency of your low-pass filter (5 minutes)',
         steps=['Build the low-pass filter of Exercise 1 again.',
                'Calculate fc = 1 / (2πRC) with your measured R.',
                'Read Vin on CH1 (Measure → Pk-Pk).',
                'Change the frequency until Vout = 0.707 × Vin.',
                'This frequency is your measured fc.'],
         ask='How far is your fc from your calculation? Why are they different?'),
    dict(bar='Capacitors', fig='p2_capcode.png',
         head='Exercise 4: Read and measure capacitors (5 minutes)',
         steps=['Each member takes one capacitor from the parts box: 103, 223 or 473.',
                'Read the code. Write the value in nF.',
                'Set the multimeter to capacitance mode.',
                'Measure the capacitor. Write it down.',
                'Put the capacitor back in its bag.'],
         ask='How far is your reading from the code? Is it within 5%?'),
    dict(bar='Inductors', fig=('You will see them in many circuits',),
         head='Exercise 5: Measure the resistance of the 1 mH inductor (2 minutes)',
         steps=['Take the 1 mH inductor. It looks like a resistor: check the bag label.',
                'Set the multimeter to resistance (Ω).',
                'Measure the inductor. Write down RL.',
                'Put it back in its bag. You need it again in Task 2.'],
         ask='An ideal inductor has 0 Ω. Why does yours not?'),
    dict(bar='Series RLC – band-pass', fig='p2_rlc.png',
         head='Exercise 6: Find the peak of a band-pass RLC filter (10 minutes)',
         steps=['Take L = 1 mH, C = 100 nF (104) and R = 100 Ω.',
                'Calculate f0 = 1 / (2π√(LC)).',
                'Build the circuit. R at the ground side. CH1 = Vin, CH2 = Vout (across R).',
                'Function generator: sine, 2 Vpp, High-Z. Start at 5 kHz.',
                'Raise the frequency slowly to 30 kHz. Find the highest Vout / Vin.'],
         ask='Is your peak at f0? Why is the highest Vout / Vin below 1?'),
],
3: [
    dict(bar='Overview of op amp', fig=('Vcc / Vee = the positive',),
         head='Exercise 1: Find the pins of the LM358 (2 minutes)',
         steps=['Find the notch or the dot on the LM358. Pin 1 is next to it.',
                'Count the pins anticlockwise, seen from the top: 1 to 8.',
                'Plug the IC across the middle gap of the breadboard.',
                'Show your teammates pin 8 (Vcc), pin 4 (Vee), pins 2 and 3 (inputs) and pin 1 (output).'],
         ask='Why must the IC sit across the middle gap?'),
    dict(bar='Constructing first amplifier', fig=('Connect 6V to Vs+ pin',),
         head='Exercise 2: Power the LM358 safely (5 minutes)',
         steps=['Switch on the trainer. Set +6 V and −6 V. Check both with the multimeter.',
                'Switch off. Connect GND to the ground rail, +6 V to pin 8, −6 V to pin 4.',
                'Ask a teammate to check your wiring.',
                'Switch on. Measure pin 8 and pin 4 against GND.',
                'Touch the IC carefully: it must not get hot.'],
         ask='What do you measure on pin 4? Why is it negative?'),
    dict(bar='Constructing first amplifier', fig='p3_follower.png',
         head='Exercise 3: A voltage follower (5 minutes)',
         steps=['Connect pin 1 (output) to pin 2 (− input).',
                'Function generator to pin 3 (+ input): sine, 1 Vpp, 1 kHz.',
                'CH2 = Vin (pin 3), CH1 = Vout (pin 1). Measure Pk-Pk of both.',
                'Remove the wire from pin 1 to pin 2 when you finish.'],
         ask='What is the gain? Is the output upside down?'),
    dict(bar='Measuring gain and phase', fig=('Your circuit should look something like this',),
         head='Exercise 4: Build the first amplifier and measure its gain (10 minutes)',
         steps=['Build the inverting amplifier: R1 = 2.2 kΩ, R2 = 30 kΩ (3 × 10 kΩ), C1 = 1 µF.',
                'Add the decoupling capacitors: 470 µF on each rail, 100 nF at pins 8 and 4.',
                'Function generator: sine, 500 Hz, 50 mVpp (easier to read than 10 mVpp).',
                'CH2 = Vin, CH1 = Vout. Measure both (Pk-Pk). Gain = Vout / Vin.',
                'Use the cursors to read Δt between the waves. Phase = 360° × Δt × f.'],
         ask='Does your gain match R2 / R1? Is the phase 180°?'),
    dict(bar='LM358 limits you will meet', fig='p3_clip.png',
         head='Exercise 5: Make the output clip (3 minutes)',
         steps=['Keep your amplifier from Exercise 4.',
                'Raise the input amplitude step by step: 0.2, 0.4, 0.6, 0.8 Vpp.',
                'Watch CH1: find when the tops of the wave become flat.',
                'Measure the highest and the lowest Vout (Measure → Max, Min).',
                'Set the input back to 50 mVpp.'],
         ask='Why can the output not reach +6 V and −6 V?'),
],
4: [
    dict(bar='Testing 8-Ω speaker', fig=('Using normal 1X cable',),
         head='Exercise 1: Test the speaker with the function generator (5 minutes)',
         steps=['Measure the resistance of the speaker with the multimeter.',
                'Connect the function generator (1X cable) to the speaker with jumper wires.',
                'Function generator: sine, 0.5 Vpp, 500 Hz. Listen.',
                'Change the amplitude, then the frequency (200 Hz – 5 kHz). Listen.',
                'CH1 across the speaker: measure Pk-Pk.'],
         ask='The generator is set to 0.5 Vpp. Why is the voltage at the speaker much smaller?'),
    dict(bar='Testing 8-Ω speaker', fig=('The schematic circuit is something like this',),
         head='Exercise 2: A coupling capacitor in front of the speaker (5 minutes)',
         steps=['Function generator: sine, 1 Vpp, 500 Hz.',
                'Put a 100 nF capacitor in series with the speaker. Measure Pk-Pk at the speaker.',
                'Replace it with 1 µF, then with 470 µF (+ lead to the generator side).',
                'Listen and measure each time.'],
         ask='Which capacitor lets the most signal through? Why?'),
    dict(bar='Testing only electret', fig=('Watch out for the correct pin',),
         head='Exercise 3: Find the ground pin of the microphone (2 minutes)',
         steps=['Set the multimeter to continuity mode.',
                'Touch the metal case of the microphone with one probe.',
                'Touch the male end of each jumper with the other probe.',
                'The one that beeps is the ground pin. It must be on the black jumper.'],
         ask='What happens if the microphone is connected the wrong way round?'),
    dict(bar='Testing only electret', fig=('Electret test circuit on the PB-503A',),
         head='Exercise 4: See your voice on the oscilloscope (10 minutes)',
         steps=['Build the electret test circuit: bias resistor from +6 V, ground pin to GND.',
                'Connect CH1 after the 1 µF coupling capacitor.',
                'Set CH1 to 100 mV/div and the time base to 1 ms/div.',
                'Talk or hum into the microphone through a rolled piece of paper.',
                'Press Run/Stop to freeze the wave. Measure Pk-Pk.'],
         ask='How big is your signal? Is it big enough for the speaker?'),
    dict(bar='Constructing second amplifier', fig=('Measurement at oscilloscope', -1),
         head='Exercise 5: Test both stages with the generator first (10 minutes)',
         steps=['Do not connect the microphone yet.',
                'Function generator to the input of stage 1: sine, 10 mVpp, 1 kHz.',
                'CH1 = output of stage 1, CH2 = output of stage 2.',
                'Both outputs must be clean sine waves: no clipping, no noise.',
                'Only then connect the microphone and the speaker.'],
         ask='Why do we test with the generator before we use the microphone?'),
],
}


def text_of(s):
    out = []
    for sh in s.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
    return '\n'.join(out)


def find(p, snippet):
    hits = [s for s in p.slides if snippet in text_of(s) and not is_practice(s)]
    if len(hits) > 1:                                  # the same words in a title and in a review slide: take the title
        hits = [s for s in hits if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == snippet for sh in s.shapes)]
    assert len(hits) == 1, (snippet, len(hits))
    return hits[0]


def is_practice(s):
    return any(sh.name == MARK for sh in s.shapes)


def picture(p, spec):
    """('text',) = largest picture on the slide with that text; ('text', -1) = on the slide before it."""
    s = find(p, spec[0])
    if len(spec) > 1:
        slides = list(p.slides); s = slides[slides.index(s) + spec[1]]
    pics = [sh for sh in s.shapes if sh.shape_type == 13 and not sh.name.startswith('Logo') and sh.width > Inches(2.2)]
    best = max(pics, key=lambda sh: sh.width * sh.height)
    return io.BytesIO(best.image.blob), best.image.size


def textbox(s, x, y, w, h, name=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    return tf


def run_(para, text, size=16, color=DARK):
    r = para.add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = 'Rockwell'; r.font.color.rgb = color
    return r


def numbered(para):
    pPr = para._p.get_or_add_pPr(); pPr.set('marL', '457200'); pPr.set('indent', '-317500')
    for tag in ('a:buFont', 'a:buAutoNum'):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    etree.SubElement(pPr, qn('a:buFont')).set('typeface', '+mj-lt')
    etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')


def add(p, model, ex):
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element)); bar = s.shapes[-1]
            bar.text_frame.paragraphs[0].runs[0].text = ex['bar']
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    tf = textbox(s, 0.6, 1.45, 7.3, 0.45, MARK); run_(tf.paragraphs[0], ex['head'])
    tf = textbox(s, 0.65, 2.15, 7.0, 3.4)
    for i, t in enumerate(ex['steps']):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); numbered(para); run_(para, t)
        para.space_after = Pt(6)
    tf = textbox(s, 0.65, 5.75, 7.0, 0.8); run_(tf.paragraphs[0], ex['ask'])
    # figure on the right, fitted into the free area, top-aligned like the Lab 1 slides
    if isinstance(ex['fig'], str):
        from PIL import Image
        src = FIG + ex['fig']; size = Image.open(src).size
    else:
        src, size = picture(p, ex['fig'])
    bx, by, bw, bh = 8.0, 1.2, 5.0, 5.8
    k = min(bw / size[0], bh / size[1]); w, h = size[0] * k, size[1] * k
    s.shapes.add_picture(src, Inches(bx + (bw - w) / 2), Inches(by + (0 if h > 3.5 else (bh - h) / 3)), Inches(w), Inches(h))
    # notes = exactly what is on the slide (lecturer, 24/10/2026)
    s.notes_slide.notes_text_frame.text = '\n'.join([ex['head']] + [f'{i}. {t}' for i, t in enumerate(ex['steps'], 1)] + [ex['ask']])
    return s


def header(p, model, n, exercises):
    """First slide of the block: what to do, in which order, how long."""
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            s.shapes._spTree.append(copy.deepcopy(sh._element)); bar = s.shapes[-1]
            bar.text_frame.paragraphs[0].runs[0].text = 'Practice exercises'
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
        elif sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    mins = [int(re.search(r'\((\d+) minutes?\)', ex['head']).group(1)) for ex in exercises]
    names = [re.sub(r'^Exercise \d+: | \(\d+ minutes?\)$', '', ex['head']) for ex in exercises]
    first = f'Now do all the exercises, one after the other (about {sum(mins)} minutes). They are not marked.'
    tf = textbox(s, 0.6, 1.45, 12.0, 0.5, MARK); run_(tf.paragraphs[0], first, 18)
    tf = textbox(s, 0.65, 2.2, 12.0, 4.0)
    for i, (name, m) in enumerate(zip(names, mins)):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); numbered(para); run_(para, f'{name} ({m} min)', 18)
        para.space_after = Pt(8)
    last = BLOCK[n][1]
    tf = textbox(s, 0.65, 6.2, 12.0, 0.6); run_(tf.paragraphs[0], last, 18)
    s.notes_slide.notes_text_frame.text = '\n'.join([first] + [f'{i}. {t}' for i, t in enumerate(names, 1)] + [last])
    return s


def run():
    for n, exercises in LAB.items():
        path = f'{HCM}/Lab{n}/HCM Electronic Circuit Lab {n}.pptx'
        p = Presentation(path); lst = p.slides._sldIdLst
        for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
            if is_practice(sl):
                p.part.drop_rel(sid.rId); lst.remove(sid)
        p.part.rename_slide_parts([x.rId for x in lst])
        model = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == MODEL for sh in s.shapes))
        new = [header(p, model, n, exercises)] + [add(p, model, ex) for ex in exercises]
        ref = find(p, BLOCK[n][0])                         # the block goes right before this slide
        sids = [lst[list(p.slides).index(s)] for s in new]
        for sid in sids: lst.remove(sid)
        pos = list(p.slides).index(ref)
        for k, sid in enumerate(sids): lst.insert(pos + k, sid)
        p.save(path)
        where = [i + 1 for i, s in enumerate(p.slides) if is_practice(s)]
        print(f'Lab {n}: {len(p.slides)} slides; practice block at {where[0]}-{where[-1]}')


if __name__ == '__main__':
    run()
