"""Lab 3, Exercise 4 (lecturer, 24/10/2026: Lab 3 was too short): the gain of the Exercise 2 amplifier (R1 = 2.2 kOhm,
R2 = 30 kOhm, C1 = 1 uF, LM358) at 9 frequencies, the same way as the tables of Lab 2: a table + an empty graph, then
"What do you notice?" with the theory curve (it also answers the question of Exercise 4). Placed right after Exercise 4.
"Use low (audio) frequencies" after it becomes the explanation of the high end.
Theory: C1 and R1 are an RC high-pass (fL = 1/(2 pi R1 C1) = 72 Hz); the LM358 (gain-bandwidth about 1 MHz) gives
fH = 1 MHz / (1 + R2/R1) = 68 kHz. Run after practice_slides.py and lab3_fixes.py, before subscripts.py.
Running again replaces the two slides."""
import math, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from rlc_table import add_runs, textbox, picture, WHITE
import rlc_table

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab3/HCM Electronic Circuit Lab 3.pptx'
FIG = rlc_table.FIG
R1, R2, C1, GBW = 2.2e3, 30e3, 1e-6, 1e6
G, FL, FH = R2 / R1, 1 / (2 * math.pi * R1 * C1), GBW / (1 + R2 / R1)
ROWS = [20, 50, 100, 500, 1e3, 10e3, 50e3, 100e3, 200e3]
MARK = rlc_table.MARK
MODEL = 'Devices and components for this session'
rlc_table.BAR = 'Gain and frequency'
LOW = 'Use low (audio) frequencies'
hz = lambda f: f'{f / 1e3:g} kHz' if f >= 1e3 else f'{f:g} Hz'
TABLE_LINE = 'Vin = 50 mVpp. At each frequency, write Vin and Vout. Calculate the gain and draw the graph in your notebook.'
NOTICE = [f'In the **middle** (about 200 Hz – 20 kHz), the gain is about **{G:.1f}** = R2 / R1.',
          f'At **low frequency** (20 Hz, 50 Hz), the gain is small. C1 and R1 are an RC **high-pass filter** (Lab 2): '
          f'fc = 1 / (2π × R1 × C1) ≈ {FL:.0f} Hz.',
          'At **high frequency** (100 kHz, 200 kHz), the gain is small again. The next slide explains why.',
          'So only the **middle frequencies** get the full gain.']
LOW_TEXT = ['Your graph shows it: at **high frequency**, the gain goes down. The LM358 is **not fast** enough.',
            'With a **bigger gain**, the gain goes down at an even lower frequency.',
            'In this lab, use **100 Hz – 2 kHz**.']


def gain(f):
    f = np.asarray(f, float)
    return G * np.abs((1j * f / FL) / (1 + 1j * f / FL) / (1 + 1j * f / FH))


def axes(ax):
    ax.set_xscale('log'); ax.set_xlim(14, 300e3); ax.set_ylim(0, 16)
    ax.set_xticks(ROWS); ax.set_xticklabels([hz(f).replace(' ', '\n') for f in ROWS]); ax.minorticks_off()
    ax.set_xlabel('frequency (log scale)'); ax.set_ylabel('gain = $V_{out}$ / $V_{in}$')
    ax.grid(True, color='#ddd')
    for s in ('top', 'right'): ax.spines[s].set_visible(False)


def empty_graph():
    fig, ax = plt.subplots(figsize=(5.6, 4.4), dpi=170)
    axes(ax); ax.set_title('Draw your graph', fontsize=12)
    fig.tight_layout(); path = FIG + 'p3_gain_empty.png'; fig.savefig(path); plt.close(fig)
    return path


def theory_graph():
    f = np.logspace(math.log10(14), math.log10(300e3), 400)
    fig, ax = plt.subplots(figsize=(6.0, 4.6), dpi=170)
    axes(ax)
    ax.plot(f, gain(f), color='#3f7a2a', lw=2.5)
    ax.axhline(G, color='#c0392b', ls=':', lw=1.5)
    ax.text(16, G + 0.4, f'R2 / R1 ≈ {G:.1f}', color='#c0392b', fontsize=11)
    for x in ROWS:
        ax.plot(x, gain(x), 'o', color='#C2410C', ms=7)
        ax.annotate(f'{gain(x):.1f}', (x, gain(x)), textcoords='offset points', xytext=(4, -14), fontsize=9, color='#444')
    ax.set_title('R1 = 2.2 kΩ, R2 = 30 kΩ, C1 = 1 µF, LM358 (theory)', fontsize=11)
    fig.tight_layout(); path = FIG + 'p3_gain_graph.png'; fig.savefig(path); plt.close(fig)
    return path


def run():
    p = Presentation(DECK); lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if any(sh.name == MARK + ' bar' for sh in sl.shapes):
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    model = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == MODEL for sh in s.shapes))
    ex4 = next(s for s in p.slides if any(sh.name == 'Practice Exercise' for sh in s.shapes) and
               any(sh.has_text_frame and sh.text_frame.text.startswith('Exercise 4:') for sh in s.shapes))

    t = rlc_table.new_slide(p, model)
    textbox(t, 0.5, 1.2, 12.3, 0.5, [TABLE_LINE])
    rows = [('f', 'Vin (CH2, Vpp)', 'Vout (CH1, Vpp)', 'Gain = Vout / Vin')] + [(hz(f), '', '', '') for f in ROWS]
    tb = t.shapes.add_table(len(rows), 4, *(int(x * 914400) for x in (0.5, 1.9, 6.6, 5.0))); tb.name = MARK
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tb.table.cell(i, j); cell.text = ''
            add_runs(cell.text_frame.paragraphs[0], f'**{v}**' if i == 0 or j == 0 else v, 16, WHITE if i == 0 else rlc_table.DARK)
    picture(t, empty_graph(), 7.3, 1.8, 5.8, 5.1)
    t.notes_slide.notes_text_frame.text = (TABLE_LINE + '\nThe frequencies are 20, 50, 100 and 500 hertz, then 1, 10, 50, 100 '
                                           'and 200 kilohertz.')

    n = rlc_table.new_slide(p, model)
    textbox(n, 0.5, 1.2, 6.3, 0.5, ['What do you notice?'])
    textbox(n, 0.5, 1.9, 6.3, 5.0, NOTICE, gap=12)
    picture(n, theory_graph(), 6.9, 1.3, 6.1, 5.6)
    n.notes_slide.notes_text_frame.text = 'What do you notice?\n' + '\n'.join(x.replace('**', '') for x in NOTICE)

    for s in (n, t):                                       # Exercise 4, table, notice
        sid = lst[list(p.slides).index(s)]; lst.remove(sid)
        lst.insert(list(p.slides).index(ex4) + 1, sid)

    # "Use low (audio) frequencies" now explains the high end of the graph
    low = next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == LOW for sh in s.shapes))
    body = next(sh for sh in low.shapes if sh.has_text_frame and sh.name != 'Title Bar' and 'LM358' in sh.text_frame.text)
    paras = list(body.text_frame.paragraphs)
    size, font = paras[0].runs[0].font.size, paras[0].runs[0].font.name
    for para in paras[1:]:
        para._p.getparent().remove(para._p)
    tf = body.text_frame
    for r in list(paras[0].runs): r._r.getparent().remove(r._r)
    for k, line in enumerate(LOW_TEXT):
        para = paras[0] if k == 0 else tf.add_paragraph(); para.space_after = paras[0].space_after
        for part in re.split(r'(\*\*.*?\*\*)', line):
            if part:
                r = para.add_run(); r.text = part.strip('*'); r.font.size = size; r.font.name = font
                r.font.bold = part.startswith('**'); r.font.color.rgb = rlc_table.DARK
    low.notes_slide.notes_text_frame.text = '\n'.join(x.replace('**', '') for x in LOW_TEXT).replace(
        '100 Hz – 2 kHz', '100 hertz to 2 kilohertz')
    p.save(DECK)
    idx = lambda s: list(p.slides).index(s) + 1
    print(f'Lab 3: Exercise 4 at {idx(ex4)}, table at {idx(t)}, notice at {idx(n)}, "{LOW}" at {idx(low)} | {len(p.slides)} slides |',
          f'fL {FL:.0f} Hz, fH {FH / 1e3:.0f} kHz, gains', [f'{gain(f):.1f}' for f in ROWS])


if __name__ == '__main__':
    run()
