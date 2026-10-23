import os, runpy, zipfile, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins

W = os.path.dirname(__file__)
OUT = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab1'
os.makedirs(OUT, exist_ok=True)
V = runpy.run_path(f'{W}/variants.py')['D1']

# Canvas questions: (key used for MATCH, full question text, unit, marks, answer key fn)
Q = [('Variant number', 'Variant number (printed on your card)', '', 0, None),
     ('R_eq', 'R_eq – total resistance between the two supply nodes, supply disconnected. Unit: Ω (ohm). Numbers only.', 'Ω', 3, lambda d: d['Req']),
     ('Total current I', 'Total current I from the 5 V supply. Unit: mA. Numbers only.', 'mA', 3, lambda d: d['I'] * 1e3),
     ('I_A', 'I_A – current through R_A. Unit: mA. Numbers only.', 'mA', 3, lambda d: d['IA'] * 1e3),
     ('V_X', 'V_X – voltage at node X with respect to 0 V. Unit: V. Numbers only.', 'V', 3, lambda d: d['VX'])]

H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E78')
thin = Side(style='thin', color='999999'); B = Border(top=thin, bottom=thin, left=thin, right=thin)
GRN = PatternFill('solid', fgColor='E2EFDA'); YEL = PatternFill('solid', fgColor='FFF2CC')
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
# ---------------- Hướng dẫn ----------------
ws = wb.active; ws.title = 'Hướng dẫn'
guide = [
    ('LAB 1 – CÁCH DÙNG FILE NÀY', None),
    ('Trước buổi học', '1) In sheet "Phiếu đề (in)" của lớp mình, cắt thành phiếu, phát mỗi SV 1 phiếu.\n2) In sheet "Đáp án GV (in)" của lớp mình để cầm theo khi đi kiểm tra.'),
    ('Trong giờ', 'SV gọi → GV chọn 1 số bất kỳ (R_eq, I, I_A hoặc V_X), nhìn số trên VOM của SV, so với đáp án (lệch ≤ 10% là được) → đánh dấu ✓ vào cột "Tick" trên tờ đáp án.'),
    ('Cuối giờ', 'SV nộp Canvas quiz "Lab 1". Quiz đóng khi hết giờ.'),
    ('Sau giờ', '1) Canvas → Quiz Lab 1 → Quiz Statistics → Student Analysis → tải file CSV.\n2) Mở CSV, Ctrl+A, Ctrl+C → dán vào ô A1 của sheet "CSV Lớp 1" (hoặc Lớp 2/3).\n3) Sang sheet "Điểm Lớp 1": nhập 1 vào cột Tick cho SV có ✓ trên tờ đáp án. Điểm tự tính.'),
    ('Luật chấm', 'Mỗi số: lệch ≤ 10% so với đáp án → đủ điểm; ≤ 20% → nửa điểm; > 20% → 0. Không có Tick → 0 cả buổi (vì cả 4 số đều là số đo).\nLab 1: 4 số × 3 điểm = 12 điểm.'),
    ('Mã đề', 'Lớp 1: 101–118 · Lớp 2: 201–218 · Lớp 3: 301–318. Đáp án tính theo giá trị danh định của điện trở.'),
    ('', ''),
    ('TẠO QUIZ TRÊN CANVAS', None),
    ('Import', 'Course → Settings → Import Course Content → Content Type: "QTI .zip file" → chọn file Canvas_Lab1_QTI.zip → Import. Quiz "Lab 1" xuất hiện trong Quizzes (dạng Classic Quiz).'),
    ('Cài đặt nên chỉnh', '• Allowed attempts: 1\n• Bỏ chọn "Let Students See Their Quiz Responses" (quiz 0 điểm, điểm thật tính bằng file này)\n• Available from/Until: giờ bắt đầu → giờ kết thúc buổi học\n• Có thể đặt Access code đọc trong lớp để SV không nộp ở nhà'),
    ('Nội dung câu hỏi', '\n'.join(f'Q{i}. {q[1]}' for i, q in enumerate(Q, 1))),
]
for r, (a, b) in enumerate(guide, 1):
    ws.cell(r, 1, a).font = Font(bold=True, size=13 if b is None else 11)
    if b: ws.cell(r, 2, b).alignment = wrap
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110


def fmt_ohm(r):
    return f'{round(r):,} Ω'.replace(',', ' ')


for cls in (1, 2, 3):
    vs = {k: v for k, v in V.items() if k // 100 == cls}
    # ---------------- Phiếu đề ----------------
    pc = wb.create_sheet(f'Phiếu đề Lớp {cls} (in)')
    pc.page_setup.orientation = 'portrait'; pc.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
    pc.column_dimensions['A'].width = 44; pc.column_dimensions['B'].width = 3; pc.column_dimensions['C'].width = 44
    for i, (code, d) in enumerate(sorted(vs.items())):
        col = 1 if i % 2 == 0 else 3; r0 = (i // 2) * 9 + 1
        lines = [(f'Electric Circuit Lab – LAB 1   |   Variant {code}', True),
                 (f'R_A = {fmt_ohm(d["A"][0])}', True), (f'R_B = {fmt_ohm(d["B"][0])}', True), (f'R_C = {fmt_ohm(d["C"][0])}', True),
                 ('Supply: 5.00 V (see circuit on the slide)', False),
                 ('Make each resistor from standard resistors in series/parallel.', False),
                 ('Submit R_eq, I, I_A, V_X on Canvas quiz "Lab 1".', False),
                 ('Name: ______________________________', False)]
        for j, (t, b) in enumerate(lines):
            c = pc.cell(r0 + j, col, t); c.font = Font(bold=b, size=12 if j == 0 else 11)
            c.border = Border(left=thin, right=thin, top=thin if j == 0 else None, bottom=thin if j == len(lines) - 1 else None)
            if j == 0: c.fill = YEL
    # ---------------- Đáp án GV ----------------
    ak = wb.create_sheet(f'Đáp án Lớp {cls} (in)')
    ak.append([f'Lab 1 – Đáp án Lớp {cls} (nguồn 5 V) – chấp nhận ±10%'])
    ak['A1'].font = Font(bold=True, size=13)
    hdr = ['Mã đề', 'R_A', 'Ghép gợi ý', 'R_B', 'Ghép gợi ý', 'R_C', 'Ghép gợi ý', 'R_eq (Ω)', 'I (mA)', 'I_A (mA)', 'V_X (V)', 'Tên SV', 'Tick ✓']
    ak.append(hdr)
    for c in ak[2]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', wrap_text=True)
    f = lambda r: f'{r/1000:g}k' if r >= 1000 else f'{r:g}'
    for code, d in sorted(vs.items()):
        row = [code]
        for k in 'ABC':
            v, t, a, b = d[k]; row += [round(v), f'{f(a)}{" + " if t == "+" else " // "}{f(b)}']
        row += [round(d['Req']), round(d['I'] * 1e3, 3), round(d['IA'] * 1e3, 3), round(d['VX'], 3), '', '']
        ak.append(row)
    for row in ak.iter_rows(min_row=2, max_row=ak.max_row):
        for c in row: c.border = B
    for r in range(3, ak.max_row + 1): ak.row_dimensions[r].height = 24
    for col, w in zip('ABCDEFGHIJKLM', [7, 7, 11, 7, 11, 7, 11, 9, 8, 8, 8, 22, 7]): ak.column_dimensions[col].width = w
    ak.page_setup.orientation = 'landscape'; ak.page_setup.fitToWidth = 1

    # ---------------- Key (hidden lookup) ----------------
    key = wb.create_sheet(f'Key{cls}')
    key.append(['Variant'] + [q[0] for q in Q[1:]])
    for code, d in sorted(vs.items()):
        key.append([code] + [q[4](d) for q in Q[1:]])
    key.sheet_state = 'hidden'

    # ---------------- CSV paste + scores ----------------
    csv = wb.create_sheet(f'CSV Lớp {cls}')
    csv['A1'] = 'DÁN TOÀN BỘ FILE CSV "Student Analysis" CỦA CANVAS VÀO Ô A1 (ghi đè dòng này)'
    csv['A1'].font = Font(bold=True, color='C00000')
    sc = wb.create_sheet(f'Điểm Lớp {cls}')
    cols = ['Tên SV', 'Mã đề', 'Tick (1 = có ✓)'] + [f'{q[0]} ({q[2]})' for q in Q[1:]] + [f'Điểm {q[0]}' for q in Q[1:]] + ['TỔNG /12', 'Ghi chú']
    sc.append(cols)
    for c in sc[1]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center', wrap_text=True)
    CS = f"'CSV Lớp {cls}'"; KS = f'Key{cls}'
    nq = len(Q) - 1
    for r in range(2, 62):
        get = lambda pat: f'IFERROR(VALUE(INDEX({CS}!$A$1:$ZZ$200,ROW(),MATCH("*{pat}*",{CS}!$1:$1,0))),"")'
        sc.cell(r, 1, f'=IF({CS}!A{r}="","",{CS}!A{r})')
        sc.cell(r, 2, '=' + get(Q[0][0]))
        sc.cell(r, 3).fill = YEL
        for i, q in enumerate(Q[1:]):
            vc = 4 + i; scol = 4 + nq + i
            sc.cell(r, vc, '=' + get(q[0]))
            vref = sc.cell(r, vc).coordinate
            k = f'VLOOKUP($B{r},{KS}!$A:${chr(65 + nq)},{i + 2},0)'
            sc.cell(r, scol, f'=IFERROR(IF($C{r}<>1,0,IF(ABS({vref}/{k}-1)<=0.1,{q[3]},IF(ABS({vref}/{k}-1)<=0.2,{q[3] / 2},0))),0)')
        first = sc.cell(r, 4 + nq).coordinate; last = sc.cell(r, 3 + 2 * nq).coordinate
        sc.cell(r, 4 + 2 * nq, f'=IF(A{r}="","",SUM({first}:{last}))').font = Font(bold=True)
        sc.cell(r, 5 + 2 * nq, f'=IF(A{r}="","",IF(ISERROR(VLOOKUP($B{r},{KS}!$A:$A,1,0)),"Sai mã đề / chưa nhập",IF(C{r}<>1,"Chưa có tick","")))')
    for i, w in enumerate([24, 8, 9] + [10] * nq + [9] * nq + [9, 22], 1):
        sc.column_dimensions[sc.cell(1, i).column_letter].width = w
    sc.freeze_panes = 'D2'

wb.move_sheet('Hướng dẫn', -len(wb.sheetnames))
wb.save(f'{OUT}/HCM_Lab1_Phieu_de_va_Cham_diem.xlsx')

# ---------------- Canvas QTI 1.2 ----------------
def item(i, title, text):
    return f'''<item ident="lab1_q{i}" title="{html.escape(title)}">
<itemmetadata><qtimetadata>
<qtimetadatafield><fieldlabel>question_type</fieldlabel><fieldentry>numerical_question</fieldentry></qtimetadatafield>
<qtimetadatafield><fieldlabel>points_possible</fieldlabel><fieldentry>0.0</fieldentry></qtimetadatafield>
</qtimetadata></itemmetadata>
<presentation><material><mattext texttype="text/html">{html.escape('<p>' + html.escape(text) + '</p>')}</mattext></material>
<response_str ident="response1" rcardinality="Single"><render_fib fibtype="Decimal"><response_label ident="answer1"/></render_fib></response_str>
</presentation>
<resprocessing><outcomes><decvar maxvalue="100" minvalue="0" varname="SCORE" vartype="Decimal"/></outcomes>
<respcondition continue="No"><conditionvar><vargte respident="response1">-1000000000</vargte><varlte respident="response1">1000000000</varlte></conditionvar>
<setvar action="Set" varname="SCORE">100</setvar></respcondition></resprocessing>
</item>'''

items = '\n'.join(item(i, q[0], q[1]) for i, q in enumerate(Q, 1))
assess = f'''<?xml version="1.0" encoding="UTF-8"?>
<questestinterop xmlns="http://www.imsglobal.org/xsd/ims_qtiasiv1p2" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:schemaLocation="http://www.imsglobal.org/xsd/ims_qtiasiv1p2 http://www.imsglobal.org/xsd/ims_qtiasiv1p2p1.xsd">
<assessment ident="hcm_lab1" title="Lab 1">
<qtimetadata><qtimetadatafield><fieldlabel>cc_maxattempts</fieldlabel><fieldentry>1</fieldentry></qtimetadatafield></qtimetadata>
<section ident="root_section">
{items}
</section></assessment></questestinterop>'''
manifest = '''<?xml version="1.0" encoding="UTF-8"?>
<manifest identifier="hcm_lab1_manifest" xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<metadata><schema>IMS Content</schema><schemaversion>1.1.3</schemaversion></metadata>
<organizations/>
<resources><resource identifier="hcm_lab1" type="imsqti_xmlv1p2"><file href="hcm_lab1/hcm_lab1.xml"/></resource></resources>
</manifest>'''
with zipfile.ZipFile(f'{OUT}/Canvas_Lab1_QTI.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('imsmanifest.xml', manifest)
    z.writestr('hcm_lab1/hcm_lab1.xml', assess)
print('ok')
