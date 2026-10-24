"""Align the 4 HCM decks with the web workflow (run on the final decks, after add_qr.py):
submit on the student page first, then the lecturer's tick, then tidy the table; one photo; parts box;
Lab 2 Task 1 individual / Task 2 team; microphone on female-male jumpers; a 'tidy your table' slide with the photo."""
import copy, io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
PHOTO = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/GradingApp/tools/content/fig/bench_tidy.jpg'
TIDY_TITLE = 'Before you leave – make your table look like this'


def runs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                yield from p.runs
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        yield from p.runs


def rep(slide, pairs):
    """Replace whole runs (exact text); each must exist once. Already-new text is skipped (safe to run twice)."""
    rs = list(runs(slide))
    if all(any(r.text == new for r in rs) for _, new in pairs):     # already done (some pairs swap texts)
        return
    for old, new in pairs:
        hit = [r for r in rs if r.text == old]
        if not hit and any(r.text == new for r in rs):
            continue
        assert len(hit) == 1, (old, len(hit))
        hit[0].text = new


def by_title(p, *titles):
    """The slide whose title bar shows one of these titles (old or new wording), so inserted slides do not shift it."""
    for s in p.slides:
        for sh in s.shapes:
            if sh.name == 'Title Bar' and sh.text_frame.text.strip() in titles:
                return s
    raise LookupError(titles)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def tidy_slide(p, model, extra):
    """New last slide: title bar + logo copied from `model`, the table photo, a 6-step checklist."""
    s = p.slides.add_slide(model.slide_layout)
    for sh in model.shapes:
        if sh.name == 'Title Bar':
            el = copy.deepcopy(sh._element); s.shapes._spTree.append(el)
            bar = s.shapes[-1]; bar.text_frame.paragraphs[0].runs[0].text = TIDY_TITLE
            for r in bar.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
        if sh.name.startswith('Logo'):
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
    s.shapes.add_picture(PHOTO, Inches(1.4), Inches(1.1), Inches(10.5))
    cols = [['Wait for the lecturer’s tick before you pack up.', 'Switch off every instrument and the PB-503A.', 'Clear the breadboard: no parts, no wires.'],
            [f'Resistors you used: in the bin. Every other part back in its bag{extra}.', 'Test leads, probes and BNC cables into the box behind the trainer.', 'Instruments in their place, like the photo. Push your chair in.']]
    for c, items in enumerate(cols):
        tb = s.shapes.add_textbox(Inches(0.55 + 6.15 * c), Inches(4.85), Inches(6.0), Inches(1.9)); tf = tb.text_frame; tf.word_wrap = True
        for i, t in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = para.add_run(); r.text = f'{3 * c + i + 1}.  {t}'; r.font.size = Pt(15); r.font.name = 'Arial'; r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            para.space_after = Pt(4)
    notes(s, 'Before you leave, your table must look like this photo.\n'
             'First wait for my tick. Then switch off every instrument and the trainer.\n'
             'Clear the breadboard. The resistors you used go into the bin. Every other part goes back to its bag in the parts box'
             + (extra.replace(' (', '. ').replace(')', '') if extra else '') + '.\n'
             'Put the leads, probes and cables in the box behind the trainer. Put the instruments back in their place, like the photo, and push your chair in.')


def strip_tidy(p):
    """Remove the tidy-bench slides made by an earlier run (they are added again, so running twice is safe)."""
    for sid, sl in reversed(list(zip(list(p.slides._sldIdLst), list(p.slides)))):
        if any(sh.has_text_frame and sh.text_frame.text == TIDY_TITLE for sh in sl.shapes):
            p.part.drop_rel(sid.rId); p.slides._sldIdLst.remove(sid)
    # number the slide files 1..n again: a new slide is named slide<n+1>.xml and must not hit an existing file
    p.part.rename_slide_parts([x.rId for x in p.slides._sldIdLst])


def move(p, slide, index):
    lst = p.slides._sldIdLst
    sid = next(x for x, sl in zip(list(lst), list(p.slides)) if sl == slide)
    lst.remove(sid); lst.insert(index, sid)


def lab1(p):
    s = by_title(p, 'How each lab session is marked')
    rep(s, [('Call the lecturer', 'Submit on the student page'),
            ('The lecturer checks ', 'Enter your values, add '),
            ('one reading', 'one photo'),
            (' on your instrument and ticks your name.', ' of your circuit, tap Submit.'),
            ('Submit your values', 'Call the lecturer'),
            ('On the student page, enter your values and add 1–3 ', 'The lecturer checks '),
            ('photos of your circuit', 'one reading'),
            (' before you leave.', ' on your instrument and ticks you. Pack up only after the tick.')])
    notes(s, 'This is how I give marks in every lab. Only 4 steps.\n'
             'Step 1: your card is on the student page. It has your own variant number. Your circuit is different from your friends.\n'
             'Step 2: you build and measure.\n'
             'Step 3: on the student page, type your numbers, add one photo of your circuit, and tap Submit.\n'
             'Step 4: call me. I check one number on your instrument. It must be the same as what you submitted. Then I give you a tick. Do not pack up before the tick.\n'
             'The rule is simple: error 10% or less = full marks. 10 to 20% = half. More than 20%, or no tick = 0.\n'
             'If your answer is 1.27 mA, what do you type? … The answer is: 1.27 — number only.')
    s = by_title(p, 'Devices and components for this session')
    rep(s, [('Reusable parts (LED…) go back to the tray at the end of the session.', 'Reusable parts (LED…) go back to the parts box at the end of the session.'),
            ('Resistors for your variant (listed on your card)', 'Resistors from the parts box: 220, 330, 470, 680 Ω, 1k, 1.5k, 2.2k, 3.3k, 4.7k, 6.8k, 10k, 15k, 22k Ω')])
    notes(s, 'Each team has one PB-503A trainer. It has a breadboard, DC supplies and a speaker.\n'
             'You also have a multimeter, an oscilloscope and a function generator.\n'
             'Take the resistors you need from the parts box of your table. There are 13 standard values.\n'
             'Look at the PB-503A on your table. Here is the breadboard, here are the supply outputs, and here is the speaker.')
    s = by_title(p, 'Lab 1 exercise – Equivalent resistance')
    rep(s, [('Call the lecturer. Then submit your numbers and a photo of your circuit on the student page.',
             'Submit your numbers and one photo of your circuit on the student page. Then call the lecturer for the tick.')])
    notes(s, 'Look at the circuit. R A and R B are in parallel. Then R C and R D are in series.\n'
             'Your card gives R A, R B, R C and R D. These values are not sold. You must make each one from two standard resistors. For example, 5.7 kilo-ohm is 1 kilo-ohm plus 4.7 kilo-ohm.\n'
             'Follow the steps. Measure R eq before you connect the supply.\n'
             'Open your card on the student page now. Take the resistors from the parts box.')
    s = by_title(p, 'Lab 1 – What to submit')
    rep(s, [('Add 1–3 photos of your real circuit.', 'Add one photo of your real circuit.'),
            ('Before you leave: (1) submit your values on the ', 'Before you leave: (1) submit on the '),
            (', (2) switch off all instruments, (3) put parts back in the tray, (4) clear your breadboard.', ', (2) call the lecturer for the tick, (3) tidy your table (next slide).')])
    notes(s, 'Type these 4 numbers on the student page, under your card. Use the unit in the table.\n'
             'Add one photo of your real circuit. The photo is required.\n'
             'Use your measured values, not calculated values.\n'
             'Tap Submit, then call me for the tick. To change a number before the tick, tap Unsubmit, edit, and submit again.\n'
             'After the tick, tidy your table. The next slide shows how.')
    tidy_slide(p, by_title(p, 'Lab 1 – What to submit'), '')
    # Lab 1 also shows it at the start, after "Your team and the student page", so the first class knows the rule
    tidy_slide(p, by_title(p, 'Lab 1 – What to submit'), '')
    team = list(p.slides).index(by_title(p, 'Your team and the student page (10 minutes)', 'Make your team (10 minutes)'))
    early = p.slides[len(p.slides) - 1]; move(p, early, team + 1)
    notes(early, 'One more rule for every lab. At the end, your table must look like this photo.\n'
                 'First wait for my tick. Then switch off everything, clear the breadboard, put the used resistors in the bin and every other part back in its bag, '
                 'put the cables in the box behind the trainer, and put the instruments in their place.\n'
                 'You will see this slide again at the end of every lab.')


def lab2(p):
    s = by_title(p, 'Devices and components for this session')
    rep(s, [('For each student:', 'For each student (Task 1) and each team (Task 2):'),
            ('RC warm-up: 1 resistor + 1 capacitor (on your variant card)', 'Task 1, each student: 1 resistor + 1 capacitor (on your card)'),
            ('RLC: 1 inductor 1 mH, capacitors 22 nF and 2.2 nF', 'Task 2, each team: 1 inductor 1 mH, capacitors 22 nF and 2.2 nF')])
    notes(s, 'Each team uses the oscilloscope, function generator, multimeter and PB-503A.\n'
             'Task 1 is individual: each student gets one resistor and one capacitor for the RC warm-up.\n'
             'Task 2 is team work: each team gets one inductor, the capacitors and the resistors for one RLC filter.\n'
             'Download the Excel template from the student page now.')
    s = by_title(p, 'Warm-up exercise – RC cut-off frequency', 'Task 1 (individual) – RC cut-off frequency')
    rep(s, [('Warm-up exercise – RC cut-off frequency', 'Task 1 (individual) – RC cut-off frequency')])
    notes(s, 'Task 1 is individual: 2 points. Your card gives the type, R and C.\n'
             'Calculate fc. Build it. Change the frequency until Vout = 0.707 × Vin. That is your measured fc. Write it in kHz.\n'
             'One student uses the oscilloscope at a time. The others build or calculate. When you finish, give the oscilloscope to the next student.\n'
             'Open your card on the student page now.')
    s = by_title(p, 'Lab 2 exercise – design and build', 'Task 2 (team) – design and build')
    rep(s, [('Lab 2 exercise – design and build', 'Task 2 (team) – design and build'),
            ('Your variant card gives the ', 'Your card gives the ')])
    notes(s, 'Task 2 is team work. The whole team builds one filter. Everybody in the team has the same Task 2 on the card.\n'
             'Your card gives the filter type and Q. Measure R L. Calculate R total = 203 / Q. Then R = R total − R L. Make R from 47, 100 and 220 Ω.\n'
             'Measure R and the coil together with the ohmmeter. That is your R total for the student page.\n'
             'Show me your calculation before you build.')
    s = by_title(p, 'Lab 2 exercise – measure and plot', 'Task 2 (team) – measure and plot')
    rep(s, [('Lab 2 exercise – measure and plot', 'Task 2 (team) – measure and plot'),
            (' from the student page (under your Lab 2 card) and open it. For each frequency on your card, measure ',
             ' from the student page (under your Lab 2 card). Share the frequencies on your card: each member measures some of them. Measure '),
            (' (CH2), peak-to-peak. Type them in. The graph draws itself.', ' (CH2) with Measure → Pk-Pk. Type them in. The graph draws itself.'),
            ('. Add more points near the peak. Write down ', '. Add at most 2 points near the peak. Write down '),
            ('Call the lecturer. Show one reading on the oscilloscope ', 'Each member: submit the team values and one screenshot of the graph on the student page. Then call the lecturer: show one reading on the oscilloscope '),
            (' your graph. Add a photo of your circuit on the student page.', ' your graph.')])
    notes(s, 'Download the Excel file from the student page, under your Lab 2 card.\n'
             'Share the work: each member measures some of the frequencies on the card. Use Measure, Pk-Pk on CH1 and CH2. Do not count squares.\n'
             'Type the values into the Excel file. The graph draws itself. Add at most 2 points near the peak.\n'
             'Low-pass and high-pass: write the highest Vout / Vin. Band-pass: find f L and f H, then BW.\n'
             'Write 1 to 3 sentences of conclusion.\n'
             'Each member submits the team values and one screenshot of the graph on the student page. Not a photo of the circuit. Then call me.')
    s = by_title(p, 'Lab 2 – What to submit')
    rep(s, [('Warm-up: measured cut-off frequency fc', 'Task 1 (individual): measured cut-off frequency fc'),
            ('Measured ', 'Task 2: measured '),
            ('Frequency of maximum output ', 'Task 2: frequency of maximum output '),
            ('Low-pass / high-pass: highest ', 'Task 2 – low-pass / high-pass: highest '),
            ('Band-pass: bandwidth BW', 'Band-pass: bandwidth BW   (this row: full marks within ±20 %, half within ±35 %)'),
            ('Graph + conclusion (checked by the lecturer at your table)', 'Task 2: graph + conclusion (one screenshot on the student page, checked at your table)'),
            ('Before you leave: submit your values on the ', 'Before you leave: submit on the '),
            (', switch off all instruments and clear your breadboard.', ', get the lecturer’s tick, then tidy your table (next slide).')])
    notes(s, 'Task 1 is your own value. Task 2 values are the same for the whole team.\n'
             'Type these numbers on the student page, under your card. Use kHz for frequencies. Add one screenshot of your RLC graph.\n'
             'Then call me for the tick. After the tick, tidy your table. The next slide shows how.')
    tidy_slide(p, by_title(p, 'Lab 2 – What to submit'), '')


def lab3(p):
    s = p.slides[4]
    rep(s, [('Resistors for your amplifier: take them from the team tray once you know R1 and R2', 'Resistors for your amplifier: take them from the parts box once you know R1 and R2')])
    notes(s, 'Each student gets one LM358 and two 100 nF capacitors. Each team gets two 470 µF capacitors.\n'
             'The LM358 and the 470 µF capacitors go back to the parts box at the end. We use them again in Lab 4.')
    s = p.slides[35]
    rep(s, [('Call the lecturer. Then submit your values and a photo of your circuit on the student page.',
             'Submit your values and one photo of your circuit on the student page. Then call the lecturer for the tick.')])
    s = p.slides[36]
    rep(s, [('Before you leave: submit your values on the ', 'Before you leave: submit on the '),
            (', return the LM358 and the 470 µF capacitors to the tray, switch off all instruments and clear your breadboard.',
             ', get the lecturer’s tick, return the LM358 and the 470 µF capacitors to the parts box, then tidy your table (next slide).')])
    notes(s, 'Type these numbers on the student page, under your card, and add one photo of your circuit. Gain and R2/R1 have no unit. Vout in volts. Phase in degrees.\n'
             'Then call me for the tick. After the tick: LM358 and 470 µF capacitors back to the parts box, then tidy your table. The next slide shows how.')
    tidy_slide(p, p.slides[36], ' (LM358 and 470 µF too)')


def lab4(p):
    s = p.slides[4]
    rep(s, [('LM358 op-amp, electret microphone with soldered leads', 'LM358 op-amp, electret microphone on 2 female–male jumper wires'),
            ('Return the op-amp, the microphone and the 470 µF capacitors to the tray at the end.', 'Return the op-amp, the microphone (with its jumpers) and the 470 µF capacitors to the parts box at the end.')])
    notes(s, 'You need the PB-503A with its speaker, the instruments, one LM358, one microphone and these parts.\n'
             'The microphone has two short legs. A female–male jumper wire is pushed onto each leg. Keep the jumpers on the microphone and handle it carefully.')
    s = p.slides[39]
    rep(s, [('The microphone leads are short: they are soldered to wires so that it fits the breadboard. Handle it carefully.',
             'The microphone legs are short: a female–male jumper is pushed onto each leg; plug the male ends into the breadboard. Black jumper = case pin = ground.')])
    notes(s, 'This is the circuit on the PB-503A. Use the +6 V output for the bias resistor.\n'
             'Your microphone has a jumper wire on each leg. The black jumper is on the pin connected to the metal case. It goes to ground.\n'
             'Do not pull the jumpers off the microphone. Show me your connections before you switch on.')
    s = p.slides[84]
    rep(s, [('Submit your value and a photo of your circuit on the ', 'Submit your value and one photo of your circuit on the '),
            (' before the exam review.', ', then call the lecturer for the tick – before the exam review.')])
    notes(s, 'Team demo: 4 points. Speak into the microphone. I must hear you in the speaker.\n'
             'Then each member does one measurement alone: 6 points. Look at your card: member 1, 2, 3, or 4 in a team of 4.\n'
             'Disconnect the microphone. Put 10 mVpp from the function generator into stage 1. Measure the Vpp asked on your card.\n'
             'Submit your value and one photo of your circuit on the student page. Then call me for the tick.')
    s = p.slides[91]
    rep(s, [('Before you leave: switch off all instruments, return the op-amp, the microphone and the 470 µF capacitors to the tray, clear your breadboard.',
             'Before you leave: return the op-amp, the microphone (with its jumpers) and the 470 µF capacitors to the parts box, then tidy your table (next slide).')])
    tidy_slide(p, p.slides[84], ' (LM358, microphone, 470 µF too)')


for n, fix in ((1, lab1), (2, lab2), (3, lab3), (4, lab4)):
    path = f'{HCM}/Lab{n}/HCM Electronic Circuit Lab {n}.pptx'
    p = Presentation(path); strip_tidy(p); fix(p); p.save(path); print(n, len(p.slides), 'slides')
