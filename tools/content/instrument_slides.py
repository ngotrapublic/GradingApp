"""Lab 2: how to use the instruments, step by step, with photos of the HCM instruments and numbered marks (lecturer,
24/10/2026: "experiments must be detailed, concrete, exact"; small slides; a table to fill in).
- before Exercise 1: "Oscilloscope – measure Vin and Vout" (TBS1052C) and "Function generator – set the frequency" (AFG1022)
- after "Task 2 (team) – measure and plot": "Task 2 – find the peak (all teams)", "Task 2 – find fL, fH and BW (band-pass
  teams)", "Task 2 – BW table" (table + common mistakes)
Figures: fig_instruments.py. Run after rc_section.py / rlc_table.py, before subscripts.py. Running again replaces the slides."""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from rlc_slides import add_runs, text_box, new_slide, FIG

DECK = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM/Lab2/HCM Electronic Circuit Lab 2.pptx'
MODEL = 'Devices and components for this session'
PRACTICE = 'Practice Exercise'

SCOPE = ('Oscilloscope – measure Vin and Vout', 'inst_scope.png', [
    'Use a **10X probe** on each channel: CH1 to V_{in}, CH2 to V_{out}. Both ground clips go to the circuit ground.',
    '1. Press **(4)**, the **1** button. Press the side button **(3) Probe Setup**. Turn **(2)** to **Attenuation**, push **(2)**, '
    'turn it to **10X**, push **(2)** again. Do the same for CH2 with **(5)**.',
    '2. Press **(1) Measure**, then the side button **(3)** next to **CH1**. Turn **(2)** to **Peak-to-Peak** and push **(2)** to select it.',
    '3. Press the side button next to **CH2** and select **Peak-to-Peak** in the same way.',
    '4. The bottom of the screen now shows **CH1 Pk-Pk = V_{in}** and **CH2 Pk-Pk = V_{out}**.',
    '5. If a wave is too small or too big, turn the **Scale** knob of that channel.'])
GEN = ('Function generator – set the frequency', 'inst_gen.png', [
    '1. Press **(1) Sine**.',
    '2. Press **(2)**, the button next to **Freq**.',
    '3. Type the number on the **(3)** number pad, for example 32.4.',
    '4. Press the button **(4)** next to the unit, for example **kHz**.',
    '5. Press **(7) On/Off** (Out1). Its light must be on.',
    '6. Small changes: turn the **(5)** round knob. The **(6) ◄ ►** keys choose which digit the knob changes.',
    'Amplitude: the button next to **Ampl**, type 2, the button next to **Vpp**. Load: High-Z (Utility → Output Setup → Load, as in Lab 1).'])
PEAK = ('Task 2 – find the peak (all teams)', 'p2_bp_scope.png', [
    '1. Set the function generator to **f_{0} = 32.4 kHz** (see "Function generator – set the frequency").',
    '2. Turn the **round knob** slowly: a little down, then a little up. Change the 1 kHz digit first, then the 0.1 kHz digit.',
    '3. **What you see:** near the peak, the CH2 wave grows. At the peak it is the **biggest**. After the peak it becomes smaller again.',
    '4. Stop where the CH2 wave is the biggest compared with CH1. Write f, V_{in} and V_{out} in **row 1** of the table. '
    'Calculate V_{out}/V_{in}. This frequency is your **f_{max}**.',
    '**Band-pass:** at the peak, CH2 is almost as big as CH1, but a little smaller. '
    '**Low-pass and high-pass:** at the peak, CH2 can be **bigger** than CH1.',
    '5. Band-pass teams only: calculate the **target = 0.707 × V_{out}/V_{in} at the peak**. Write it in **row 2**.'])
EDGES = ('Task 2 – find fL, fH and BW (band-pass teams)', 'rlc_bp_bw.png', [
    '1. Turn the round knob to a **lower** frequency. The CH2 wave becomes smaller. Stop when V_{out}/V_{in} = the target.',
    '2. Read V_{in} and V_{out} again, because V_{in} changes near the peak. Write f, V_{in} and V_{out} in **row 3**. This frequency is **f_{L}**.',
    '3. Turn the knob to a **higher** frequency, past the peak, until V_{out}/V_{in} = the target again. Write the values in **row 4**. This is **f_{H}**.',
    '4. Calculate **BW = f_{H} − f_{L}**. Write it in **row 5**.',
    'Tip: near the target, change the 0.1 kHz digit, so that you do not jump past it.'])
TABLE_TITLE = 'Task 2 – BW table'
ROWS = [('Step', 'f (kHz)', 'V_{in} (CH1, Vpp)', 'V_{out} (CH2, Vpp)', 'V_{out} / V_{in}'),
        ('1. Peak (all teams)', '', '', '', ''),
        ('2. Target = 0.707 × peak', '–', '–', '–', ''),
        ('3. f_{L} (below the peak)', '', '', '', ''),
        ('4. f_{H} (above the peak)', '', '', '', ''),
        ('5. BW = f_{H} − f_{L}', '', '–', '–', '–')]
MISTAKES = ['Do not use 0.707 × V_{in}. Use 0.707 × the **V_{out}/V_{in} at the peak**.',
            'Do not read V_{in} only once. Near the peak V_{in} changes. Read **CH1 and CH2 again** at every frequency.',
            'Do not take f_{L} and f_{H} from the 9 frequencies on your card. They are too far apart. **Turn the knob** to find them.']
TITLES = {SCOPE[0], GEN[0], PEAK[0], EDGES[0], TABLE_TITLE}


def spoken(lines):
    return '\n'.join(re.sub(r'_\{(.*?)\}', r'\1', l.replace('**', '')) for l in lines)


def find_title(p, title):
    return next(s for s in p.slides if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() == title for sh in s.shapes))


def how_to(p, model, spec, photo_w=6.9):
    title, fig, lines = spec
    s = new_slide(p, model, title)
    from PIL import Image
    iw, ih = Image.open(FIG + '/' + fig).size
    h = min(5.9, photo_w * ih / iw); w = h * iw / ih
    s.shapes.add_picture(f'{FIG}/{fig}', Inches(13.0 - w), Inches(1.15), Inches(w), Inches(h))
    text_box(s, 0.45, 1.1, 12.55 - w - 0.45, 6.0, lines, size=15 if w < 7 else 14)
    s.notes_slide.notes_text_frame.text = spoken(lines)
    return s


def move(p, slides, index):
    lst = p.slides._sldIdLst
    sids = [lst[list(p.slides).index(s)] for s in slides]
    for sid in sids: lst.remove(sid)
    for k, sid in enumerate(sids): lst.insert(index + k, sid)


def run():
    p = Presentation(DECK); lst = p.slides._sldIdLst
    for sid, sl in reversed(list(zip(list(lst), list(p.slides)))):
        if any(sh.name == 'Title Bar' and sh.text_frame.text.strip() in TITLES for sh in sl.shapes):
            p.part.drop_rel(sid.rId); lst.remove(sid)
    p.part.rename_slide_parts([x.rId for x in lst])
    model = find_title(p, MODEL)

    scope, gen = how_to(p, model, SCOPE, photo_w=7.9), how_to(p, model, GEN, photo_w=7.9)
    peak, edges = how_to(p, model, PEAK, photo_w=4.4), how_to(p, model, EDGES, photo_w=5.6)
    t = new_slide(p, model, TABLE_TITLE)
    tb = t.shapes.add_table(len(ROWS), 5, Inches(0.5), Inches(1.25), Inches(12.3), Inches(3.4))
    for r, row in enumerate(ROWS):
        for c, v in enumerate(row):
            cell = tb.table.cell(r, c); cell.text = ''
            add_runs(cell.text_frame.paragraphs[0], f'**{v}**' if r == 0 or c == 0 else v, 15)
            if r == 0:
                for run_ in cell.text_frame.paragraphs[0].runs: run_.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for c, w in enumerate((3.5, 1.8, 2.3, 2.4, 2.3)):
        tb.table.columns[c].width = Inches(w)
    text_box(t, 0.5, 4.95, 12.3, 2.0, ['•  ' + m for m in MISTAKES], size=15, box=True, title='Common mistakes')
    t.notes_slide.notes_text_frame.text = ('Fill in this table step by step: row 1 is the peak, row 2 is the target, row 3 is f L, '
                                           'row 4 is f H, and row 5 is the bandwidth.\nCommon mistakes.\n' + spoken(MISTAKES))

    ex1 = next(s for s in p.slides if any(sh.name == PRACTICE for sh in s.shapes) and
               any(sh.has_text_frame and sh.text_frame.text.startswith('Exercise 1:') for sh in s.shapes))
    move(p, [scope, gen], list(p.slides).index(ex1))
    plot = find_title(p, 'Task 2 (team) – measure and plot')
    move(p, [peak, edges, t], list(p.slides).index(plot) + 1)

    p.save(DECK)
    idx = lambda s: list(p.slides).index(s) + 1
    print('Lab 2:', ', '.join(f'{idx(s)} {s.shapes[0].text_frame.text.strip() if False else n}' for s, n in
          [(scope, 'scope'), (gen, 'generator'), (ex1, 'Exercise 1'), (plot, 'measure and plot'), (peak, 'peak'), (edges, 'fL fH'), (t, 'table')]),
          '|', len(p.slides), 'slides')


if __name__ == '__main__':
    run()
