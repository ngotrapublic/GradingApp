"""Packing list for the 6 bench boxes (one bag per part type) + printable labels."""
import json, math, re, sys
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.pagebreak import Break

HCM = '/Users/mac/Tra/SE9_ElecCircuitLab/HCM'
K = json.load(open(f'{HCM}/GradingApp/courses/SE009/keys.json'))
bom = openpyxl.load_workbook(f'{HCM}/docs/HCM_Linh_kien_can_mua.xlsx', data_only=True).worksheets[0]
parts, order, kind, thi = [], {}, {}, {}
for r in bom.iter_rows(min_row=5, values_only=True):
    if not r[0]: continue
    parts.append(r[1]); order[r[1]] = r[13]; kind[r[1]] = r[2]; thi[r[1]] = r[7]

def R(v):  # '6.8k' / '220' / '1.2kΩ' -> part name
    v = v.replace('Ω', '').strip()
    return f'Điện trở {v}Ω'
for p in parts: pass
CAP = {'10 nF': 'Tụ film 10nF', '22 nF': 'Tụ film 22nF', '47 nF': 'Tụ film 47nF'}

need = defaultdict(lambda: defaultdict(int))     # need[bench][part]
labs = defaultdict(set)
def add(b, p, n, lab):
    assert p in order, p
    need[b][p] += n; labs[p].add(lab)

for b in range(1, 7):
    vs = [c * 100 + 3 * (b - 1) + m for c in (1, 2, 3) for m in (1, 2, 3)] + [200 + 18 + b]   # + one member 4 (class of 19)
    for v in vs:
        v = str(v)
        for line in K['labs']['1']['variants'][v]['info']:
            for x in re.findall(r'\(([^)]*)\)', line)[0].replace('//', '+').split('+'):
                add(b, R(x), 1, 'Lab 1')
        add(b, R('2.2k'), 1, 'Lab 1'); add(b, R('4.7k'), 1, 'Lab 1')
        rc = K['labs']['2']['variants'][v]['info'][1]
        m = re.search(r'R = ([\d.]+k)Ω, C = (\d+ nF)', rc)
        add(b, R(m.group(1)), 1, 'Lab 2'); add(b, CAP[m.group(2)], 1, 'Lab 2')
        m = re.search(r'R1 = ([\d.]+k?)Ω, R2 = ([\d.]+k?)Ω', K['labs']['3']['variants'][v]['info'][0])
        add(b, R(m.group(1)), 1, 'Lab 3'); add(b, R(m.group(2)), 1, 'Lab 3'); add(b, 'Tụ gốm 100nF', 2, 'Lab 3')
    for p, n in [('Cuộn cảm 1mH 0410-102', 1), ('Tụ film 22nF', 1), ('Tụ film 2.2nF', 1), (R('47'), 2), (R('100'), 2), (R('220'), 1)]:
        add(b, p, 3 * n, 'Lab 2')                       # Task 2 (RLC) per team: 3 classes, one team per bench
    for p, n in [(R('10k'), 5), (R('2.2k'), 2), (R('4.7k'), 1), (R('47k'), 1), ('Tụ gốm 1nF', 1), ('Tụ gốm 100nF', 4), ('Tụ hóa 1µF 50V', 1), ('Tụ hóa 10µF 25V', 1)]:
        add(b, p, 3 * n, 'Lab 4')                       # 3 classes, one team per bench
    # parts collected back after each class: enough for one class on one bench
    for p, n, lab in [('IC LM358L (DIP-8)', 4, 'Lab 3, 4'), ('Tụ hóa 470µF 25V', 4, 'Lab 3, 4'), ('Micro Electret', 1, 'Lab 4'), ('LED 5mm', 3, 'Lab 1')]:
        add(b, p, n, lab)
    for p in parts:
        if thi[p]: add(b, p, math.ceil(thi[p] / 6), 'Thi')

# put need + spare in each bag; whatever is left goes to the reserve bag
put, short = {}, {}
for p in parts:
    n = [need[b][p] for b in range(1, 7)]
    if not sum(n): continue
    spare = [0 if kind[p] == 'Thu hồi' else max(1, math.ceil(0.1 * x)) for x in n]
    while sum(n) + sum(spare) > order[p] and any(spare):
        i = spare.index(max(spare)); spare[i] -= 1
    put[p] = [a + s for a, s in zip(n, spare)]
    if sum(n) > order[p]: short[p] = sum(n) - order[p]

COL = ['đen', 'nâu', 'đỏ', 'cam', 'vàng', 'lục', 'lam', 'tím', 'xám', 'trắng']
def bands(p):
    m = re.match(r'Điện trở ([\d.]+)(k?)Ω', p)
    if not m: return ''
    ohm = round(float(m.group(1)) * (1000 if m.group(2) else 1))
    e = 0
    while ohm >= 1000: ohm //= 10; e += 1
    d = f'{ohm * 10 if ohm < 100 else ohm:03d}'
    if ohm < 100: e -= 1
    mult = 'nhũ vàng' if e == -1 else COL[e]
    return ' – '.join([COL[int(c)] for c in d] + [mult, 'nâu'])

def short_name(p):
    return (p.replace('Điện trở ', '').replace(' 0410-102', '').replace('IC ', '').replace(' (DIP-8)', ''))

thin = Side(style='thin', color='999999'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
G = PatternFill('solid', fgColor='549E39'); LG = PatternFill('solid', fgColor='E2EFDA'); Y = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='F8CBAD')
W = Font(bold=True, color='FFFFFF'); C = Alignment(horizontal='center', vertical='center', wrap_text=True)
def head(ws, row, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v); c.fill = G; c.font = W; c.alignment = C; c.border = B

wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Tổng hợp'
ws['A1'] = 'Chia linh kiện vào 6 hộp (mỗi cụm máy 1 hộp, dùng chung cho cả 3 lớp, Lab 1–4 + thi)'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = ('Mỗi loại linh kiện 1 túi zip 6×8 cm, dán nhãn (sheet "Nhãn in"). Số trong ô = số con bỏ vào túi đó (đã gồm ~10% dư). '
            'Phần còn lại để nguyên trong túi của shop làm DỰ TRỮ – cụm nào hết thì lấy ở đây. Nền xanh = linh kiện thu hồi cuối mỗi buổi.')
ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:L2'); ws.row_dimensions[2].height = 45
head(ws, 4, ['STT', 'Linh kiện', 'Mã màu (5 vòng)', 'Dùng ở', 'Cụm 1', 'Cụm 2', 'Cụm 3', 'Cụm 4', 'Cụm 5', 'Cụm 6', 'Đã đặt mua', 'Còn lại (dự trữ)'])
row = 5
for i, p in enumerate(put, 1):
    vals = [i, p, bands(p), ', '.join(sorted(labs[p])), *put[p], order[p], order[p] - sum(put[p])]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v); c.border = B; c.alignment = C if j != 2 and j != 3 else Alignment(vertical='center', wrap_text=True)
        if kind[p] == 'Thu hồi': c.fill = LG
    if p in short:
        ws.cell(row, 12).fill = RED; ws.cell(row, 12, f'THIẾU {short[p]}')
    row += 1
n = len(put)
ws.cell(row + 1, 2, 'Số túi 6×8 cm cần dùng').font = Font(bold=True)
ws.cell(row + 1, 5, f'{n} túi/hộp × 6 hộp = {6 * n} túi → mua {math.ceil((6 * n * 1.1) / 50) * 50} túi (dự trữ để nguyên túi của shop)')
ws.cell(row + 2, 2, 'Hộp ngoài').font = Font(bold=True)
ws.cell(row + 2, 5, '6 hộp giấy/nhựa có sẵn, ghi "CỤM 1" … "CỤM 6" + 1 hộp "DỰ TRỮ" (túi của shop)')
for col, w in zip('ABCDEFGHIJKL', [5, 26, 34, 18, 8, 8, 8, 8, 8, 8, 11, 12]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'C5'

# one checklist per bench (print and tape inside the box lid)
for b in range(1, 7):
    s = wb.create_sheet(f'Cụm {b}')
    s['A1'] = f'HỘP CỤM {b} – danh sách túi ({n} túi)'; s['A1'].font = Font(bold=True, size=14)
    head(s, 3, ['✓', 'Linh kiện', 'Mã màu (5 vòng)', 'Số con', 'Dùng ở'])
    for i, p in enumerate(put, 4):
        for j, v in enumerate(['☐', short_name(p), bands(p), put[p][b - 1], ', '.join(sorted(labs[p]))], 1):
            c = s.cell(i, j, v); c.border = B; c.alignment = C if j in (1, 4) else Alignment(vertical='center')
            if kind[p] == 'Thu hồi': c.fill = LG
    for col, w in zip('ABCDE', [5, 22, 38, 8, 18]): s.column_dimensions[col].width = w
    s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 1; s.sheet_properties.pageSetUpPr.fitToPage = True

# labels: 4 per row, one page per bench (cut and stick on the bags)
L = wb.create_sheet('Nhãn in')
L['A1'] = 'Nhãn dán túi – in A4, cắt theo viền. Mỗi cụm 1 trang.'; L['A1'].font = Font(italic=True)
r0 = 2
for b in range(1, 7):
    items = list(put)
    for k, p in enumerate(items):
        rr, cc = r0 + k // 4, 1 + k % 4
        qty = put[p][b - 1] if b != 'DỰ TRỮ' else order[p] - sum(put[p])
        tag = f'CỤM {b}' if b != 'DỰ TRỮ' else 'DỰ TRỮ'
        txt = f'{tag}\n{short_name(p)}\n{bands(p)}\nSL: {qty}'
        c = L.cell(rr, cc, txt); c.alignment = C; c.border = Border(left=Side('dashed'), right=Side('dashed'), top=Side('dashed'), bottom=Side('dashed'))
        c.font = Font(size=10, bold=False)
        if kind[p] == 'Thu hồi': c.fill = LG
        L.row_dimensions[rr].height = 62
    r0 += math.ceil(len(items) / 4)
    L.row_breaks.append(Break(id=r0 - 1))
for col in 'ABCD': L.column_dimensions[col].width = 24
L.page_setup.fitToWidth = 1; L.page_setup.fitToHeight = 0; L.sheet_properties.pageSetUpPr.fitToPage = True
L.print_options.horizontalCentered = True

out = f'{HCM}/docs/HCM_Chia_tui_linh_kien.xlsx'
wb.save(out)
print(out, n, 'types; short:', short)
for p in put: print(f'{p:28s}', put[p], order[p], order[p] - sum(put[p]), bands(p))
